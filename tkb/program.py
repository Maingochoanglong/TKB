"""Đọc chương trình học: cột 'Môn học' và các cột 'Khối <tên>' (tên khối tùy ý, vd "Khối 1", "Khối Lá").

Chương trình học nằm ở sheet "CHƯƠNG TRÌNH HỌC" của file vào (mẫu V8).
Danh sách môn và tên môn lấy nguyên từ file. Khi dựng bài toán, môn nào trùng tên một môn có luật (sheet QUY ĐỊNH
hoặc mặc định trong tkb/config.py, xem `canonical_subject`) thì nhận luật đó, môn khác vẫn được xếp bình thường.
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from . import config
from .staff import InputError, clean_name, find_sheet, grade_key, normalize, parse_grade, subject_key

_GRADE_HEAD = re.compile(r"khối\s+(\S.*)", re.IGNORECASE)


def grade_head(value) -> int | str | None:
    """Tiêu đề cột "Khối <tên>" -> tên khối (staff.parse_grade); cột khác: None."""
    m = _GRADE_HEAD.fullmatch(clean_name(value)) if value is not None else None
    return parse_grade(m.group(1)) if m else None


def grade_columns(ws) -> tuple[int, int, dict[int | str, int]] | None:
    """Dòng tiêu đề, cột Môn học và {khối: cột} của sheet chương trình học; None nếu không có dòng tiêu đề có cột Môn
    học và ít nhất một cột Khối."""
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        subject = None
        cols: dict[int | str, int] = {}
        for cell in row:
            if cell.value is None:
                continue
            if normalize(cell.value) == "môn học":
                subject = cell.column
            elif (grade := grade_head(cell.value)) is not None:
                cols.setdefault(grade, cell.column)
        if subject and cols:
            return row[0].row, subject, cols
    return None


def _rule_names() -> dict[str, str]:
    names = {subject_key(label): s for s, label in config.DISPLAY_NAMES.items()}
    names.update({subject_key(s): s for s in config.rule_subjects()})
    return names


def canonical_subject(name) -> str:
    """Tên môn dùng khi dựng bài toán: tên trong config nếu khớp một môn có luật, không thì giữ tên trong file."""
    return _rule_names().get(subject_key(name), clean_name(name))


def read_program(path: str | Path) -> dict[int | str, dict[str, int]]:
    """Đọc sheet "CHƯƠNG TRÌNH HỌC" của file vào.

    Trả về {khối: {tên môn như trong file: số tiết}}, giữ thứ tự dòng của file.
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = find_sheet(wb, config.PROGRAM_SHEET)
    if ws is None:
        raise InputError(f"File vào thiếu sheet {config.PROGRAM_SHEET} (Môn học, Khối 1, Khối 2...)")
    found = grade_columns(ws)
    if found is None:
        raise InputError(f"Sheet {config.PROGRAM_SHEET} phải có cột 'Môn học' và các cột 'Khối <tên>' (vd 'Khối 1')")
    header_row, subject_col, grade_cols = found

    curriculum: dict[int | str, dict[str, int]] = {g: {} for g in grade_cols}
    seen: dict[str, int] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        subject = ws.cell(r, subject_col).value
        if subject is None or str(subject).strip() == "" or normalize(subject) in ("tổng", "tổng cộng"):
            continue
        name = clean_name(subject)
        key = canonical_subject(subject)
        if key in seen:
            raise InputError(f"Dòng {r}: môn '{name}' bị lặp với dòng {seen[key]} trong chương trình học")
        seen[key] = r
        for grade, col in grade_cols.items():
            value = ws.cell(r, col).value
            value = 0 if value is None or str(value).strip() == "" else value
            try:
                number = float(value)
            except (TypeError, ValueError):
                raise InputError(f"Dòng {r}, Khối {grade}: số tiết không hợp lệ {value!r}") from None
            if number != int(number) or number < 0:
                raise InputError(f"Dòng {r}, Khối {grade}: số tiết phải là số nguyên không âm")
            curriculum[grade][name] = int(number)
    if not any(curriculum.values()):
        raise InputError(f"Sheet {config.PROGRAM_SHEET} chưa có môn nào")
    return curriculum


def subjects_in_order(curriculum: dict[int | str, dict[str, int]]) -> list[str]:
    """Các môn theo thứ tự xuất hiện trong file."""
    return list(dict.fromkeys(s for g in sorted(curriculum, key=grade_key) for s in curriculum[g]))


def missing_rule_subjects(curriculum: dict[int | str, dict[str, int]]) -> list[str]:
    """Môn có luật (sheet QUY ĐỊNH hoặc config) nhưng không có trong chương trình học (thường do gõ khác tên)."""
    present = {canonical_subject(s) for req in curriculum.values() for s in req}
    return [s for s in config.rule_subjects() if s not in present]
