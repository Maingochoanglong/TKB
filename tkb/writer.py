"""Xuất ra Excel: TKB (chỉ các sheet Khối); file thống kê (số tiết từng môn của mỗi giáo viên); file vào
cập nhật.

Style (phông, cỡ chữ, viền, căn lề, chiều cao dòng) chép từ file vào (xem tkb/style.py).
"""
from __future__ import annotations

from collections import Counter, defaultdict
from copy import copy
from pathlib import Path

import re

import openpyxl
from openpyxl.styles import Alignment, PatternFill
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidationList

from . import config
from .solver import Solution, session_of
from .staff import Teacher, _find_columns, class_sort_key, clean_name, find_sheet, grade_of, normalize, staff_sheet
from .style import CellStyle, Style
from .rules import code as rules_code, subject_columns
from . import luat_rieng
from .template import GUIDE_SHEET, write_custom_sheet, write_guide, write_roles_sheet, write_rules_sheet

MAX_DAY_WIDTH = 30  # cột ngày trong TKB: tên dài hơn thì xuống dòng
BLOCK_GAP = 2  # số dòng trống giữa hai lớp (giống template)
LABEL_PAD = 4  # cột BUỔI, TIẾT của TKB: rộng thêm so với chữ dài nhất cho dễ nhìn
HIRE_LABEL = "tuyển thêm"  # tên của người cần tuyển trong file thống kê
CODE_HEADER = "Mã GV"
LOAD_HEADER = "Số Tiết Thực Dạy"
OVERTIME_HEADER = "Số Tiết Bù"
OVERTIME_DETAIL_HEADER = "Môn Dạy Bù"  # file thống kê: số tiết bù từng môn, vd "Tiếng Việt 1, TNXH 2"
SPARE_HEADER = "Số Tiết Dư"  # file vào cập nhật: định mức − thực dạy (người dạy ít hơn định mức)
STATS_SHEET = "Thống kê"
SHORTAGE_SHEET = "Thiếu tiết"
TOTAL_HEADER = "Tổng Tiết"
# File thống kê, trường có lớp ở cơ sở 2: ai dạy ở cả hai cơ sở (các buổi ở cơ sở 2), ai đổi cơ sở trong ngày.
MOVE_HEADERS = ("Buổi Ở Cơ Sở 2", "Đổi Cơ Sở Trong Ngày")
# File thống kê: tô nền cả dòng để biết ai dạy bù (chế độ bù giờ), ai là người cần tuyển (chế độ tuyển thêm).
OVERTIME_FILL = "FFEB9C"  # vàng nhạt
HIRE_FILL = "C6EFCE"  # xanh lá nhạt
OVERTIME_LEGEND = "Dạy bù (vượt định mức)"
HIRE_LEGEND = "Cần tuyển thêm"
# File thống kê và file vào cập nhật: dòng người còn dư tiết (dạy ít hơn định mức).
SPARE_FILL = "DDEBF7"  # xanh dương nhạt
SPARE_LEGEND = "Dạy ít hơn định mức (còn dư tiết)"
# Trong dòng người dạy bù: ô môn có tiết bù tô màu riêng; số tiết bù từng môn ghi ở cột OVERTIME_DETAIL_HEADER.
OVERTIME_CELL_FILL = "F4B183"  # cam
OVERTIME_CELL_LEGEND = "Môn có tiết dạy bù"
# Trường có lớp ở cơ sở 2 (cột Cơ sở 2): TKB tách thành hai file, hậu tố tên file -> lớp ở cơ sở 2?
CAMPUS_FILES = (("diem_chinh", False), ("diem_phu", True))


def teacher_labels(teachers: dict[str, Teacher], with_codes: bool = False) -> dict[str, str]:
    """Chức vụ -> tên hiển thị dưới tên môn trong TKB.

    Ghi tên giáo viên; tên để trống hoặc người cần tuyển thêm thì ghi Mã GV (vd "Bộ Môn 6"), tên trùng
    nhau thì kèm Mã GV. with_codes: thêm một dòng Mã GV dưới tên (file TKB có chức vụ).
    """
    counts = Counter(t.name.strip() for t in teachers.values() if t.name.strip())
    labels = {}
    for title, t in teachers.items():
        name = t.name.strip()
        if not name or t.supplementary:
            labels[title] = t.code
        elif with_codes:
            labels[title] = f"{name}\n{t.code}"
        elif counts[name] > 1:
            labels[title] = f"{name} ({t.code})"
        else:
            labels[title] = name
    return labels


def session_rows() -> list[tuple[config.Session, int]]:
    """Các hàng của bảng TKB: (buổi, tiết trong ngày). Cột TIẾT ghi tiết trong ngày: sáng 1–4, chiều 5–7."""
    sessions: dict[str, config.Session] = {}
    for d in sorted(config.DAY_SESSIONS):
        for s in config.DAY_SESSIONS[d]:
            sessions.setdefault(s.name, s)
    ordered = sorted(sessions.values(), key=lambda s: s.periods[0])
    return [(s, p) for s in ordered for p in s.periods]


def _merge(ws, style: Style, r1: int, c1: int, r2: int, c2: int, value) -> None:
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            style.body_cell(ws, r, c, None, bold=True)
    ws.cell(r1, c1, value)
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def _grade_sheets(wb, solution: Solution, style: Style, with_codes: bool = False,
                  classes: list[str] | None = None) -> None:
    grid = {(l.class_name, l.day, l.period): l for l in solution.lessons}
    days = sorted(config.DAY_SESSIONS)
    rows = session_rows()
    day_periods = {d: {p for s in config.DAY_SESSIONS[d] for p in s.periods} for d in days}
    first_day_col = 4
    problem = solution.problem
    names = teacher_labels(problem.teachers, with_codes)
    header = ["LỚP", "BUỔI", "TIẾT", *[config.DAYS[d].upper() for d in days]]
    # Cột ngày: cùng độ rộng ở mọi sheet Khối, nới theo dòng dài nhất của cả trường (tối đa MAX_DAY_WIDTH).
    texts = {line for les in solution.lessons
             for t in (problem.subject_label(les.subject), names[les.teacher]) for line in t.split("\n")}
    day_width = min(MAX_DAY_WIDTH, max(style.text_width(t) for t in [*header[3:], config.OFF_LABEL, *texts]))
    title = {c: c for c in problem.classes}  # cột LỚP chỉ ghi tên lớp (cơ sở 2 đã tách file riêng)
    chosen = problem.classes if classes is None else classes
    for grade in sorted({grade_of(c) for c in chosen}):
        ws = wb.create_sheet(f"Khối {grade}")
        classes = sorted((c for c in chosen if grade_of(c) == grade), key=class_sort_key)
        widths = [max(style.text_width(t) for t in [header[0], *(title[c] for c in classes)]),
                  max(style.text_width(t) for t in [header[1], *(s.name.upper() for s, _ in rows)]) + LABEL_PAD,
                  style.text_width(header[2]) + LABEL_PAD]
        for i, width in enumerate(widths + [day_width] * len(days), start=1):
            ws.column_dimensions[get_column_letter(i)].width = round(width, 1)
        # In: khổ ngang, co vừa 1 trang theo chiều rộng.
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        top = 1
        for cls in classes:
            ws.row_dimensions[top].height = style.row_height
            for col, text in enumerate(header, start=1):
                style.header_cell(ws, top, col, text)
            first = top + 1
            _merge(ws, style, first, 1, first + len(rows) - 1, 1, title[cls])
            r = first
            for session in dict.fromkeys(s for s, _ in rows):
                n = sum(1 for s, _ in rows if s is session)
                _merge(ws, style, r, 2, r + n - 1, 2, session.name.upper())
                r += n
            for j, (_, period) in enumerate(rows):
                r = first + j
                lines = 2  # môn + tên giáo viên
                style.body_cell(ws, r, 3, period)
                for i, d in enumerate(days):
                    value = None
                    if period not in day_periods[d]:
                        value = config.OFF_LABEL
                    elif (les := grid.get((cls, d, period))) is not None:
                        value = f"{problem.subject_label(les.subject)}\n{names[les.teacher]}"
                        lines = max(lines, style.lines(value, day_width))
                    style.body_cell(ws, r, first_day_col + i, value)
                ws.row_dimensions[r].height = max(style.row_height, style.line_height * lines)
            top = first + len(rows) + BLOCK_GAP


def staff_rows(solution: Solution) -> list[Teacher]:
    """GV thật theo thứ tự file gốc, sau đó GV bổ sung được dùng."""
    real = [t for t in solution.problem.teachers.values() if not t.supplementary]
    real.sort(key=lambda t: (t.row is None, t.row or 0))
    return real + solution.used_supplements()


def campus_paths(path: str | Path, problem) -> list[tuple[Path, list[str] | None]]:
    """Các file TKB cần ghi: (đường dẫn, các lớp; None = cả trường). Trường có lớp ở cơ sở 2 thì tách thành
    <tên>_diem_chinh (lớp cơ sở 1) và <tên>_diem_phu (lớp cơ sở 2); cơ sở nào không có lớp thì không có file."""
    path = Path(path)
    if not problem.campus2:
        return [(path, None)]
    out = []
    for suffix, at2 in CAMPUS_FILES:
        classes = [c for c in problem.classes if (c in problem.campus2) == at2]
        if classes:
            out.append((path.with_name(f"{path.stem}_{suffix}{path.suffix}"), classes))
    return out


def write_timetable(solution: Solution, path: str | Path, style: Style | None = None,
                    with_codes: bool = False, classes: list[str] | None = None) -> None:
    """File TKB: chỉ các sheet Khối. Nhân sự và thống kê ghi ở file thống kê (write_statistics).
    with_codes: mỗi ô thêm dòng Mã GV (chức vụ) dưới tên giáo viên. classes: chỉ ghi các lớp này (vd các lớp
    của một cơ sở, xem campus_paths); None = cả trường. Độ rộng cột ngày vẫn tính theo cả trường."""
    style = style or Style()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _grade_sheets(wb, solution, style, with_codes, classes)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _stats_name(t: Teacher) -> str:
    return HIRE_LABEL if t.supplementary else t.name or None


def campus_moves(solution: Solution) -> dict[str, tuple[list[str], list[str]]]:
    """GV dạy ở cả hai cơ sở -> (các buổi ở cơ sở 2, vd "Sáng T3"; các ngày sáng một cơ sở, chiều cơ sở kia,
    vd "T5: sáng cơ sở 1, chiều cơ sở 2"). Luật cứng: mỗi buổi chỉ một cơ sở."""
    campus2 = solution.problem.campus2
    sess = session_of()
    where = defaultdict(set)  # (GV, ngày, buổi) -> {lớp ở cơ sở 2?}
    for les in solution.lessons:
        where[les.teacher, les.day, sess[les.day, les.period]].add(les.class_name in campus2)
    campuses = defaultdict(set)
    for (g, _, _), cs in where.items():
        campuses[g] |= cs
    day = lambda d: config.DAYS[d].replace("Thứ ", "T")  # noqa: E731
    out = {}
    for g in sorted(g for g, cs in campuses.items() if len(cs) > 1):
        keys = sorted((d, config.DAY_SESSIONS[d].index(s), s) for (h, d, s) in where if h == g)
        at2 = [f"{s.name} {day(d)}" for d, _, s in keys if True in where[g, d, s]]
        switches = []
        for d in sorted({d for d, _, _ in keys}):
            parts = [(s.name.lower(), 2 if True in where[g, d, s] else 1) for dd, _, s in keys if dd == d]
            if len({c for _, c in parts}) > 1:
                switches.append(f"{day(d)}: " + ", ".join(f"{n} cơ sở {c}" for n, c in parts))
        out[g] = (at2, switches)
    return out


def subject_table(solution: Solution, style: Style) -> tuple[list[str], list[list]]:
    """Họ và Tên | Chức Vụ (Mã GV) | số tiết từng môn | Tổng Tiết | Số Tiết/Tuần | Số Tiết Bù | (Môn Dạy Bù) |
    Số Tiết Dư, mỗi giáo viên một dòng; cuối bảng có dòng Tổng. Chỉ chữ và số, không ghi chú, không công thức.

    Chỉ có cột cho các môn có người dạy, theo thứ tự môn trong chương trình học; ô trống là không dạy môn đó.
    Số Tiết/Tuần là định mức (người cần tuyển: định mức tuyển); Số Tiết Bù là số tiết dạy bù vượt định mức
    (chế độ bù giờ); Môn Dạy Bù (chỉ khi có tiết bù) ghi số tiết bù từng môn; Số Tiết Dư là định mức − Tổng Tiết
    khi dạy ít hơn định mức. Trường có lớp ở cơ sở 2: thêm hai
    cột MOVE_HEADERS cho người dạy ở cả hai cơ sở (xem campus_moves).
    """
    problem = solution.problem
    count = Counter((les.teacher, les.subject) for les in solution.lessons)
    order = {s: i for i, s in enumerate(problem.subject_order)}
    subjects = sorted({s for _, s in count}, key=lambda s: (order.get(s, len(order)), s))
    extra = overtime_cells(solution)
    detail = [OVERTIME_DETAIL_HEADER] if extra else []
    header = [style.staff_headers["name"], "Chức Vụ", *(problem.subject_label(s) for s in subjects), TOTAL_HEADER,
              style.staff_headers["lessons"], OVERTIME_HEADER, *detail, SPARE_HEADER]
    moves = campus_moves(solution) if problem.campus2 else None
    if moves is not None:
        header += MOVE_HEADERS
    text = {OVERTIME_DETAIL_HEADER, *MOVE_HEADERS}  # cột chữ: dòng Tổng không cộng
    overtime = solution.overtime()
    rows = []
    for t in staff_rows(solution):
        per = [count[t.title, s] for s in subjects]
        total = sum(per)
        subject_extra = ", ".join(f"{problem.subject_label(s)} {extra[t.title, s]}" for s in subjects
                                  if extra.get((t.title, s)))
        rows.append([_stats_name(t), t.code, *(n or None for n in per), total, t.max_lessons,
                     overtime.get(t.title) or None, *([subject_extra or None] if extra else []),
                     max(0, t.max_lessons - total) or None])
        if moves is not None:
            at2, switches = moves.get(t.title, ([], []))
            rows[-1] += [", ".join(at2) or None, "; ".join(switches) or None]
    summary = {OVERTIME_DETAIL_HEADER: f"{len(extra)} ô"}
    if moves is not None:
        summary.update({MOVE_HEADERS[0]: f"{len(moves)} người",
                        MOVE_HEADERS[1]: f"{sum(len(sw) for _, sw in moves.values())} lần"})
    rows.append(["Tổng", None, *(summary[h] if h in text else sum(r[c] or 0 for r in rows)
                                 for c, h in enumerate(header) if c >= 2)])
    return header, rows


def _fill(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def row_marks(solution: Solution, spare: bool = False) -> dict[str, tuple[str, int]]:
    """GV -> (màu nền, số tiết): người dạy bù (số tiết bù) và người cần tuyển thêm (số tiết thực dạy); spare: thêm
    người còn dư tiết (số tiết dư), trừ hai nhóm trên."""
    load = solution.teacher_load()
    marks = {}
    if spare:
        marks = {t.title: (SPARE_FILL, t.max_lessons - load[t.title]) for t in staff_rows(solution)
                 if not t.supplementary and t.max_lessons > load[t.title]}
    marks.update({g: (OVERTIME_FILL, n) for g, n in solution.overtime().items()})
    marks.update({t.title: (HIRE_FILL, load[t.title]) for t in solution.used_supplements()})
    return marks


def overtime_cells(solution: Solution) -> Counter:
    """(GV, môn) -> số tiết dạy bù (vượt định mức) của GV đó trong môn đó (chế độ bù giờ)."""
    return Counter((les.teacher, les.subject) for les in solution.lessons if les.overtime)


def _mark_rows(ws, solution: Solution, header: list[str], top: int, style: Style) -> list[int]:
    """Tô nền dòng người dạy bù, người cần tuyển và người còn dư tiết (bảng bắt đầu ở dòng 1); trong dòng người
    dạy bù, ô môn có tiết bù tô màu riêng (số tiết ở cột Môn Dạy Bù). Dưới bảng, từ dòng `top`, ghi chú thích màu
    kèm số người, số tiết (chữ thường, không dùng ghi chú trong ô). Trả về các dòng chú thích."""
    marks = row_marks(solution, spare=True)
    width = len(header)
    extra = overtime_cells(solution)
    labels = {solution.problem.subject_label(s): s for s in {s for _, s in extra}}
    n_cells = 0
    for r, t in enumerate(staff_rows(solution), start=2):
        if t.title in marks:
            for c in range(1, width + 1):
                ws.cell(r, c).fill = _fill(marks[t.title][0])
        for c, head in enumerate(header, start=1):
            n = extra.get((t.title, labels.get(head)), 0)
            if n:
                ws.cell(r, c).fill = _fill(OVERTIME_CELL_FILL)
                n_cells += 1
    texts = []
    for color, label in ((OVERTIME_FILL, OVERTIME_LEGEND), (HIRE_FILL, HIRE_LEGEND), (SPARE_FILL, SPARE_LEGEND)):
        lessons = [n for mark, n in marks.values() if mark == color]
        if lessons:
            texts.append((color, f"{label}: {len(lessons)} người, {sum(lessons)} tiết"))
        if color == OVERTIME_FILL and n_cells:
            texts.append((OVERTIME_CELL_FILL, f"{OVERTIME_CELL_LEGEND}: {n_cells} ô, {sum(extra.values())} tiết "
                                              f"(số tiết từng môn ở cột {OVERTIME_DETAIL_HEADER})"))
    moves = campus_moves(solution) if solution.problem.campus2 else {}
    if moves:
        texts.append((None, f"Dạy ở cả hai cơ sở: {len(moves)} người (cột {MOVE_HEADERS[0]}); đổi cơ sở trong "
                            f"ngày: {sum(len(sw) for _, sw in moves.values())} lần (cột {MOVE_HEADERS[1]})"))
    legend = []
    for color, text in texts:
        cell = style.body_cell(ws, top, 1, None)
        if color:
            cell.fill = _fill(color)
        cell = ws.cell(top, 2, text)
        cell.font = copy(style.body.font)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[top].height = style.row_height
        legend.append(top)
        top += 1
    return legend


def write_statistics(solution: Solution, path: str | Path, style: Style | None = None) -> None:
    """File thống kê: một bảng số tiết từng môn của mỗi giáo viên, kèm định mức, số tiết bù, số tiết dư (xem
    subject_table). Dòng người dạy bù tô vàng (ô môn có tiết bù tô cam, số tiết bù từng môn ở cột Môn Dạy Bù),
    dòng người cần tuyển tô xanh lá, dòng người còn dư tiết tô xanh dương; chú thích dưới bảng. Chỉ chữ, số và
    màu: không cố định dòng/cột, không ghi chú, không công thức."""
    style = style or Style()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = STATS_SHEET
    header, rows = subject_table(solution, style)
    last = style.table(ws, header, rows, bold_last=True)
    legend = _mark_rows(ws, solution, header, last + 2, style)
    style.fit_columns(ws, skip_rows=legend)  # chú thích tràn sang các ô trống bên phải, không nới cột
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_shortage(rows: list[tuple[str, str, int, str]], path: str | Path, style: Style | None = None) -> None:
    """File thống kê khi chế độ bù giờ không đủ: sheet SHORTAGE_SHEET liệt kê các tiết không ai dạy được."""
    style = style or Style()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHORTAGE_SHEET
    body = [list(r) for r in rows] + [["Tổng", None, sum(r[2] for r in rows), None]]
    style.table(ws, ["Lớp", "Môn", "Số Tiết Thiếu", "Lý Do"], body, bold_last=True)
    style.fit_columns(ws)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _copy_style(src, dst) -> None:
    if src.has_style:
        dst._style = copy(src._style)


OLD_NOTES_SHEET = "Chú thích"  # sheet giải thích của bản trước, nay gộp vào sheet HƯỚNG DẪN
LIST_SHEET = "Danh mục"  # sheet danh mục ẩn của file mẫu cũ (danh sách thả xuống), bỏ khi chép file vào
_ROW_FORMULA = re.compile(r"^=\s*ROW\(\)\s*([+-])\s*(\d+)\s*$", re.IGNORECASE)


def plain_values(wb, cached) -> None:
    """Chỉ giữ chữ, số và màu: bỏ mọi ghi chú (comment), cố định dòng/cột, lọc, danh sách thả xuống, định dạng theo
    điều kiện và sheet danh mục ẩn của file mẫu (LIST_SHEET); đổi mọi công thức thành giá trị: giá trị Excel đã lưu
    trong file (`cached`: cùng file mở với data_only=True); file chưa từng mở bằng Excel thì không có giá trị lưu
    sẵn, khi đó công thức dạng =ROW()±k (hay dùng cho cột STT) được tự tính, công thức khác để trống."""
    if LIST_SHEET in wb.sheetnames and wb[LIST_SHEET].sheet_state != "visible":
        del wb[LIST_SHEET]
    for ws in wb.worksheets:
        ws.freeze_panes = None
        ws.auto_filter.ref = None
        ws.data_validations = DataValidationList()
        ws.conditional_formatting = ConditionalFormattingList()
        for row in ws.iter_rows():
            for cell in row:
                cell.comment = None
                if cell.data_type != "f":
                    continue
                value = cached[ws.title][cell.coordinate].value
                if value is None and (m := _ROW_FORMULA.match(str(cell.value))):
                    value = cell.row + int(m.group(2)) * (1 if m.group(1) == "+" else -1)
                cell.value = value


def _result_notes(overtime_mode: bool) -> list[tuple[str, str]]:
    """Các dòng sheet HƯỚNG DẪN giải thích phần kết quả của file vào cập nhật."""
    sheet = config.STAFF_SHEET
    return [(f"{sheet}: {CODE_HEADER}", "Kết quả: mã giáo viên dùng trong TKB và file thống kê."),
            (f"{sheet}: {LOAD_HEADER}", "Kết quả: số tiết thực dạy trong tuần."),
            *([(f"{sheet}: {OVERTIME_HEADER}", "Kết quả: số tiết dạy vượt định mức.")] if overtime_mode else []),
            (f"{sheet}: {SPARE_HEADER}", "Kết quả: định mức trừ số tiết thực dạy, khi dạy ít hơn định mức."),
            (f"{sheet}: màu dòng", f"Kết quả: vàng là {OVERTIME_LEGEND.lower()}; xanh lá là {HIRE_LEGEND.lower()}; "
                                   f"xanh dương là {SPARE_LEGEND.lower()}. Các cột kết quả và màu được ghi lại mỗi "
                                   f"lần chạy, không cần xóa."),
            (config.SAVED_SHEET, f"Kết quả: TKB đã xếp, dạng lưới như TKB; mỗi ô ghi môn, xuống dòng ghi Mã GV (thêm "
                                 f"{config.SAVED_OVERTIME} ở tiết dạy bù); dòng đầu ghi mã kết quả và mã quy định. Nạp "
                                 f"lại file này làm file vào (vd chỉ đổi tên người \"chưa có\" thành tên người mới "
                                 f"tuyển) thì chương trình giữ nguyên TKB nếu vẫn đúng mọi luật và quy định không đổi; "
                                 f"muốn xếp lại từ đầu thì đặt GIU_TKB_DA_XEP = False trong main.py.")]


def _add_subject_rules(wb) -> None:
    """Sheet CHƯƠNG TRÌNH HỌC của file vào cập nhật ghi đủ các quy định của môn đã dùng: thêm các cột quy định còn
    thiếu, và các môn có quy định mà sheet chưa có (số tiết để trống). Dòng, cột mới chép style của file vào."""
    ws = find_sheet(wb, config.PROGRAM_SHEET)
    if ws is None:
        return
    found = next(((r, c) for r in range(1, min(ws.max_row, 20) + 1) for c in range(1, ws.max_column + 1)
                  if normalize(ws.cell(r, c).value or "") == "môn học"), None)
    if found is None:
        return
    header_row, key_col = found
    rows = {r: str(ws.cell(r, key_col).value) for r in range(header_row + 1, ws.max_row + 1)
            if ws.cell(r, key_col).value not in (None, "")
            and normalize(ws.cell(r, key_col).value) not in ("tổng", "tổng cộng")}
    heads, values, rest = subject_columns(list(rows.values()))
    existing = {normalize(ws.cell(header_row, c).value): c for c in range(1, ws.max_column + 1)
                if ws.cell(header_row, c).value not in (None, "")}
    new_cols = []
    for h in heads:
        if normalize(h) not in existing:
            c = ws.max_column + 1
            existing[normalize(h)] = c
            new_cols.append(c)
            ws.cell(header_row, c, h)
            _copy_style(ws.cell(header_row, key_col), ws.cell(header_row, c))
            ws.column_dimensions[get_column_letter(c)].width = max(10, len(h) + 4)
    last = max(rows, default=header_row)
    new_rows = []
    for s in rest:
        last += 1
        new_rows.append(last)
        rows[last] = s
        ws.cell(last, key_col, s)
    src = max((r for r in rows if r not in new_rows), default=header_row)  # dòng mẫu để chép style
    for r, name in rows.items():
        for i, h in enumerate(heads):
            c = existing[normalize(h)]
            if r in new_rows or c in new_cols:
                ws.cell(r, c).value = values[clean_name(name)][i]
                _copy_style(ws.cell(src, key_col + 1 if c in new_cols else c), ws.cell(r, c))
        if r in new_rows:
            for c in range(1, ws.max_column + 1):
                if c == key_col or c not in existing.values():
                    _copy_style(ws.cell(src, c), ws.cell(r, c))
            if ws.row_dimensions[src].height:
                ws.row_dimensions[r].height = ws.row_dimensions[src].height


def _write_saved(wb, solution: Solution, style: Style) -> None:
    """Sheet config.SAVED_SHEET: TKB đã xếp dạng lưới như TKB (Lớp | Tiết | Thứ 2 …), mỗi ô ghi môn, xuống dòng ghi Mã
    GV (thêm config.SAVED_OVERTIME ở tiết dạy bù); dòng đầu ghi mã kết quả và mã các quy định đã dùng. Nạp lại file
    vào cập nhật thì chương trình dùng lại TKB này (solver.reuse) nếu quy định không đổi."""
    if config.SAVED_SHEET in wb.sheetnames:
        del wb[config.SAVED_SHEET]
    ws = wb.create_sheet(config.SAVED_SHEET)
    problem = solution.problem
    result, rules = config.SAVED_CODES
    for c, value in enumerate((result, solution.fingerprint(), rules, rules_code()), start=1):
        ws.cell(1, c, value).font = copy(style.body.font)
    days = sorted(config.DAY_SESSIONS)
    periods = [p for _, p in session_rows()]
    day_periods = {d: {p for s in config.DAY_SESSIONS[d] for p in s.periods} for d in days}
    header = ["Lớp", "Tiết", *[config.DAYS[d] for d in days]]
    top = 3
    for c, text in enumerate(header, start=1):
        style.header_cell(ws, top, c, text)
    grid = {(l.class_name, l.day, l.period): l for l in solution.lessons}
    r = top
    for cls in sorted(problem.classes, key=class_sort_key):
        for i, p in enumerate(periods):
            r += 1
            style.body_cell(ws, r, 1, cls if i == 0 else None)
            style.body_cell(ws, r, 2, p)
            for j, d in enumerate(days, start=3):
                les = grid.get((cls, d, p))
                value = config.OFF_LABEL if p not in day_periods[d] else None
                if les is not None:
                    code = problem.teachers[les.teacher].code
                    value = f"{problem.subject_label(les.subject)}\n{code}" + \
                        (f" {config.SAVED_OVERTIME}" if les.overtime else "")
                cell = style.body_cell(ws, r, j, value)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[r].height = style.line_height * 2 + 4
    for c, width in enumerate((10, 6, *[24] * len(days)), start=1):
        ws.column_dimensions[get_column_letter(c)].width = width


def role_rows(solution: Solution) -> list[list[str]]:
    """Các dòng sheet CHỨC VỤ cho file vào chưa có sheet này: các chức vụ GV chuyên biệt có người giữ (cả người cần
    tuyển), mỗi dòng [tên chức vụ, các môn được dạy]. Đó là các chức vụ trùng tên môn: ghi ra thì như không ghi
    (rules.code không đổi), chỉ để nhà trường thấy mỗi chức vụ dạy môn nào."""
    problem = solution.problem
    held = {t.role: t.label or t.role for t in reversed(staff_rows(solution)) if t.role in problem.specialists}
    return [[held[role], ", ".join(problem.subject_labels.get(s, s) for s in subjects)]
            for role, subjects in problem.specialists.items() if role in held]


def write_updated_staff(solution: Solution, source: str | Path, path: str | Path) -> None:
    """Chép file vào, thêm người cần tuyển vào cuối danh sách nhân sự và các cột Mã GV, số tiết thực dạy
    (số tiết bù ở chế độ bù giờ) và số tiết dư. Dòng, cột mới chép style của file vào; đọc lại file này làm file
    vào vẫn được, và sheet config.SAVED_SHEET lưu TKB đã xếp để lần nạp lại giữ nguyên TKB (solver.reuse). Đây là
    bản thống kê gọn theo mẫu file vào: tô nền cả dòng người dạy bù (vàng), người cần tuyển
    (xanh lá), người còn dư tiết (xanh dương); sheet HƯỚNG DẪN (ghi lại mỗi lần) giải thích cả các cột kết quả và
    màu. File cũng ghi đủ các quy định đã dùng (cột quy định của sheet CHƯƠNG TRÌNH HỌC, sheet QUY ĐỊNH). File ra chỉ
    có chữ, số và màu (plain_values): không ghi chú, không công thức, không cố định dòng/cột, không danh sách thả
    xuống.

    Mẫu V8 (có cột Lớp) ghi chức vụ không kèm số thứ tự, vd "Bộ Môn": khi đọc lại, chương trình tự
    đánh số tiếp theo (bộ môn 6, 7...). Mẫu cũ ghi đủ chức vụ, vd "bộ môn 6".
    """
    wb = openpyxl.load_workbook(source)
    plain_values(wb, openpyxl.load_workbook(source, data_only=True))
    ws = staff_sheet(wb)
    header_row, cols = _find_columns(ws)
    title_col = cols["title"]
    last = max((r for r in range(header_row + 1, ws.max_row + 1)
                if ws.cell(r, title_col).value not in (None, "")), default=header_row)
    v8 = "class" in cols
    by_row = {t.row: t for t in solution.problem.teachers.values() if not t.supplementary and t.row}
    for i, t in enumerate(solution.used_supplements(), start=1):
        r = last + i
        for c in range(1, ws.max_column + 1):
            _copy_style(ws.cell(last, c), ws.cell(r, c))
        if ws.row_dimensions[last].height:
            ws.row_dimensions[r].height = ws.row_dimensions[last].height
        if "stt" in cols:  # STT: số của dòng trên + 1 (công thức đã đổi thành giá trị)
            above = ws.cell(r - 1, cols["stt"]).value
            ws.cell(r, cols["stt"], above + 1 if isinstance(above, int) else None)
        ws.cell(r, cols["name"], t.name)
        ws.cell(r, title_col, t.label if v8 else t.title)
        ws.cell(r, cols["lessons"], t.max_lessons)
        by_row[r] = t

    # Cột kết quả: ghi đè nếu file đã có (chạy lại trên file cập nhật), không thì thêm vào bên phải.
    load = solution.teacher_load()
    overtime = solution.overtime()
    spare = lambda t: max(0, t.max_lessons - load[t.title]) or None  # noqa: E731
    values = {CODE_HEADER: lambda t: t.code, LOAD_HEADER: lambda t: load[t.title],
              OVERTIME_HEADER: lambda t: overtime.get(t.title), SPARE_HEADER: spare}
    wanted = [CODE_HEADER, LOAD_HEADER, *([OVERTIME_HEADER] if solution.problem.overtime_mode() else []),
              SPARE_HEADER]
    headers = {normalize(ws.cell(header_row, c).value): c for c in range(1, ws.max_column + 1)
               if ws.cell(header_row, c).value not in (None, "")}
    next_col = max(headers.values(), default=0) + 1
    style = Style(body=CellStyle.of(ws.cell(header_row, title_col)))
    for h in values:
        c = headers.get(normalize(h))
        if c is None and h not in wanted:
            continue
        if c is None:
            c, next_col = next_col, next_col + 1
        ws.cell(header_row, c, h if h in wanted else None)
        _copy_style(ws.cell(header_row, title_col), ws.cell(header_row, c))
        for r in range(header_row + 1, max(ws.max_row, last) + 1):
            t = by_row.get(r)
            ws.cell(r, c, values[h](t) if t is not None and h in wanted else None)
            if t is not None:
                _copy_style(ws.cell(r, title_col), ws.cell(r, c))
        letter = get_column_letter(c)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0,
                                                 round(style.text_width(h), 1))
    # Tô nền cả dòng (đến cột tiêu đề cuối); chạy lại trên file cập nhật thì bỏ màu cũ của chương trình.
    width = max(c for c in range(1, ws.max_column + 1) if ws.cell(header_row, c).value not in (None, ""))
    marks = {g: color for g, (color, _) in row_marks(solution).items()}
    ours = {OVERTIME_FILL, HIRE_FILL, SPARE_FILL}
    for r, t in by_row.items():
        color = marks.get(t.title) or (SPARE_FILL if spare(t) else None)
        for c in range(1, width + 1):
            cell = ws.cell(r, c)
            if color:
                cell.fill = _fill(color)
            elif cell.fill.fill_type == "solid" and str(cell.fill.start_color.rgb)[-6:] in ours:
                cell.fill = PatternFill()
    # Ghi đủ các quy định đã dùng: cột quy định của môn, sheet CHỨC VỤ, QUY ĐỊNH (nếu file vào chưa có, sau sheet
    # chương trình học); sheet HƯỚNG DẪN viết lại, kèm giải thích phần kết quả; cuối cùng là TKB đã xếp.
    _add_subject_rules(wb)
    program = find_sheet(wb, config.PROGRAM_SHEET)
    after = wb.worksheets.index(program) + 1 if program is not None else None
    if find_sheet(wb, config.ROLES_SHEET) is None:
        write_roles_sheet(wb, role_rows(solution), after)
        after = after and after + 1
    if find_sheet(wb, config.RULES_SHEET) is None:
        write_rules_sheet(wb, after)
    if find_sheet(wb, luat_rieng.SHEET) is None:  # sheet LUẬT RIÊNG trống để nhà trường biết mà điền
        write_custom_sheet(wb, index=wb.worksheets.index(find_sheet(wb, config.RULES_SHEET)) + 1)
    for name in (GUIDE_SHEET, OLD_NOTES_SHEET):
        old = find_sheet(wb, name)
        if old is not None:
            del wb[old.title]
    write_guide(wb, _result_notes(solution.problem.overtime_mode()))
    _write_saved(wb, solution, Style.from_file(source))
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
