"""Máy chủ HTTP của giao diện (chỉ thư viện chuẩn) và việc xếp TKB ở tiến trình con.

An toàn khi chạy trên máy của trường:
- chỉ nghe trên 127.0.0.1 và chỉ nhận yêu cầu có Host là 127.0.0.1/localhost (trang web lạ không gọi được);
- mọi lệnh /api/ cần mã phiên (X-TKB-Token, sinh mới mỗi lần mở) nên trang web khác không gửi lệnh thay được;
- chỉ mở, ghi các file trong thư mục TKB.

Mọi file trang đọc, ghi đều ở thư mục TKB (thư mục của chương trình; bản đóng gói: thư mục có TKB.exe; xem FOLDERS,
docs/Thiet_Ke_Ba_Thuc_The.md 2.1): ba thực thể chương trình học, nhân sự, luật mỗi thực thể một thư mục; file lấy từ
chỗ khác được chép bản gốc vào tep_nhap/ rồi mới đọc; kết quả ở ket_qua/; bản nháp của trang, file tạm và hồ sơ trình
duyệt ở du_lieu/ (xem data_dir, serve). Xóa thư mục là hết.

Xếp TKB chạy `python -m tkb <file vào> ...` (đúng như main.py) ở tiến trình con: giao diện không treo, in ra từng
dòng như khi chạy dòng lệnh, và nút Dừng gửi Ctrl+C (Windows: Ctrl+Break) để chương trình dừng sớm mà vẫn ghi TKB tốt
nhất (tkb/lns.py).
"""
from __future__ import annotations

import json
import os
import re
import secrets
import signal
import subprocess
import sys
import tempfile
import threading
import shutil
import time
import webbrowser
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse

from .. import config, kich_ban, nhap_tkb
from ..staff import InputError

STATIC = Path(__file__).resolve().parent / "static"
PROJECT = Path(__file__).resolve().parents[2]  # thư mục dự án khi chạy từ mã nguồn
PORT = 8765
TOKEN_MARK = "__TKB_TOKEN__"  # chỗ ghi mã phiên trong index.html
MAX_BODY = 20 * 1024 * 1024
TYPES = {".html": "text/html; charset=utf-8", ".css": "text/css; charset=utf-8",
         ".js": "text/javascript; charset=utf-8", ".svg": "image/svg+xml"}
MODES = [{"key": config.MODE_OVERTIME, "label": "Bù giờ",
          "note": "GVCN và bộ môn dạy bù vượt định mức; bù tối đa vẫn thiếu thì báo lỗi, ghi bảng tiết thiếu."},
         {"key": config.MODE_HIRE, "label": "Tuyển thêm",
          "note": f"Thêm GV \"{config.SUPPLEMENT_NAME}\" dạy đúng các ô mà ở chế độ bù là tiết bù (cùng TKB)."}]
# Cài đặt chạy mặc định (như main.py; bù tối đa theo mức được duyệt trong tkb/config.py).
RUN_DEFAULTS = {"name": "kich_ban", "mode": config.MODE_OVERTIME, "overtime_max": config.OVERTIME_MAX,
                "student_rules": True, "time_limit": 1200, "reproducible": True, "keep_saved": True, "workers": 8}
FILE_NAMES = {"tkb": "TKB.xlsx", "roles": "TKB_chuc_vu.xlsx", "teachers": "TKB_giao_vien.xlsx",
              "school": "TKB_toan_truong.xlsx", "stats": "Thong_Ke.xlsx"}
# Thư mục con của thư mục TKB: khóa -> tên thư mục. Ba thư mục đầu là ba thực thể (kich_ban.ENTITIES).
FOLDERS = {"program": "chuong_trinh_hoc", "staff": "nhan_su", "rules": "luat", "imported": "tep_nhap",
           "results": "ket_qua"}
FOLDER_LABELS = {"program": "Chương trình học", "staff": "Nhân sự", "rules": "Luật",
                 "imported": "File lấy từ chỗ khác (bản gốc)", "results": "Kết quả (file vào đủ, file cập nhật)"}
DATA_FOLDER = "du_lieu"  # của riêng chương trình: bản nháp của trang, file tạm, hồ sơ trình duyệt
# File mẫu trống của từng thực thể (python -m tkb.template chia theo thực thể); mẫu luật: các luật có sẵn.
TEMPLATE_NAMES = {"program": "Mau_Chuong_Trinh_Hoc.xlsx", "staff": "Mau_Nhan_Su.xlsx", "rules": "Mau_Luat.xlsx"}
RULES_TEMPLATE_NAME = TEMPLATE_NAMES["rules"]


STORE_NAME = "trang.json"  # bản nháp và lựa chọn của trang (thay cho bộ nhớ trình duyệt), trong data_dir


def app_dir() -> Path:
    """Thư mục TKB, của chương trình: bản đóng gói là thư mục có TKB.exe; chạy từ mã nguồn là out/ của dự án (đã bỏ
    qua trong git vì có tên giáo viên)."""
    if getattr(sys, "frozen", False) or not (PROJECT / "main.py").is_file():
        return Path(sys.executable).resolve().parent
    return PROJECT / "out"


def default_out_dir() -> Path:
    """Thư mục kết quả: ket_qua/ của thư mục TKB."""
    return app_dir() / FOLDERS["results"]


def data_dir(home=None) -> Path:
    """Chỗ để bản nháp của trang, file tạm và hồ sơ trình duyệt: du_lieu/ của thư mục TKB (`home`; không ghi: thư
    mục của chương trình)."""
    return Path(home or app_dir()) / DATA_FOLDER


def cli_command(argv: list[str]) -> list[str]:
    """Lệnh chạy `python -m tkb`; bản đóng gói (PyInstaller) gọi lại chính nó với --cli (xem __main__.py)."""
    if getattr(sys, "frozen", False):
        return [sys.executable, "--cli", *argv]
    return [sys.executable, "-m", "tkb", *argv]


def safe_name(name) -> str:
    """Tên file vào: bỏ ký tự Windows không cho phép, bỏ đuôi .xlsx."""
    text = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", str(name or "")).strip().strip(".")
    text = re.sub(r"\.xlsx$", "", text, flags=re.IGNORECASE).strip()
    return text[:80] or RUN_DEFAULTS["name"]


def run_argv(source: Path, out: Path, name: str, run: dict) -> list[str]:
    """Tham số dòng lệnh `python -m tkb` theo cài đặt chạy (như main.run); báo lỗi nếu cài đặt sai."""
    mode = run.get("mode", RUN_DEFAULTS["mode"])
    if mode not in config.MODES:
        raise InputError(f"Chế độ phải là một trong {', '.join(config.MODES)}")
    numbers = {}
    for key, low, kind in (("overtime_max", 0, int), ("time_limit", 0, float), ("workers", 1, int)):
        value = run.get(key, RUN_DEFAULTS[key])
        value = 0 if key == "time_limit" and value in (None, "") else value
        try:
            number = kind(value)
        except (TypeError, ValueError):
            number = None
        if number is None or isinstance(value, bool) or number < low or (kind is int and number != float(value)):
            raise InputError({"overtime_max": "Số tiết bù tối đa phải là số nguyên >= 0",
                              "time_limit": "Thời gian phải là số >= 0 (0 = không giới hạn)",
                              "workers": "Số luồng phải là số nguyên >= 1"}[key])
        numbers[key] = number
    argv = [str(source), "-o", str(out / FILE_NAMES["tkb"]), "--roles-out", str(out / FILE_NAMES["roles"]),
            "--teachers-out", str(out / FILE_NAMES["teachers"]), "--school-out", str(out / FILE_NAMES["school"]),
            "--stats-out", str(out / FILE_NAMES["stats"]), "--staff-out", str(out / f"{name}_cap_nhat.xlsx"),
            "--mode", mode, "--max-overtime", str(numbers["overtime_max"]),
            "--time-limit", f"{numbers['time_limit']:g}", "--workers", str(numbers["workers"])]
    if not run.get("student_rules", True):
        argv.append("--no-student-rules")
    if not run.get("reproducible", True):
        argv.append("--non-reproducible")
    if not run.get("keep_saved", True):
        argv.append("--xep-lai")
    return argv


def summary(lines: list[str]) -> dict:
    """Kết quả đọc từ các dòng in ra (tkb/__main__.py): mã kết quả, đạt luật bắt buộc hay không, các file đã ghi, lỗi
    (từ dòng "LỖI:" đầu tiên đến hết, trừ các dòng "Đã ghi"), và các số để trang kết quả tóm tắt: vì sao dừng, giây
    xếp giờ, dùng lại TKB cũ, số ô đổi so với TKB cũ, người cần tuyển, tiết dạy bù, các mục tiêu mềm in ra."""
    out = {"code": None, "passed": None, "files": [], "error": None, "stop": None, "seconds": None, "reused": False,
           "changes": None, "hires": None, "overtime": None, "soft": []}
    first = next((i for i, line in enumerate(lines) if line.startswith("LỖI:")), None)
    if first is not None:
        out["error"] = "\n".join(line for line in lines[first:] if not line.startswith("Đã ghi:")).strip()
    for i, line in enumerate(lines):
        if m := re.match(r"Mã kết quả: (\S+)", line):
            out["code"] = m.group(1)
            if s := re.search(r"xếp giờ mất (\d+) giây", line):
                out["seconds"] = int(s.group(1))
        if "kiểm tra luật bắt buộc:" in line:
            out["passed"] = line.rstrip().endswith("ĐẠT") and "KHÔNG ĐẠT" not in line
        if m := re.match(r"Đã ghi: (.+)", line):
            out["files"].append(m.group(1).strip())
        if m := re.match(r"\s*Dừng xếp giờ: (.+)", line):
            out["stop"] = m.group(1).strip()
        if line.startswith("Dùng lại TKB đã xếp"):
            out["reused"] = True
        if m := re.match(r"So với TKB đã xếp trong file vào: đổi (\d+)/(\d+) ô", line):
            out["changes"] = [int(m.group(1)), int(m.group(2))]
        if m := re.match(r"Cần bổ sung (\d+) GV cho (\d+) tiết", line):
            out["hires"] = [int(m.group(1)), int(m.group(2))]
        elif line.startswith("Không cần bổ sung giáo viên"):
            out["hires"] = [0, 0]
        if m := re.match(r"Dạy bù (\d+) tiết: GVCN (\d+) tiết \((\d+) người\), bộ môn (\d+) tiết \((\d+) người\)", line):
            levels = lines[i + 1].strip() if i + 1 < len(lines) and "người bù +" in lines[i + 1] else ""
            out["overtime"] = {"total": int(m.group(1)), "homeroom": [int(m.group(2)), int(m.group(3))],
                               "general": [int(m.group(4)), int(m.group(5))], "levels": levels}
        if m := re.match(r"(.+?) \((mục tiêu mềm|ưu tiên, càng ít càng tốt)[^)]*\)$", line):
            out["soft"].append(m.group(1))
    return out


def quality(path) -> dict | None:
    """Sheet Chất lượng của file thống kê (writer.quality_rows): tổng số lần không theo luật ưu tiên, tổng điểm trừ, và
    các luật ưu tiên trừ nhiều điểm nhất (tối đa 5) cho trang kết quả. Không có file, sheet: None."""
    from openpyxl import load_workbook

    from ..writer import QUALITY_HEADERS, QUALITY_SHEET
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError):
        return None
    if QUALITY_SHEET not in wb.sheetnames:
        return None
    rows = [list(r) for r in wb[QUALITY_SHEET].iter_rows(values_only=True)]
    wb.close()
    head = next((i for i, r in enumerate(rows) if r[:len(QUALITY_HEADERS)] == list(QUALITY_HEADERS)), None)
    if head is None:
        return None
    body = [r for r in rows[head + 1:] if r and r[0] != "Tổng"]
    soft = [r for r in body if isinstance(r[3], int) and r[3] > 0 and isinstance(r[4], (int, float)) and r[4] > 0]
    soft.sort(key=lambda r: -r[4])
    return {"count": sum(r[3] for r in soft), "points": sum(r[4] for r in soft), "rules": len(soft),
            "top": [{"rule": r[1], "level": r[2], "count": r[3], "points": r[4], "example": r[5]} for r in soft[:5]]}


class Job:
    """Một lần xếp TKB: tiến trình con, đọc từng dòng in ra."""

    def __init__(self, argv: list[str], cwd: Path):
        env = {**os.environ, "PYTHONUNBUFFERED": "1", "PYTHONIOENCODING": "utf-8"}
        if tempfile.tempdir:  # file tạm của tiến trình xếp cũng trong thư mục chương trình (serve)
            env.update(TMP=tempfile.tempdir, TEMP=tempfile.tempdir, TMPDIR=tempfile.tempdir)
        if not getattr(sys, "frozen", False):
            env["PYTHONPATH"] = os.pathsep.join(p for p in (str(PROJECT), env.get("PYTHONPATH")) if p)
        flags = subprocess.CREATE_NEW_PROCESS_GROUP if os.name == "nt" else 0
        self.started = time.time()
        self.ended: float | None = None
        self.lines: list[str] = []
        self.stopping = False
        self._summary: dict | None = None  # kết quả tóm tắt, đọc một lần khi xong
        self.proc = subprocess.Popen(cli_command(argv), cwd=cwd, env=env, stdin=subprocess.DEVNULL,
                                     stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8",
                                     errors="replace", creationflags=flags)
        self._reader = threading.Thread(target=self._read, daemon=True)
        self._reader.start()

    def _read(self) -> None:
        for line in self.proc.stdout:
            self.lines.append(line.rstrip("\r\n"))
        self.proc.wait()
        self.ended = time.time()

    @property
    def running(self) -> bool:
        return self.ended is None

    def stop(self) -> None:
        """Dừng sớm như bấm Ctrl+C: chương trình xếp xong vùng đang xếp rồi ghi TKB tốt nhất."""
        if self.running and self.proc.poll() is None:
            try:
                self.proc.send_signal(signal.CTRL_BREAK_EVENT if os.name == "nt" else signal.SIGINT)
            except OSError as exc:  # vd Windows không cùng cửa sổ dòng lệnh với tiến trình con
                raise InputError(f"Không gửi được lệnh dừng sớm ({exc}); bấm Dừng hẳn nếu cần dừng ngay") from None
            self.stopping = True

    def kill(self) -> None:
        if self.proc.poll() is None:
            self.proc.kill()

    def state(self, since: int = 0) -> dict:
        lines = list(self.lines)
        done = not self.running
        if done and self._summary is None:  # một lần: đọc thêm sheet Chất lượng của file thống kê vừa ghi
            self._summary = summary(lines)
            stats = next((p for p in self._summary["files"] if Path(p).name == FILE_NAMES["stats"]), None)
            self._summary["quality"] = quality(stats) if stats else None
        return {"running": not done, "stopping": self.stopping, "lines": lines[since:], "next": len(lines),
                "elapsed": round((self.ended or time.time()) - self.started),
                "exit": self.proc.returncode if done else None, "summary": self._summary if done else None}


class App:
    """Trạng thái của giao diện: thư mục TKB, mã phiên, việc xếp TKB đang chạy."""

    def __init__(self, home: Path, token: str):
        self.home = Path(home).expanduser().resolve()
        self.folders = {k: self.home / name for k, name in FOLDERS.items()}
        for folder in self.folders.values():  # tạo sẵn để nhà trường thấy chỗ để file
            folder.mkdir(parents=True, exist_ok=True)
        self.out_dir = self.folders["results"]
        self.data = data_dir(self.home).resolve()  # bản nháp của trang (store)
        self.token = token
        self.lock = threading.Lock()  # kịch bản <-> Excel đổi tạm config (rules.applied): mỗi lúc một việc
        self.job: Job | None = None
        self.grid: tuple[str, kich_ban.Grid] | None = None  # bài toán của TKB đang xem (đổi ô thì dùng lại)

    def inside(self, path) -> Path:
        """Đường dẫn trong thư mục TKB (trừ du_lieu/ của chương trình), không thì báo lỗi."""
        p = Path(str(path)).expanduser()
        p = (p if p.is_absolute() else self.home / p).resolve()
        if (p != self.home and self.home not in p.parents) or p == self.data or self.data in p.parents:
            raise InputError(f"Chỉ mở được file trong thư mục TKB ({self.home})")
        return p

    def _keep(self, body: bytes, name: str) -> Path:
        """Chép bản gốc của file lấy từ chỗ khác vào tep_nhap/ (giữ tên; trùng tên khác nội dung thì thêm số)."""
        folder = self.folders["imported"]
        folder.mkdir(parents=True, exist_ok=True)
        stem = safe_name(Path(name).stem)
        path, n = folder / f"{stem}.xlsx", 2
        while path.exists() and path.read_bytes() != body:
            path, n = folder / f"{stem} ({n}).xlsx", n + 1
        if not path.exists():
            path.write_bytes(body)
        return path

    @staticmethod
    def _write(paths: list[Path], write) -> list[str]:
        """Ghi các file bằng `write()`; file đang mở trong Excel (Windows khóa file) thì báo rõ."""
        try:
            return [str(p) for p in write()]
        except PermissionError as exc:
            name = Path(exc.filename or "").name or ", ".join(p.name for p in paths)
            raise InputError(f"Không ghi được {name}: file đang mở trong Excel? Đóng file rồi bấm lại") from None

    # ---- các lệnh /api/ ----
    def schema(self, _):
        return {**kich_ban.schema(), "modes": MODES, "run_defaults": RUN_DEFAULTS, "out_dir": str(self.out_dir),
                "home": str(self.home), "folders": {k: str(p) for k, p in self.folders.items()},
                "files": FILE_NAMES, "running": bool(self.job and self.job.running)}

    def files(self, _):
        """Các file Excel đọc được trong thư mục TKB (hộp Mở của trang), mới nhất trước: theo thư mục ({key, label,
        path, files: [{name, path, kind, saved, time}]}; kind: thực thể, "full" là file vào đủ), và các bộ: file cùng
        tên ở từ hai thư mục thực thể trở lên (Lưu ghi ba file cùng tên)."""
        folders, sets = [], {}
        for key, folder in self.folders.items():
            items = []
            paths = [p for p in folder.glob("*.xlsx") if not p.name.startswith("~$")]  # ~$: file khóa của Excel
            for p in sorted(paths, key=lambda p: (-p.stat().st_mtime, p.name)):
                try:
                    sheets = kich_ban.sheets_in(p)
                except Exception:  # không phải file Excel đọc được: không hiện
                    continue
                kind = kich_ban.entity_of(sheets)
                if kind is None:  # vd TKB.xlsx, Thong_Ke.xlsx: file ra, không mở vào trang được
                    continue
                item = {"name": p.name, "path": str(p), "kind": kind, "saved": config.SAVED_SHEET in sheets,
                        "time": int(p.stat().st_mtime)}
                items.append(item)
                if key in kich_ban.ENTITIES and kind == key:
                    entry = sets.setdefault(p.stem.lower(), {"name": p.stem, "parts": {}, "time": 0})
                    entry["parts"][key] = p.name
                    entry["time"] = max(entry["time"], item["time"])
            folders.append({"key": key, "label": FOLDER_LABELS[key], "path": str(folder), "files": items})
        full_sets = sorted((v for v in sets.values() if len(v["parts"]) > 1), key=lambda v: (-v["time"], v["name"]))
        return {"home": str(self.home), "folders": folders, "sets": full_sets}

    def open_set(self, data: dict):
        """Mở một bộ: các file cùng tên `name` ở ba thư mục thực thể, đọc như một file vào đủ."""
        name = str(data.get("name") or "")
        paths = [p for key in kich_ban.ENTITIES for p in self.folders[key].glob("*.xlsx")
                 if p.stem.lower() == name.lower() and not p.name.startswith("~$")]
        if not paths:
            raise InputError(f"Không có file {name}.xlsx trong các thư mục {', '.join(FOLDERS[k] for k in kich_ban.ENTITIES)}")
        try:
            with self.lock:
                scenario, warnings, sheets = kich_ban.read_files(paths)
        except InputError as exc:
            raise InputError(f"Không đọc được bộ {name}: {exc}") from None
        return {"scenario": scenario, "warnings": warnings, "name": safe_name(name), "sheets": sheets,
                "files": [str(p.relative_to(self.home)) for p in paths]}

    def new(self, query: dict):
        """Kịch bản mới theo bộ luật mẫu `mau` (tkb/bo_mau.py; không ghi: Tiểu học Việt Nam)."""
        with self.lock:
            return {"scenario": kich_ban.default_scenario((query.get("mau") or [None])[0])}

    def sample(self, _):
        """Trường mẫu tên giả (nút "Xem thử với trường mẫu" của trang bắt đầu)."""
        with self.lock:
            return {"scenario": kich_ban.sample_scenario(), "name": "truong_mau"}

    def _import(self, path: Path, name: str):
        try:
            with self.lock:
                scenario, warnings = kich_ban.from_excel(path)
        except InputError as exc:
            raise InputError(f"Không đọc được file {name}: {exc}") from None
        except Exception as exc:  # không phải file Excel (.xlsx) hợp lệ
            from ..__main__ import _not_xlsx, not_xlsx_message
            if _not_xlsx(exc):
                text = not_xlsx_message(name)
                raise InputError(text[0].upper() + text[1:]) from None
            raise InputError(f"Không đọc được file {name} (cần file Excel .xlsx): {exc}") from None
        # File cập nhật "x_cap_nhat.xlsx": xếp lại vẫn ghi x.xlsx, x_cap_nhat.xlsx (không thành x_cap_nhat_cap_nhat).
        stem = re.sub(r"(_cap_nhat)+$", "", Path(name).stem, flags=re.IGNORECASE) or Path(name).stem
        return {"scenario": scenario, "warnings": warnings, "name": safe_name(stem), "sheets": kich_ban.sheets_in(path)}

    def import_file(self, body: bytes, name: str):
        """File lấy từ chỗ khác (hộp chọn file của trình duyệt): đọc được thì chép bản gốc vào tep_nhap/."""
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "vao.xlsx"
            path.write_bytes(body)
            result = self._import(path, name or "vào")
        result["kept"] = str(self._keep(body, name or "vao.xlsx"))
        return result

    def import_path(self, data: dict):
        path = self.inside(data.get("path", ""))
        return self._import(path, path.name)

    def board(self, data: dict):
        """Nhập TKB đang dùng (dạng bảng tin, tkb/nhap_tkb.py) theo kịch bản đang soạn: lưới TKB đã xếp (để thay
        scenario["saved"]), báo cáo và các dòng báo cáo; data: {scenario, run, file (base64), name, mapping}."""
        import base64
        import binascii
        name = str(data.get("name") or "TKB")
        try:
            body = base64.b64decode(data["file"], validate=True)
        except (KeyError, TypeError, binascii.Error):
            raise InputError("Không đọc được file gửi lên") from None
        from ..rules import applied
        with self.lock:
            grid = self._grid(data)
            if grid.problem is None:
                raise InputError("Dữ liệu đang soạn còn lỗi, sửa xong (các bước có dấu ⚠) mới nhập TKB được: "
                                 + "; ".join(grid.errors[:3]))
            with tempfile.TemporaryDirectory() as tmp:
                path = Path(tmp) / "tkb.xlsx"
                path.write_bytes(body)
                try:
                    with applied(grid.values):
                        rows, report = nhap_tkb.read_board(path, grid.problem, data.get("mapping"))
                except InputError as exc:
                    raise InputError(f"Không đọc được TKB trong file {name}: {exc}") from None
                except Exception as exc:  # không phải file Excel (.xlsx) hợp lệ
                    from ..__main__ import _not_xlsx, not_xlsx_message
                    if _not_xlsx(exc):
                        text = not_xlsx_message(name)
                        raise InputError(text[0].upper() + text[1:]) from None
                    raise
            kept = self._keep(body, name)
            problem = grid.problem
            subjects = list(dict.fromkeys(problem.subject_label(c.subject) for c in problem.courses))
            teachers = [{"code": t.code, "name": t.name} for t in sorted(
                (t for t in problem.teachers.values() if not t.supplementary), key=lambda t: t.row or 0)]
        return {"grid": rows, "report": report, "lines": nhap_tkb.describe(report), "subjects": subjects,
                "teachers": teachers, "kept": str(kept)}

    def upgrade(self, data: dict):
        """Bản nháp kịch bản của bản trước (khung giờ cách cũ) -> bản hiện tại (kich_ban.upgrade)."""
        return {"scenario": kich_ban.upgrade(data["scenario"])}

    def describe(self, data: dict):
        with self.lock:
            return kich_ban.describe(data["scenario"], data.get("rows"), bool(data.get("student_rules", True)))

    def check(self, data: dict):
        run = {**RUN_DEFAULTS, **(data.get("run") or {})}
        run_argv(Path("x"), self.out_dir, "x", run)  # cài đặt chạy sai thì báo lỗi luôn
        with self.lock:
            return kich_ban.check(data["scenario"], run["mode"], int(run["overtime_max"]), bool(run["student_rules"]),
                                  quick=bool(data.get("quick")))

    def _grid(self, data: dict) -> kich_ban.Grid:
        """Bài toán của kịch bản để xem, đổi ô TKB (gọi trong self.lock): dựng lại khi kịch bản (trừ TKB) hay cài
        đặt chạy đổi."""
        run = {**RUN_DEFAULTS, **(data.get("run") or {})}
        run_argv(Path("x"), self.out_dir, "x", run)
        args = (run["mode"], int(run["overtime_max"]), bool(run["student_rules"]))
        key = kich_ban.grid_key(data["scenario"], *args)
        if self.grid is None or self.grid[0] != key:
            self.grid = key, kich_ban.Grid(data["scenario"], *args)
        return self.grid[1]

    def timetable(self, data: dict):
        """TKB đã xếp của kịch bản để vẽ trên trang, kèm lỗi luật bắt buộc và các ô có lỗi (kich_ban.Grid.view)."""
        with self.lock:
            return self._grid(data).view(data["scenario"].get("saved"))

    def swaps(self, data: dict):
        """Các ô cùng lớp đổi chỗ được với một ô mà không thêm lỗi luật bắt buộc (kich_ban.Grid.swaps)."""
        with self.lock:
            return {"swaps": self._grid(data).swaps(data["scenario"].get("saved"), str(data.get("cls")),
                                                    int(data.get("d", -1)), int(data.get("p", -1)))}

    def save(self, data: dict):
        """Tệp › Lưu: ba file cùng tên <tên>.xlsx, mỗi thực thể một file ở thư mục của nó, và file vào đủ (có cả TKB
        đã xếp) ở ket_qua/ như file vào của lần xếp."""
        name = safe_name(data.get("name"))
        paths = {**{k: self.folders[k] / f"{name}.xlsx" for k in kich_ban.ENTITIES}, "full": self.out_dir / f"{name}.xlsx"}
        with self.lock:
            files = self._write(list(paths.values()), lambda: kich_ban.to_entity_files(data["scenario"], paths))
        return {"files": files}

    def rules_file(self, data: dict):
        """File chỉ có các luật (sheet LUẬT, HƯỚNG DẪN) ở luat/: các luật của kịch bản (<tên>.xlsx), hoặc (mẫu) các
        luật có sẵn (Mau_Luat.xlsx)."""
        name = RULES_TEMPLATE_NAME if data.get("template") else f"{safe_name(data.get('name'))}.xlsx"
        path = self.folders["rules"] / name
        with self.lock:
            scenario = kich_ban.default_scenario() if data.get("template") else data["scenario"]
            files = self._write([path], lambda: [kich_ban.rules_to_excel(scenario, path) or path])
        return {"files": files}

    def store(self, data=None):
        """Bản nháp và lựa chọn của trang ({khóa: chữ}), ghi ở data/trang.json thay cho bộ nhớ trình duyệt: xóa thư
        mục chương trình là hết. GET (data là query): đọc; POST {items}: ghi đè cả bản."""
        path = self.data / STORE_NAME
        if not isinstance(data, dict) or "items" not in data:
            try:
                return json.loads(path.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                return {}
        items = {str(k): str(v) for k, v in dict(data["items"]).items()}
        self.data.mkdir(parents=True, exist_ok=True)
        tmp = path.with_suffix(".tmp")
        tmp.write_text(json.dumps(items, ensure_ascii=False), encoding="utf-8")
        os.replace(tmp, path)
        return {"ok": True}

    def template(self, _=None):
        """Mẫu trống của ba thực thể (file mẫu python -m tkb.template chia theo thực thể), mỗi file ở thư mục của nó:
        nhà trường điền trong Excel rồi mở lại ở đây."""
        from ..template import write_staff_template

        paths = {k: self.folders[k] / name for k, name in TEMPLATE_NAMES.items()}
        with tempfile.TemporaryDirectory() as tmp:
            full = Path(tmp) / "mau.xlsx"
            with self.lock:
                write_staff_template(full)
                files = self._write(list(paths.values()), lambda: kich_ban.split_file(full, paths))
        return {"files": files}

    def run(self, data: dict):
        if self.job and self.job.running:
            raise InputError("Đang xếp TKB, chờ xong hoặc bấm Dừng trước khi xếp lại")
        run = {**RUN_DEFAULTS, **(data.get("run") or {})}
        name = safe_name(run.get("name"))
        out = self.out_dir
        source = out / f"{name}.xlsx"
        argv = run_argv(source, out, name, run)
        out.mkdir(parents=True, exist_ok=True)
        with self.lock:
            self._write([source], lambda: [kich_ban.to_excel(data["scenario"], source) or source])
        self.job = Job(argv, out)
        return {"input": str(source), "command": " ".join(cli_command(argv))}

    def status(self, query: dict):
        if self.job is None:
            return {"running": False, "lines": [], "next": 0, "exit": None, "summary": None}
        return self.job.state(int((query.get("since") or ["0"])[0] or 0))

    def stop(self, _):
        if self.job:
            self.job.stop()
        return {"ok": True}

    def kill(self, _):
        if self.job:
            self.job.kill()
        return {"ok": True}

    def open(self, data: dict):
        path = self.inside(data.get("path") or self.out_dir)
        if not path.exists():
            raise InputError(f"Không có {path}")
        try:
            if os.name == "nt":
                os.startfile(path)  # noqa: S606 (mở bằng ứng dụng mặc định, vd Excel)
            else:
                subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", str(path)],
                                 stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        except OSError as exc:
            raise InputError(f"Không mở được {path.name}: {exc}") from None
        return {"ok": True}

    def forget(self, _):
        """Tệp › Xóa dữ liệu trên máy này: bỏ TKB đang xem và nhật ký lần chạy (có tên giáo viên) khỏi bộ nhớ.

        File đã ghi ra thư mục kết quả để nguyên: trang chỉ chỗ cho người dùng tự xóa."""
        if self.job and self.job.running:
            raise InputError("Đang xếp TKB: bấm Dừng, chờ dừng hẳn rồi xóa")
        with self.lock:
            self.job = None
            self.grid = None
        (self.data / STORE_NAME).unlink(missing_ok=True)  # bản nháp có tên giáo viên
        return {"ok": True}

    def shutdown(self) -> None:
        if self.job:
            self.job.kill()


class Handler(BaseHTTPRequestHandler):
    app: App  # gán khi tạo máy chủ (make_server)
    server_version = "TKB"

    def log_message(self, fmt, *args):  # không in mỗi yêu cầu ra màn hình
        pass

    # ---- trả lời ----
    def _send(self, status: int, body: bytes, ctype: str, headers: dict | None = None) -> None:
        self.send_response(status)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        for k, v in (headers or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def _json(self, status: int, payload) -> None:
        self._send(status, json.dumps(payload, ensure_ascii=False).encode("utf-8"), "application/json; charset=utf-8")

    def _host_ok(self) -> bool:
        host = (self.headers.get("Host") or "").rsplit(":", 1)[0].strip("[]").lower()
        return host in ("127.0.0.1", "localhost")

    def _token_ok(self, query: dict) -> bool:
        given = self.headers.get("X-TKB-Token") or (query.get("t") or [""])[0]
        return secrets.compare_digest(given.encode(), self.app.token.encode())

    def _body(self) -> bytes:
        length = int(self.headers.get("Content-Length") or 0)
        if length > MAX_BODY:
            raise InputError("Dữ liệu quá lớn")
        return self.rfile.read(length) if length else b""

    # ---- yêu cầu ----
    def do_GET(self):
        url = urlparse(self.path)
        if not self._host_ok():
            return self._send(HTTPStatus.FORBIDDEN, b"Forbidden", "text/plain")
        if url.path.startswith("/api/"):
            return self._api("GET", url)
        self._static(url.path)

    def do_POST(self):
        url = urlparse(self.path)
        if not self._host_ok() or not url.path.startswith("/api/"):
            return self._send(HTTPStatus.FORBIDDEN, b"Forbidden", "text/plain")
        self._api("POST", url)

    def _static(self, path: str) -> None:
        name = "index.html" if path in ("", "/") else path.lstrip("/")
        file = (STATIC / name).resolve()
        if STATIC not in file.parents or not file.is_file() or file.suffix not in TYPES:
            return self._send(HTTPStatus.NOT_FOUND, "Không có trang này".encode(), "text/plain; charset=utf-8")
        body = file.read_bytes()
        if file.name == "index.html":
            body = body.replace(TOKEN_MARK.encode(), self.app.token.encode())
        self._send(HTTPStatus.OK, body, TYPES[file.suffix])

    def _api(self, method: str, url) -> None:
        query = parse_qs(url.query)
        if not self._token_ok(query):
            return self._json(HTTPStatus.FORBIDDEN, {"error": "Phiên đã hết hạn: mở lại trang bằng đường dẫn in ở "
                                                              "cửa sổ chương trình"})
        app = self.app
        name = url.path[len("/api/"):]
        try:
            if method == "POST" and name == "import":
                return self._json(HTTPStatus.OK, app.import_file(self._body(),
                                                                 unquote(self.headers.get("X-File-Name") or "")))
            routes = {("GET", "schema"): app.schema, ("GET", "new"): app.new, ("GET", "sample"): app.sample,
                      ("GET", "status"): app.status, ("GET", "ping"): lambda _: {"ok": True},  # trang hỏi còn chạy
                      ("POST", "import_path"): app.import_path, ("POST", "check"): app.check, ("POST", "run"): app.run,
                      ("GET", "files"): app.files, ("POST", "open_set"): app.open_set, ("POST", "save"): app.save,
                      ("POST", "template"): app.template, ("POST", "rules_file"): app.rules_file,
                      ("POST", "describe"): app.describe, ("POST", "timetable"): app.timetable,
                      ("POST", "upgrade"): app.upgrade, ("POST", "board"): app.board,
                      ("POST", "swaps"): app.swaps,
                      ("POST", "stop"): app.stop, ("POST", "kill"): app.kill, ("POST", "open"): app.open,
                      ("POST", "forget"): app.forget,
                      ("GET", "store"): app.store, ("POST", "store"): app.store}
            handler = routes.get((method, name))
            if handler is None:
                return self._json(HTTPStatus.NOT_FOUND, {"error": f"Không có lệnh {name}"})
            arg = query if method == "GET" else json.loads(self._body() or b"{}")
            return self._json(HTTPStatus.OK, handler(arg))
        except InputError as exc:
            return self._json(HTTPStatus.BAD_REQUEST, {"error": str(exc)})
        except (ValueError, KeyError, TypeError) as exc:
            return self._json(HTTPStatus.BAD_REQUEST, {"error": f"Dữ liệu gửi lên không hợp lệ: {exc}"})


def make_server(port: int = PORT, home=None) -> tuple[ThreadingHTTPServer, App, str]:
    """Máy chủ trên 127.0.0.1 (cổng bận thì lấy cổng khác): (máy chủ, trạng thái, đường dẫn trang). `home`: thư mục
    TKB (không ghi: thư mục của chương trình)."""
    app = App(Path(home) if home else app_dir(), secrets.token_urlsafe(16))
    handler = type("TkbHandler", (Handler,), {"app": app})
    try:
        httpd = ThreadingHTTPServer(("127.0.0.1", port), handler)
    except OSError:
        httpd = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    return httpd, app, f"http://127.0.0.1:{httpd.server_port}/"


# Trình duyệt mở trang ở cửa sổ riêng với hồ sơ trong thư mục chương trình (lịch sử, bộ nhớ đệm, file tải về đều ở
# đó). Edge có sẵn trên mọi máy Windows 10/11; không có thì Chrome, Chromium; không có nữa thì trình duyệt mặc định.
BROWSERS = [r"Microsoft\Edge\Application\msedge.exe", r"Google\Chrome\Application\chrome.exe"]
BROWSER_NAMES = ["msedge", "microsoft-edge", "google-chrome", "chrome", "chromium", "chromium-browser"]


def find_browser() -> str | None:
    if os.name == "nt":
        for root in (os.environ.get(k) for k in ("PROGRAMFILES(X86)", "PROGRAMFILES", "LOCALAPPDATA")):
            for rel in BROWSERS:
                if root and (Path(root) / rel).is_file():
                    return str(Path(root) / rel)
    return next((found for name in BROWSER_NAMES if (found := shutil.which(name))), None)


def open_page(url: str, app: App) -> None:
    """Mở trang: trình duyệt họ Chromium với --user-data-dir trong data/trinh_duyet (file tải về vào thư mục kết
    quả); không có thì trình duyệt mặc định (khi đó trình duyệt nhớ lịch sử trang như mọi trang web)."""
    exe = find_browser()
    if not exe:
        webbrowser.open(url)
        return
    profile = app.data / "trinh_duyet"
    prefs = profile / "Default" / "Preferences"
    try:
        if not prefs.is_file():
            prefs.parent.mkdir(parents=True, exist_ok=True)
            prefs.write_text(json.dumps({"download": {"default_directory": str(app.out_dir), "prompt_for_download": False},
                                         "savefile": {"default_directory": str(app.out_dir)}}), encoding="utf-8")
        subprocess.Popen([exe, f"--user-data-dir={profile}", "--no-first-run", "--no-default-browser-check",
                          f"--app={url}"], stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except OSError:
        webbrowser.open(url)


def serve(port: int = PORT, open_browser: bool = True, home=None) -> int:
    httpd, app, url = make_server(port, home)
    tmp = app.data / "tam"  # file tạm (đọc, kiểm tra file Excel) cũng trong thư mục chương trình; tắt thì xóa
    shutil.rmtree(tmp, ignore_errors=True)
    tmp.mkdir(parents=True, exist_ok=True)
    tempfile.tempdir = str(tmp)
    print(f"Giao diện xếp TKB: {url}")
    print(f"Thư mục TKB     : {app.home}")
    print("Giữ cửa sổ này mở trong khi dùng giao diện; bấm Ctrl+C để tắt.")
    if open_browser:
        open_page(url, app)
    try:
        httpd.serve_forever()
    except KeyboardInterrupt:
        print("Đã tắt giao diện.")
    finally:
        app.shutdown()
        httpd.server_close()
        shutil.rmtree(tmp, ignore_errors=True)
    return 0
