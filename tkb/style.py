"""Style của các file ra, chép từ sheet NHÂN SỰ của file vào (font, viền, căn lề, nền, chiều cao dòng).

Không viết cứng phông chữ hay cỡ chữ trong code: đổi style của file vào thì các file ra đổi theo.
"""
from __future__ import annotations

import math
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill
from openpyxl.utils import get_column_letter

from .staff import _find_columns, clean_name, staff_sheet

DEFAULT_ROW_HEIGHT = 25  # file vào không đặt chiều cao dòng
MAX_COLUMN_WIDTH = 60  # chữ dài hơn thì xuống dòng
STAFF_HEADERS = {"name": "Họ và Tên", "title": "Chức Vụ", "class": "Lớp", "lessons": "Số Tiết/Tuần",
                 "maternity": "Chế Độ"}


def with_bold(font: Font, bold: bool) -> Font:
    return Font(name=font.name, sz=font.sz, b=bold, i=font.i, u=font.u, strike=font.strike, color=font.color,
                vertAlign=font.vertAlign, charset=font.charset, family=font.family, scheme=font.scheme)


@dataclass
class CellStyle:
    font: Font = field(default_factory=Font)
    border: Border = field(default_factory=Border)
    alignment: Alignment = field(default_factory=lambda: Alignment(horizontal="center", vertical="center"))
    fill: PatternFill | None = None

    @classmethod
    def of(cls, cell) -> CellStyle:
        return cls(copy(cell.font), copy(cell.border), copy(cell.alignment),
                   copy(cell.fill) if cell.fill is not None and cell.fill.fill_type else None)

    def apply(self, cell, bold: bool | None = None, horizontal: str | None = None, wrap: bool = True) -> None:
        cell.font = copy(self.font) if bold is None else with_bold(self.font, bold)
        cell.border = copy(self.border)
        cell.alignment = Alignment(horizontal=horizontal or self.alignment.horizontal,
                                   vertical=self.alignment.vertical or "center", wrap_text=wrap)
        if self.fill is not None:
            cell.fill = copy(self.fill)


@dataclass
class Style:
    header: CellStyle = field(default_factory=lambda: CellStyle(font=Font(b=True)))
    body: CellStyle = field(default_factory=CellStyle)
    row_height: float = DEFAULT_ROW_HEIGHT
    staff_headers: dict[str, str] = field(default_factory=lambda: dict(STAFF_HEADERS))

    @classmethod
    def from_file(cls, path: str | Path) -> Style:
        """Style của ô tiêu đề và ô dữ liệu đầu tiên ở cột Chức Vụ của sheet nhân sự."""
        ws = staff_sheet(openpyxl.load_workbook(path))
        header_row, cols = _find_columns(ws)
        col = cols["title"]
        height = ws.row_dimensions[header_row + 1].height or ws.row_dimensions[header_row].height
        headers = dict(STAFF_HEADERS)
        headers.update({k: clean_name(ws.cell(header_row, c).value) for k, c in cols.items() if k in headers})
        return cls(header=CellStyle.of(ws.cell(header_row, col)), body=CellStyle.of(ws.cell(header_row + 1, col)),
                   row_height=round(height) if height else DEFAULT_ROW_HEIGHT, staff_headers=headers)

    @property
    def font_size(self) -> float:
        return float(self.body.font.sz or 11)

    @property
    def line_height(self) -> float:
        """Chiều cao một dòng chữ (point)."""
        return self.font_size * 1.5

    def text_width(self, text) -> float:
        """Độ rộng ước lượng (đơn vị cột Excel) của một dòng chữ."""
        return len(str(text)) * self.font_size / 11.5 + 2

    def lines(self, text, width: float) -> int:
        return sum(max(1, math.ceil(self.text_width(part) / width)) for part in str(text).split("\n"))

    def header_cell(self, ws, row: int, col: int, value) -> None:
        self.header.apply(ws.cell(row, col, value))

    def body_cell(self, ws, row: int, col: int, value, bold: bool | None = None, horizontal: str | None = None):
        cell = ws.cell(row, col, value)
        self.body.apply(cell, bold=bold, horizontal=horizontal)
        return cell

    def title_cell(self, ws, row: int, col: int, value) -> None:
        """Dòng tựa phía trên một bảng: chữ tiêu đề, không viền."""
        cell = ws.cell(row, col, value)
        cell.font = copy(self.header.font)
        cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = self.row_height

    def table(self, ws, header: list[str], rows: list[list], top: int = 1, bold_last: bool = False) -> int:
        """Bảng như sheet nhân sự của file vào: dòng tiêu đề rồi các dòng dữ liệu. Trả về dòng cuối."""
        for c, h in enumerate(header, start=1):
            self.header_cell(ws, top, c, h)
        ws.row_dimensions[top].height = self.row_height
        r = top
        for i, values in enumerate(rows):
            r = top + 1 + i
            for c, v in enumerate(values, start=1):
                self.body_cell(ws, r, c, v, bold=True if bold_last and i == len(rows) - 1 else None)
            ws.row_dimensions[r].height = self.row_height
        return r

    def fit_columns(self, ws, first_row: int = 1, minimum: float = 8, skip_rows=()) -> None:
        """Nới độ rộng cột vừa chữ dài nhất (tối đa MAX_COLUMN_WIDTH, dài hơn thì xuống dòng).

        skip_rows: các dòng tựa (chữ tràn sang ô bên cạnh) không tính.
        """
        widths: dict[int, float] = {}
        for row in ws.iter_rows(min_row=first_row):
            if row[0].row in skip_rows:
                continue
            for cell in row:
                if cell.value is None or type(cell).__name__ == "MergedCell":
                    continue
                w = max(self.text_width(part) for part in str(cell.value).split("\n"))
                widths[cell.column] = max(widths.get(cell.column, minimum), w)
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = min(MAX_COLUMN_WIDTH, round(w, 1))
        # Chữ bị xuống dòng thì dòng cao thêm.
        for row in ws.iter_rows(min_row=first_row):
            if row[0].row in skip_rows:
                continue
            n = max((self.lines(c.value, min(MAX_COLUMN_WIDTH, widths.get(c.column, minimum)))
                     for c in row if c.value is not None and type(c).__name__ != "MergedCell"), default=1)
            if n > 1:
                ws.row_dimensions[row[0].row].height = max(self.row_height, n * self.line_height)
