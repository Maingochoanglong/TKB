"""Nhập TKB nhà trường đang dùng, dạng bảng tin (như file TKB toàn trường: mỗi lớp một cột, mỗi tiết ghi môn và,
khi người dạy không phải GVCN, tên giáo viên), thành lưới TKB đã xếp (sheet config.SAVED_SHEET, đọc bằng
staff.parse_saved_grid). Sau khi nhập, giao diện vẽ TKB ở bước 8, kiểm mọi luật bắt buộc; "Xếp lại phần còn lại"
(xếp lại ít xáo trộn) sửa ít ô nhất để TKB đúng luật.

Đọc dễ tính, không cần đúng một mẫu:
- Dòng tên lớp: dòng có nhiều tên lớp của file vào nhất (trong 15 dòng đầu của mọi sheet); mỗi lớp một cột.
- Cột tiết: cột bên trái các lớp có nhiều số nhất (số tiết trong buổi hay trong ngày). Mỗi số là một tiết; tiết
  chiếm các dòng tới số kế tiếp (ô gộp hay không). Số 1 (hay số không lớn hơn số trước) là đầu một buổi.
- Buổi theo thứ tự khung giờ (Thứ 2 sáng, Thứ 2 chiều …); có chữ ngày ở các cột bên trái (vd "2", "Thứ 3", "T4")
  thì theo ngày đó. Tiết ngoài khung giờ (vd chiều Thứ 5 có tiết 4 mà khung giờ chỉ 3 tiết) bỏ qua, có báo.
- Ô: các dòng chữ của tiết đó trong cột lớp. Môn khớp tên môn, Tên trong TKB, chữ viết tắt (writer.short_subjects),
  chữ đầu các từ viết hoa ("LS&ĐL" là Lịch Sử và Địa Lý), hay vài chữ đầu của tên ("Tin"); không phân biệt hoa
  thường, dấu, dấu câu. Giáo viên khớp Tên Viết Tắt, tên tự viết tắt (writer.short_names), họ tên, Mã GV, tên ở
  dòng GVCN, hay tên kèm chữ đầu của họ/tên đệm ("N.Ly", "Võ Ly"). Ô không ghi tên giáo viên: GVCN của lớp dạy.
- Chữ không khớp (vd "T" cho Toán, "GDTC" cho Thể Dục) trả về để người dùng ghép tay (`mapping`), không đoán bừa.

Dòng lệnh: python -m tkb.nhap_tkb <TKB đang dùng.xlsx> <file vào.xlsx> [-o <file ra.xlsx>] [--ghep CHỮ=TÊN ...]
  [--max-overtime n]
"""
from __future__ import annotations

import argparse
import difflib
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

from . import config, khung_gio
from .staff import InputError, _fold, class_sort_key, clean_name, fold_key

MAX_HEADER_ROW = 15  # dòng tên lớp nằm trong mấy dòng đầu sheet
EXAMPLES = 3  # số chỗ ghi làm ví dụ cho mỗi chữ không khớp


def _class_key(text) -> str:
    """Khóa so tên lớp: không dấu, không khoảng trắng, bỏ chữ "Lớp" đầu, "." "-" như "/" ("Lớp 3.2" là "3/2")."""
    key = re.sub(r"\s+", "", _fold(text))
    key = re.sub(r"^lop", "", key)
    return re.sub(r"[.\-]", "/", key)


def _initials(text: str) -> str:
    """Chữ đầu các từ viết hoa, từ viết hoa cả giữ nguyên: "Lịch Sử và Địa Lý" -> "LSĐL", "TV Tăng Cường" -> "TVTC"."""
    words = re.findall(r"[^\W\d_]+", text)
    return "".join(w if len(w) > 1 and w.isupper() else w[0] for w in words if w[0].isupper())


class _Names:
    """So chữ trong ô với môn và giáo viên của bài toán (khóa staff.fold_key)."""

    def __init__(self, problem, homeroom_row: dict[str, str], mapping: dict | None):
        from .writer import short_names, short_subjects
        self.problem = problem
        subjects = list(dict.fromkeys(c.subject for c in problem.courses))
        self.subjects = subjects
        short = short_subjects(problem)
        self.subject_of: dict[str, str] = {}
        # Tên, Tên trong TKB, chữ viết tắt trước; chữ đầu các từ viết hoa sau ("TV tăng cường" ra "TV" trùng Tiếng
        # Việt: không dùng). Khóa trùng hai môn ở cùng một bậc thì bỏ.
        for texts in (lambda s, label: (s, label, short[s]), lambda s, label: (_initials(label), _initials(s))):
            found: dict[str, set] = defaultdict(set)
            for s in subjects:
                for text in texts(s, problem.subject_label(s)):
                    if fold_key(text) and fold_key(text) not in self.subject_of:
                        found[fold_key(text)].add(s)
            self.subject_of.update({k: next(iter(v)) for k, v in found.items() if len(v) == 1})
        self.full = {s: {fold_key(s), fold_key(problem.subject_label(s))} for s in subjects}
        self.teachers = {title: t for title, t in problem.teachers.items() if not t.supplementary}
        self.homeroom = {t.class_name: title for title, t in self.teachers.items() if t.class_name}
        names = short_names(problem.teachers)
        found = defaultdict(set)
        for title, t in self.teachers.items():
            for text in (t.short_name, names[title], t.name, t.code):
                if fold_key(text):
                    found[fold_key(text)].add(title)
        for cls, text in homeroom_row.items():  # dòng GVCN dưới tên lớp (bỏ chữ là tên môn: không có dòng GVCN)
            if cls in self.homeroom and fold_key(text) and fold_key(text) not in self.subject_of:
                found[fold_key(text)].add(self.homeroom[cls])
        self.teacher_of = {k: next(iter(v)) for k, v in found.items() if len(v) == 1}
        self.code_of = {fold_key(t.code): title for title, t in self.teachers.items()}
        # Ghép tay: chữ -> môn hoặc Mã GV ("" = bỏ qua: môn thì bỏ cả ô, tên giáo viên thì coi như GVCN dạy).
        self.mapped_subject: dict[str, str] = {}
        self.mapped_teacher: dict[str, str] = {}
        for token, value in ((mapping or {}).get("subjects") or {}).items():
            target = self.subject(value) if value else ""
            if target is not None:
                self.mapped_subject[fold_key(token)] = target
        for token, value in ((mapping or {}).get("teachers") or {}).items():
            target = (self.code_of.get(fold_key(value)) or self.teacher(value)) if value else ""
            if target is not None:
                self.mapped_teacher[fold_key(token)] = target

    def subject(self, text) -> str | None:
        """Môn của chữ `text`; "" là chữ ghép tay để bỏ qua; None nếu không khớp."""
        key = fold_key(text)
        if not key:
            return None
        if key in self.mapped_subject:
            return self.mapped_subject[key]
        if key in self.subject_of:
            return self.subject_of[key]
        if len(key) >= 3 and key not in self.teacher_of and not self.candidates(text):
            # vài chữ đầu của tên môn, khớp đúng một môn ("Tin" là Tin Học); không áp cho tên giáo viên
            hits = [s for s, keys in self.full.items() if any(k.startswith(key) for k in keys)]
            if len(hits) == 1:
                return hits[0]
        return None

    def candidates(self, text) -> list[str]:
        """Các GV có thể là tên `text` ("Khoa", "N.Ly", "Võ Ly", "H. Vi"): tên trùng từ cuối họ tên, chữ đầu ghi trước
        tên trùng chữ đầu của một từ đứng trước, các từ ghi trước tên có trong họ tên."""
        s = re.sub(r"\s*\.\s*", ".", clean_name(text))
        if "." in s:
            *initials, given = s.split(".")
            initials, before = [i for i in initials if i], []
        else:
            words = s.split()
            given, before, initials = words[-1], words[:-1], []
        out = []
        for title, t in self.teachers.items():
            words = [_fold(w) for w in t.name.split()]
            if not words or words[-1] != _fold(given):
                continue
            rest = words[:-1]
            if any(_fold(i)[:1] not in {w[:1] for w in rest} for i in initials):
                continue
            if any(_fold(b) not in rest for b in before):
                continue
            out.append(title)
        return out

    def teacher(self, text) -> str | None:
        """Mã GV (Teacher.title) của tên `text`; "" là chữ ghép tay để coi như GVCN dạy; None nếu không khớp được
        đúng một người."""
        key = fold_key(text)
        if not key:
            return None
        if key in self.mapped_teacher:
            return self.mapped_teacher[key]
        if key in self.teacher_of:
            return self.teacher_of[key]
        found = self.candidates(text)
        if len(found) > 1:  # dòng tên dưới môn thường là GV khác GVCN của lớp
            found = [t for t in found if self.teachers[t].role != config.ROLE_HOMEROOM] or found
        return found[0] if len(found) == 1 else None


def _find_header(wb, classes: list[str]):
    """(sheet, dòng tên lớp, {cột: lớp}) có nhiều tên lớp của file vào nhất."""
    keys = {_class_key(c): c for c in classes}
    best = (None, 0, {})
    for ws in wb.worksheets:
        for r in range(1, min(ws.max_row, MAX_HEADER_ROW) + 1):
            cols = {}
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is not None and (cls := keys.get(_class_key(v))) and cls not in cols.values():
                    cols[c] = cls
            if len(cols) > len(best[2]):
                best = (ws, r, cols)
    if not best[2]:
        raise InputError("không tìm thấy dòng tên lớp: TKB cần một dòng ghi tên các lớp (mỗi lớp một cột) đúng như "
                         "danh sách lớp của file vào")
    return best


def _number(value) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and float(value).is_integer() and 0 < value <= 20:
        return int(value)
    if isinstance(value, str) and re.fullmatch(r"\s*\d{1,2}\s*", value) and 0 < int(value) <= 20:
        return int(value)
    return None


def read_board(path, problem, mapping: dict | None = None) -> tuple[list[list], dict]:
    """TKB dạng bảng tin ở `path` -> (lưới TKB đã xếp như sheet config.SAVED_SHEET, báo cáo). Gọi trong
    rules.applied của file vào (khung giờ, tên môn). mapping: {"subjects": {chữ: tên môn | ""}, "teachers": {chữ: Mã
    GV hay tên | ""}} ghép tay các chữ không tự khớp được.

    Báo cáo: {classes, missing, cells, outside, double, no_teacher, sessions, unknown_subjects, unknown_teachers}:
    unknown_* là [{token, count, where, guess/candidates}] để giao diện hỏi người dùng."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, header, class_cols = _find_header(wb, problem.classes)
    first_col = min(class_cols)
    # Cột tiết: cột bên trái các lớp có nhiều số nhất (bằng nhau thì cột gần các lớp hơn).
    counts = {c: sum(_number(ws.cell(r, c).value) is not None for r in range(header + 1, ws.max_row + 1))
              for c in range(1, first_col)}
    period_col = max(counts, key=lambda c: (counts[c], c), default=None)
    if period_col is None or counts[period_col] == 0:
        raise InputError("không tìm thấy cột tiết: cột bên trái các lớp cần ghi số tiết (1, 2, 3 …)")
    span = {(rng.min_row, rng.min_col): rng.max_row for rng in ws.merged_cells.ranges}
    starts = [r for r in range(header + 1, ws.max_row + 1) if _number(ws.cell(r, period_col).value) is not None]
    blocks = []  # (dòng đầu, dòng cuối, số tiết ghi ở cột tiết)
    for i, r in enumerate(starts):
        if i + 1 < len(starts):
            end = starts[i + 1] - 1
        else:  # tiết cuối: hết ô gộp, hoặc thêm dòng dưới nếu dòng đó có chữ ở cột lớp (dòng tên GV)
            end = span.get((r, period_col), r)
            if end == r and r < ws.max_row and any(ws.cell(r + 1, c).value not in (None, "") for c in class_cols):
                end = r + 1
        blocks.append((r, end, _number(ws.cell(r, period_col).value)))
    # Các buổi: số 1, hay số không lớn hơn số trước, mở buổi mới.
    sessions: list[list] = []
    for block in blocks:
        if not sessions or block[2] <= sessions[-1][-1][2]:
            sessions.append([])
        sessions[-1].append(block)
    frame = [(d, s) for d in khung_gio.days() for s in khung_gio.sessions(d)]

    def day_hint(rows: list[int]) -> int | None:
        for r in rows:
            for c in range(1, first_col):
                v = ws.cell(r, c).value
                if c != period_col and v is not None and (d := khung_gio.day_index(v)) is not None:
                    return d
        return None

    placed = []  # (buổi trong khung giờ hay None, các tiết của buổi)
    at = 0
    for group in sessions:
        hint = day_hint([r for r1, r2, _ in group for r in range(r1, r2 + 1)])
        if hint is not None and at < len(frame) and frame[at][0] != hint:
            at = next((i for i, (d, _) in enumerate(frame) if d == hint), at)
        placed.append((frame[at] if at < len(frame) else None, group))
        at += 1

    names = _Names(problem, {cls: str(ws.cell(header + 1, c).value or "") for c, cls in class_cols.items()},
                   mapping)
    links = {cls: [set(g) for g in groups] for cls, groups in problem.links.items()}
    lessons: dict[tuple[str, int, int], list[tuple[str, str]]] = {}
    unknown_s: dict[str, list] = {}
    unknown_t: dict[str, list] = {}
    outside, report = [], Counter()

    def note(store: dict, token: str, where: str) -> None:
        item = store.setdefault(fold_key(token), [0, [], token])  # cùng khóa ("Đ.đức", "Đ\\.đức"): một chữ
        item[0] += 1
        if len(item[1]) < EXAMPLES:
            item[1].append(where)

    for spot, group in placed:
        for r1, r2, k in group:
            d, session = spot if spot else (None, None)
            period = session.periods[k - 1] if session and k <= len(session.periods) else None
            for c, cls in class_cols.items():
                pieces = [p.strip() for r in range(r1, r2 + 1) if (v := ws.cell(r, c).value) is not None
                          for p in str(v).splitlines() if p.strip()]
                if not pieces:
                    continue
                if period is None:
                    if len(outside) < EXAMPLES:
                        outside.append(f"lớp {cls}, " + (f"{config.DAYS[d]} buổi {session.name} tiết {k}" if session
                                                         else f"dòng {r1}"))
                    report["outside"] += 1
                    continue
                where = f"lớp {cls}, {config.DAYS[d]} tiết {period}"
                subjects, teachers, unknown, skip = [], [], [], False
                # Các môn học cùng giờ ghi "A / B" (TKB toàn trường): mỗi phần một mẩu.
                pieces = [x.strip() for p in pieces for x in p.split(" / ") if x.strip()]
                for piece in pieces:
                    parts = [piece]
                    if not subjects and names.subject(piece) is None and names.teacher(piece) is None:
                        words = piece.split()  # "TA Phi", "Tiếng Anh Phi": môn rồi tên trong một dòng
                        for i in range(len(words) - 1, 0, -1):
                            if names.subject(" ".join(words[:i])) is not None:
                                parts = [" ".join(words[:i]), " ".join(words[i:])]
                                break
                    for part in parts:
                        # Đã có môn thì dòng sau là tên giáo viên trước: "An" là cô An, không phải ÂN (Âm Nhạc).
                        if (subjects or skip) and (t := names.teacher(part)) is not None:
                            teachers.append(t)
                            continue
                        s = names.subject(part)
                        if s == "":
                            skip = True
                        elif s is not None:
                            subjects.append(s)
                        elif (t := names.teacher(part)) is not None:
                            teachers.append(t)  # "" = ghép tay: coi như GVCN dạy
                        else:
                            unknown.append(part)
                if not subjects and not skip and unknown:
                    note(unknown_s, unknown.pop(0), where)
                    continue
                for token in unknown:
                    note(unknown_t, token, where)
                if skip or not subjects:
                    report["skipped"] += 1
                    continue
                if unknown:
                    report["unknown_teacher_cells"] += 1
                    continue
                subjects = list(dict.fromkeys(subjects))
                together = next((g for g in links.get(cls, []) if set(subjects) <= g), None)
                if len(subjects) > 1 and together is None:
                    report["double"] += 1
                    subjects = subjects[:1]
                who = [t for t in teachers if t] or [names.homeroom.get(cls)]
                if who[0] is None:
                    report["no_teacher"] += 1
                    continue
                lessons[cls, d, period] = [(s, who[min(i, len(who) - 1)]) for i, s in enumerate(subjects)]
    grid = _grid(problem, lessons)
    guesses = {**{k: problem.subject_label(s) for k, s in names.subject_of.items()},
               **{fold_key(problem.subject_label(s)): problem.subject_label(s) for s in names.subjects}}
    return grid, {
        "classes": [class_cols[c] for c in sorted(class_cols)],
        "missing": sorted(set(problem.classes) - set(class_cols.values()), key=class_sort_key),
        "cells": len(lessons),
        "sessions": [len(placed), len(frame)],
        "outside": report["outside"], "outside_examples": outside,
        "double": report["double"], "no_teacher": report["no_teacher"], "skipped": report["skipped"],
        "unknown_teacher_cells": report["unknown_teacher_cells"],
        "unknown_subjects": [{"token": tok, "count": n, "where": ex,
                              "guess": guesses.get(next(iter(difflib.get_close_matches(key, list(guesses), n=1,
                                                                                       cutoff=0.6)), ""), "")
                              if len(key) > 1 else ""}  # một chữ cái ("T"): không đoán
                             for key, (n, ex, tok) in sorted(unknown_s.items(), key=lambda x: -x[1][0])],
        "unknown_teachers": [{"token": tok, "count": n, "where": ex,
                              "candidates": [names.teachers[t].code for t in names.candidates(tok)]}
                             for key, (n, ex, tok) in sorted(unknown_t.items(), key=lambda x: -x[1][0])],
    }


def _grid(problem, lessons: dict[tuple[str, int, int], list[tuple[str, str]]]) -> list[list]:
    """Lưới TKB đã xếp (như writer._write_saved, không có mã kết quả): Lớp | Tiết | các ngày, ô "môn" xuống dòng
    "Mã GV"; tiết không có trong ngày ghi config.OFF_LABEL."""
    days = khung_gio.days()
    in_day = {d: {p for s in khung_gio.sessions(d) for p in s.periods} for d in days}
    rows = [["Lớp", "Tiết", *(config.DAYS[d] for d in days)]]
    for cls in sorted(problem.classes, key=class_sort_key):
        for i, p in enumerate(range(1, khung_gio.max_periods() + 1)):
            row = [cls if i == 0 else None, p]
            for d in days:
                items = lessons.get((cls, d, p))
                if p not in in_day[d]:
                    row.append(config.OFF_LABEL)
                elif items:
                    row.append("\n".join(f"{problem.subject_label(s)}\n{problem.teachers[t].code}" for s, t in items))
                else:
                    row.append(None)
            rows.append(row)
    return rows


def describe(report: dict) -> list[str]:
    """Các dòng báo cáo để in ra màn hình (dòng lệnh) hay ghi trên trang."""
    lines = [f"Đọc {report['cells']} ô của {len(report['classes'])} lớp."]
    if report["missing"]:
        lines.append(f"Không có cột của lớp: {', '.join(report['missing'])}.")
    found, want = report["sessions"]
    if found != want:
        lines.append(f"TKB có {found} buổi, khung giờ có {want} buổi: kiểm tra lại khung giờ (bước 1).")
    if report["outside"]:
        lines.append(f"{report['outside']} ô ngoài khung giờ, bỏ qua (vd {'; '.join(report['outside_examples'])}).")
    if report["double"]:
        lines.append(f"{report['double']} ô ghi hai môn: lấy môn đầu.")
    if report["no_teacher"]:
        lines.append(f"{report['no_teacher']} ô của lớp không có GVCN mà không ghi tên giáo viên: bỏ qua.")
    for item in report["unknown_subjects"]:
        lines.append(f"Chữ {item['token']!r} ({item['count']} ô, vd {item['where'][0]}) không khớp môn nào"
                     + (f" (ý là {item['guess']!r}?)" if item["guess"] else "") + ".")
    for item in report["unknown_teachers"]:
        who = f": có thể là {', '.join(item['candidates'])}" if item["candidates"] else ""
        lines.append(f"Tên {item['token']!r} ({item['count']} ô, vd {item['where'][0]}) không khớp đúng một giáo "
                     f"viên{who}.")
    return lines


def main(argv: list[str] | None = None) -> int:
    from .__main__ import use_utf8_output
    from .allocation import build_problem
    from .program import read_program
    from .rules import applied, read_rules
    from .staff import read_staff
    use_utf8_output()
    ap = argparse.ArgumentParser(prog="python -m tkb.nhap_tkb",
                                 description="Nhập TKB đang dùng (dạng bảng tin) vào file vào, làm TKB đã xếp")
    ap.add_argument("board", help="File TKB đang dùng: mỗi lớp một cột, mỗi tiết ghi môn (và tên GV khác GVCN)")
    ap.add_argument("staff", help="File vào của trường (mẫu V8)")
    ap.add_argument("-o", "--output", help="File vào kèm sheet TKB đã xếp (mặc định <file vào>_tkb_dang_dung.xlsx)")
    ap.add_argument("--max-overtime", type=int, default=config.OVERTIME_MAX,
                    help=f"Số tiết bù tối đa mỗi người, như khi xếp (mặc định {config.OVERTIME_MAX})")
    ap.add_argument("--ghep", action="append", default=[], metavar="CHỮ=TÊN",
                    help="Ghép tay một chữ không tự khớp: tên môn, hoặc Mã GV/tên giáo viên; để trống sau dấu = là "
                         "bỏ qua (vd --ghep T=Toán --ghep GDTC=\"Thể Dục\" --ghep Stem+=)")
    args = ap.parse_args(argv)
    try:
        values = read_rules(args.staff)
        with applied(values):
            curriculum = read_program(args.staff)
            staff = read_staff(args.staff, subjects=[s for req in curriculum.values() for s in req])
            problem = build_problem(staff, curriculum, {}, overtime_max=args.max_overtime)
            mapping: dict[str, dict] = {"subjects": {}, "teachers": {}}
            names = _Names(problem, {}, None)
            for item in args.ghep:
                token, _, value = item.partition("=")
                kind = "subjects" if not value or names.subject(value) is not None else "teachers"
                mapping[kind][token.strip()] = value.strip()
            grid, report = read_board(args.board, problem, mapping)
    except InputError as exc:
        print(f"LỖI: {exc}", file=sys.stderr)
        return 1
    for line in describe(report):
        print(line)
    if not report["cells"]:
        print("LỖI: không đọc được ô nào", file=sys.stderr)
        return 1
    source = Path(args.staff)
    out = Path(args.output) if args.output else source.with_name(f"{source.stem}_tkb_dang_dung.xlsx")
    wb = openpyxl.load_workbook(args.staff)
    if config.SAVED_SHEET in wb.sheetnames:
        del wb[config.SAVED_SHEET]
    ws = wb.create_sheet(config.SAVED_SHEET)
    for row in grid:
        ws.append(row)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Đã ghi: {out}")
    print(f"Xếp tiếp: python -m tkb {out} (TKB này đúng luật thì giữ nguyên, không thì sửa ít ô nhất)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
