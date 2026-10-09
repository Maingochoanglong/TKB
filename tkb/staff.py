"""Đọc và kiểm tra file Excel danh sách nhân sự.

Chỉ đọc mẫu V8: Họ và Tên | Chức Vụ | Lớp | Số Tiết/Tuần (chức vụ không ghi số, chương trình tự đánh số
theo thứ tự dòng). Cột không bắt buộc: Thai Sản, Hợp Đồng, Cơ sở 2 (ghi "Có" hoặc để trống), Lớp Đang Dạy
(các lớp dạy trong TKB cũ), Buổi Nghỉ (vd "Chiều T5, Sáng T6" hoặc "2 buổi chiều"). Các cột khác bị bỏ qua.

Chức vụ hợp lệ: Chủ Nhiệm, Bộ Môn, Quản Lý, hoặc tên một môn trong sheet CHƯƠNG TRÌNH HỌC (GV chuyên
biệt của môn đó). Danh sách môn lấy từ file vào nên không có danh sách chức vụ cố định trong code.
"""
from __future__ import annotations

import datetime
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

from . import config


class InputError(ValueError):
    """Lỗi dữ liệu đầu vào."""


@dataclass
class Teacher:
    name: str
    title: str  # chức vụ chuẩn hóa, dùng làm mã giáo viên, vd "bộ môn 5"
    role: str  # "chủ nhiệm", "bộ môn", "quản lý", hoặc một chức vụ chuyên biệt
    index: int | None  # số thứ tự trong chức vụ (không áp dụng cho chủ nhiệm)
    class_name: str | None  # lớp chủ nhiệm, vd "1/1" hoặc "1D15"
    max_lessons: int
    supplementary: bool = False
    row: int | None = None  # dòng trong file Excel gốc
    label: str = ""  # chức vụ như ghi trong file, vd "Tiếng Anh"; trống = dùng `role`
    maternity: bool = False  # đang hưởng chế độ thai sản: không dạy bù, chỉ dạy ở cơ sở 2
    contract: bool = False  # GV hợp đồng: nhận tiết bù trước GV cùng loại (config.Weights.overtime_*)
    campus2: bool = False  # GVCN: lớp ở cơ sở 2; GV khác: chỉ dạy ở cơ sở 2
    history: frozenset[str] = frozenset()  # các lớp dạy trong TKB cũ (ưu tiên giữ khối, rồi giữ lớp)
    # Xếp lại ít xáo trộn: các (lớp, môn) GV dạy trong TKB đã xếp nạp lại (solver.solve(previous=...)).
    previous: frozenset[tuple[str, str]] = frozenset()
    off_sessions: frozenset[tuple[int, str]] = frozenset()  # buổi nghỉ cố định: (ngày 0–4, tên buổi)
    off_any: tuple[tuple[str | None, int], ...] = ()  # nghỉ thêm n buổi bất kỳ: (tên buổi, None = buổi nào cũng được; n)

    @property
    def grade(self) -> int | str | None:
        return grade_of(self.class_name) if self.class_name else None

    @property
    def campus2_only(self) -> bool:
        """GV không chủ nhiệm chỉ dạy các lớp ở cơ sở 2 (đánh dấu Cơ sở 2, hoặc đang hưởng thai sản)."""
        return not self.class_name and (self.campus2 or self.maternity)

    @property
    def code(self) -> str:
        """Mã GV hiển thị trong các file ra, vd "Bộ Môn 5", "Chủ Nhiệm 1/1"."""
        return f"{self.label or self.role} {self.class_name or self.index}"


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFC", str(text)).strip().lower()
    return re.sub(r"\s+", " ", text)


def clean_name(text) -> str:
    """Chữ trong file, bỏ khoảng trắng thừa (giữ hoa thường)."""
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(text)).strip())


def subject_key(name) -> str:
    """Khóa so khớp tên môn/chức vụ: không phân biệt hoa thường, dấu câu và chữ "và"."""
    words = re.sub(r"[^\w\s]", " ", normalize(name)).split()
    return " ".join(w for w in words if w != "và")


SPECIAL_ROLES = (config.ROLE_HOMEROOM, config.ROLE_GENERAL, config.ROLE_MANAGER)


def role_errors(teachers: list[Teacher], subjects) -> list[str]:
    """Chức vụ không phải Chủ Nhiệm/Bộ Môn/Quản Lý, không có trong sheet CHỨC VỤ (config.CUSTOM_ROLES) và không trùng
    tên môn nào của chương trình học."""
    keys = {subject_key(s) for s in subjects} | {subject_key(r.name) for r in config.CUSTOM_ROLES}
    bad: dict[str, list[str]] = {}
    for t in teachers:
        if t.role not in SPECIAL_ROLES and subject_key(t.role) not in keys:
            bad.setdefault(t.label or t.role, []).append(str(t.row) if t.row else "?")
    valid = ", ".join(config.ROLE_LABELS[r] for r in SPECIAL_ROLES)
    return [f"Dòng {', '.join(rows)}: chức vụ '{label}' không xác định (hợp lệ: {valid}, một chức vụ của sheet "
            f"{config.ROLES_SHEET} hoặc đúng tên một môn trong sheet {config.PROGRAM_SHEET})"
            for label, rows in bad.items()]


def canonical_title(role: str, index: int | None, class_name: str | None) -> str:
    return f"{role} {class_name}" if role == config.ROLE_HOMEROOM else f"{role} {index}"


_CLASS_RE = re.compile(r"^(\d+)\s*/\s*(\d+)$")
_CLASS_NAMED_RE = re.compile(r"^(\d+)\s*([^\W\d_]+)\s*(\d*)$")  # khối + chữ + số, vd "1D15", "2A"


def _numbered_class(text: str) -> str | None:
    """Tên lớp theo cách ghi có khối ở đầu: "1/1", "1 / 1" -> "1/1"; "1d15" -> "1D15"; cách ghi khác: None."""
    if m := _CLASS_RE.match(text):
        return f"{int(m.group(1))}/{int(m.group(2))}"
    if m := _CLASS_NAMED_RE.match(text):
        return f"{int(m.group(1))}{m.group(2).upper()}{m.group(3)}"
    return None


def canonical_class(value) -> str:
    """Tên lớp ở sheet LỚP: "1 / 1" viết thành "1/1", tên khác giữ như ghi (vd "1 Blue", "Lá 2")."""
    if m := _CLASS_RE.match(normalize(value)):
        return f"{int(m.group(1))}/{int(m.group(2))}"
    return clean_name(value)


_school_cache: tuple = (None, {}, {})


def _school() -> tuple[dict[str, config.SchoolClass], dict[str, str]]:
    """Các lớp của sheet LỚP (config.CLASSES): ({tên: lớp}, {tên viết thường bỏ dấu: tên}); trống nếu file không có
    sheet LỚP."""
    global _school_cache
    if _school_cache[0] is not config.CLASSES:
        _school_cache = (config.CLASSES, {c.name: c for c in config.CLASSES},
                         {_fold(c.name): c.name for c in config.CLASSES})
    return _school_cache[1], _school_cache[2]


def parse_class(value) -> str:
    """Cột Lớp: một lớp của sheet LỚP (tên tùy ý); file không có sheet LỚP thì khối/số thứ tự, vd "1/1", hoặc khối +
    tên lớp, vd "1D15" (khối là các chữ số đầu)."""
    if isinstance(value, (datetime.date, datetime.datetime)):
        raise InputError(f"cột Lớp bị Excel đổi thành ngày tháng ({value:%d/%m}); hãy chọn lớp từ danh sách "
                         f"thả xuống hoặc định dạng cột Lớp là Text")
    names = _school()[1]
    if names:  # tên như ghi ở sheet LỚP, không phân biệt hoa thường; "1d15" cũng khớp lớp "1D15"
        for key in (canonical_class(value), _numbered_class(normalize(value))):
            if key and _fold(key) in names:
                return names[_fold(key)]
        raise InputError(f"lớp {clean_name(value)!r} không có trong sheet {config.CLASSES_SHEET}")
    name = _numbered_class(normalize(value))
    if name is None:
        raise InputError(f"cột Lớp phải ghi dạng khối/số thứ tự (vd '1/1') hoặc khối rồi tên lớp (vd '1D15'), "
                         f"đang ghi {value!r}; muốn đặt tên lớp tùy ý thì ghi các lớp ở sheet {config.CLASSES_SHEET}")
    return name


def parse_grade(value) -> int | str:
    """Tên khối (cột Khối <tên> của sheet CHƯƠNG TRÌNH HỌC, cột Khối của sheet LỚP, LUẬT): ghi toàn chữ số thì là số
    (vd 1), không thì giữ chữ như ghi (vd "Lá", "Year 7")."""
    if isinstance(value, bool) or (isinstance(value, (int, float)) and (value != int(value) or value < 0)):
        raise InputError(f"tên khối không hợp lệ: {value!r}")
    if isinstance(value, (int, float)):
        return int(value)
    text = clean_name("" if value is None else value)
    if not text:
        raise InputError("thiếu tên khối")
    return int(text) if text.isdigit() else text


def natural_key(text: str) -> tuple:
    """Thứ tự tự nhiên của chữ: phần số so theo số ("Lá 2" trước "Lá 10"), phần chữ không phân biệt hoa thường, dấu."""
    return tuple((0, int(p), "") if p.isdigit() else (1, 0, _fold(p)) for p in re.findall(r"\d+|\D+", text))


def grade_key(grade) -> tuple:
    """Thứ tự khối: các khối số theo số (1 < 2 < 10), rồi các khối tên chữ theo thứ tự tự nhiên."""
    return (0, grade) if isinstance(grade, int) else (1, natural_key(str(grade)))


def grade_of(class_name: str) -> int | str:
    """Khối của một lớp: cột Khối của sheet LỚP; lớp không có ở đó thì các chữ số đầu tên lớp, vd "1/2" và "1D15"
    đều là khối 1."""
    known = _school()[0].get(class_name)
    if known is not None:
        return known.grade
    m = re.match(r"\d+", class_name)
    if m is None:
        raise InputError(f"không biết lớp {class_name!r} thuộc khối nào (ghi lớp này ở sheet {config.CLASSES_SHEET})")
    return int(m.group())


def class_campus2(class_name: str) -> bool:
    """Lớp ghi Có ở cột Cơ sở 2 của sheet LỚP (lớp ở cơ sở 2 còn có thể đánh dấu ở dòng Chủ Nhiệm)."""
    known = _school()[0].get(class_name)
    return known is not None and known.campus2


def class_sort_key(class_name: str) -> tuple:
    """Sắp lớp theo khối (grade_key), rồi theo tên: "1/2" trước "1/10", "1D9" trước "1D15", "Lá 2" trước "Lá 10"."""
    try:
        grade = grade_key(grade_of(class_name))
    except InputError:
        grade = (2,)
    m = re.match(r"(\d+)\D*?([^\W\d_]*)(\d*)$", class_name)
    return grade, ((0, m.group(2), int(m.group(3) or 0)) if m else (1, natural_key(class_name)))


def _fold(text) -> str:
    """Chữ thường, bỏ dấu tiếng Việt (so khớp "Chiều T5" với "chieu thu 5")."""
    text = unicodedata.normalize("NFD", normalize(text)).replace("đ", "d")
    return "".join(ch for ch in text if not unicodedata.combining(ch))


_YES = {"co", "x", "1", "true", "yes", "c"}
_NO = {"khong", "0", "false", "no", "k"}


def parse_yes(value, column: str) -> bool:
    """Cột Có/Không: trống hoặc "Không" là không; "Có", "x", "1" là có."""
    if isinstance(value, bool):
        return value
    if _blank(value):
        return False
    key = _fold(value)
    if key in _YES:
        return True
    if key in _NO:
        return False
    raise InputError(f"cột {column} chỉ ghi 'Có' hoặc để trống, đang ghi {value!r}")


def parse_classes(value) -> frozenset[str]:
    """Cột Lớp Đang Dạy: các lớp cách nhau bằng dấu phẩy hoặc chấm phẩy."""
    if _blank(value):
        return frozenset()
    return frozenset(parse_class(part) for part in re.split(r"[,;]", str(value)) if part.strip())


_ANY_OFF_RE = re.compile(r"^(\d+)\s*buoi(?:\s+(.+?))?(?:\s+bat k[yi])?$")  # "2 buoi chieu", "1 buoi"


def off_text(t: Teacher) -> str:
    """Cột Buổi Nghỉ viết lại từ dữ liệu đã đọc, vd "Chiều T5, Sáng T6, 2 buổi chiều"."""
    from .khung_gio import short_day
    fixed = [f"{name} {short_day(day)}" for day, name in sorted(t.off_sessions)]
    extra = [f"{n} buổi" + (f" {name.lower()}" if name else "") for name, n in t.off_any]
    return ", ".join(fixed + extra)


def parse_off(value) -> tuple[frozenset[tuple[int, str]], tuple[tuple[str | None, int], ...]]:
    """Cột Buổi Nghỉ: các mục cách nhau bằng dấu phẩy/chấm phẩy. Mỗi mục là buổi cố định ("Chiều T5",
    "Sáng thứ 6") hoặc số buổi bất kỳ ("2 buổi chiều", "1 buổi sáng", "2 buổi": buổi nào cũng được).
    Buổi không có trong khung giờ (vd chiều Thứ 6) vốn đã nghỉ nên bỏ qua."""
    from .khung_gio import session_name, split_session_day
    fixed: set[tuple[int, str]] = set()
    count: dict[str | None, int] = defaultdict(int)
    if _blank(value):
        return frozenset(), ()
    for part in re.split(r"[,;]", str(value)):
        item = _fold(part)
        if not item:
            continue
        if found := split_session_day(part):
            name, day = found
            if any(s.name == name for s in config.DAY_SESSIONS.get(day, ())):
                fixed.add((day, name))
        elif (m := _ANY_OFF_RE.match(item)) and (m.group(2) is None or session_name(m.group(2))):
            count[session_name(m.group(2)) if m.group(2) else None] += int(m.group(1))
        elif session_name(item.split(" ")[0]) and not split_session_day(part):
            raise InputError(f"cột Buổi Nghỉ: không có ngày học nào ({', '.join(config.DAYS)}) là {part.strip()!r}")
        else:
            raise InputError(f"cột Buổi Nghỉ ghi buổi cố định (vd 'Chiều T5') hoặc số buổi (vd '2 buổi chiều'), "
                             f"đang ghi {part.strip()!r}")
    for name, n in count.items():
        total = sum(1 for d, ss in config.DAY_SESSIONS.items() for s in ss
                    if (name is None or s.name == name) and (d, s.name) not in fixed)
        if n > total:
            kind = f"buổi {name.lower()}" if name else "buổi"
            raise InputError(f"cột Buổi Nghỉ: xin nghỉ {n} {kind} nhưng chỉ còn {total} {kind} để chọn")
    return frozenset(fixed), tuple(sorted(count.items(), key=lambda kv: (kv[0] is None, kv[0] or "")))


def find_sheet(wb, name: str):
    """Sheet có tên `name` (không phân biệt hoa thường), hoặc None."""
    return next((ws for ws in wb.worksheets if normalize(ws.title) == normalize(name)), None)


def staff_sheet(wb):
    """Sheet nhân sự: sheet tên "NHÂN SỰ" nếu có, không thì sheet đầu tiên."""
    return find_sheet(wb, config.STAFF_SHEET) or wb.worksheets[0]


@dataclass
class SavedTimetable:
    """TKB đã xếp đọc từ file vào cập nhật (sheet config.SAVED_SHEET)."""
    rows: list[tuple]  # (lớp, thứ, tiết, môn, Mã GV, tiết bù?, ô khóa?, vị trí trong sheet) như chữ trong file
    result_code: str | None  # mã kết quả lúc xếp
    rules_code: str | None  # mã các quy định lúc xếp (rules.code)
    # Mọi ô của lưới, cả ô trống (trừ ô "Nghỉ"): (lớp, thứ, tiết, dòng, cột, chỉ số trong rows hoặc None nếu ô
    # trống), dòng và cột đếm từ 1 như Excel; giao diện đổi chữ của ô theo vị trí này.
    places: list[tuple] = field(default_factory=list)


def read_saved_timetable(path: str | Path) -> SavedTimetable | None:
    """Sheet config.SAVED_SHEET của file vào (TKB đã xếp, xem parse_saved_grid). File không có sheet này: None."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = find_sheet(wb, config.SAVED_SHEET)
    if ws is None:
        return None
    return parse_saved_grid([list(row) for row in ws.iter_rows(values_only=True)])


def parse_saved_grid(grid: list[list]) -> SavedTimetable:
    """TKB đã xếp dạng lưới (các dòng của sheet config.SAVED_SHEET; giao diện giữ chúng trong kịch bản): bảng Lớp |
    Tiết | Thứ 2 …, tên lớp ở dòng đầu mỗi khối, mỗi ô "môn" xuống dòng "Mã GV", thêm config.SAVED_OVERTIME ở tiết
    bù, config.SAVED_LOCKED ở ô khóa; mã kết quả và mã quy định ở các dòng trên bảng. Không có bảng: không có dòng
    nào."""
    def cell(r: int, c: int):
        row = grid[r - 1] if 0 < r <= len(grid) else []
        return row[c - 1] if 0 < c <= len(row) else None

    width = max((len(row) for row in grid), default=0)
    codes: dict[str, str] = {}
    header = None
    for r in range(1, min(len(grid), 10) + 1):
        cells = {c: cell(r, c) for c in range(1, width + 1) if not _blank(cell(r, c))}
        for c, value in cells.items():
            if normalize(value) in map(normalize, config.SAVED_CODES) and not _blank(cell(r, c + 1)):
                codes[normalize(value)] = str(cell(r, c + 1)).strip()
        heads = {normalize(v): c for c, v in cells.items()}
        if header is None and "lớp" in heads and "tiết" in heads:
            header = r, heads["lớp"], heads["tiết"], {c: str(v) for c, v in cells.items() if _fold(v).startswith("thu")}
    result_code, rules_code = (codes.get(normalize(k)) for k in config.SAVED_CODES)
    if header is None:
        return SavedTimetable([], result_code, rules_code)
    header_row, class_col, period_col, day_cols = header
    rows, places, cls = [], [], None
    for r in range(header_row + 1, len(grid) + 1):
        if not _blank(cell(r, class_col)):
            cls = clean_name(cell(r, class_col))
        period = cell(r, period_col)
        for c, day in day_cols.items():
            value = cell(r, c)
            if not _blank(value) and normalize(value) == normalize(config.OFF_LABEL):
                continue
            place = (cls, day, period, r, c)
            if _blank(value):
                places.append((*place, None))
                continue
            lines = [line.strip() for line in str(value).splitlines() if line.strip()]
            flags = {}
            for flag in (config.SAVED_LOCKED, config.SAVED_OVERTIME):  # "Bộ Môn 2 (bù) (khóa)": cắt từ cuối
                for i, line in enumerate(lines):
                    if _fold(line).endswith(_fold(flag)):
                        lines[i] = line[:-len(flag)].strip()
                        flags[flag] = True
            lines = [line for line in lines if line]
            if not lines:
                places.append((*place, None))
                continue
            code = lines[-1] if len(lines) > 1 else ""
            places.append((*place, len(rows)))
            rows.append((cls, day, period, lines[0], code, flags.get(config.SAVED_OVERTIME, False),
                         flags.get(config.SAVED_LOCKED, False), f"dòng {r}, {clean_name(day)}"))
    return SavedTimetable(rows, result_code, rules_code, places)


def _blank(value) -> bool:
    return value is None or str(value).strip() == ""


# Cột không bắt buộc: khóa -> tiêu đề trong file.
OPTIONAL_COLUMNS = {"maternity": "Thai Sản", "contract": "Hợp Đồng", "campus2": "Cơ sở 2",
                    "history": "Lớp Đang Dạy", "off": "Buổi Nghỉ"}


def _find_columns(ws) -> tuple[int, dict[str, int]]:
    """Dòng tiêu đề và vị trí các cột (mẫu V8). Bắt buộc: Họ và Tên, Chức Vụ, Số Tiết/Tuần; không bắt buộc: Lớp,
    STT, Thai Sản, Hợp Đồng, Cơ sở 2, Lớp Đang Dạy, Buổi Nghỉ."""
    wanted = {"họ và tên": "name", "chức vụ": "title", "số tiết/tuần": "lessons", "số tiết / tuần": "lessons",
              "lớp": "class", "stt": "stt", **{normalize(label): key for key, label in OPTIONAL_COLUMNS.items()}}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 20)):
        found = {}
        for cell in row:
            if cell.value is None:
                continue
            key = normalize(cell.value)
            if key in wanted:
                found[wanted[key]] = cell.column
        if {"name", "title", "lessons"} <= found.keys():
            return row[0].row, found
    raise InputError("Không tìm thấy dòng tiêu đề có đủ các cột 'Họ và Tên', 'Chức Vụ', 'Số Tiết/Tuần' (mẫu V8)")


def _to_lessons(value, title: str) -> int:
    if isinstance(value, bool) or value is None:
        raise InputError(f"Số tiết của '{title}' bị trống hoặc không hợp lệ")
    try:
        number = float(str(value).strip())
    except ValueError:
        raise InputError(f"Số tiết của '{title}' không phải số: {value!r}") from None
    if number != int(number) or number < 0:
        raise InputError(f"Số tiết của '{title}' phải là số nguyên không âm: {value!r}")
    return int(number)


def make_teacher(name, role: str, index: int | None, class_name: str | None, lessons,
                 row: int | None = None, label: str = "", **extra) -> Teacher:
    """`extra`: các trường không bắt buộc của Teacher (maternity, contract, campus2, history, off_sessions, off_any)."""
    title = canonical_title(role, index, class_name)
    return Teacher(
        name=str(name).strip() if name is not None else "",
        title=title,
        role=role,
        index=index,
        class_name=class_name,
        max_lessons=_to_lessons(lessons, f"{label or role} {class_name or index}"),
        row=row,
        label=label,
        **extra,
    )


def _check_extras(teachers: list[Teacher]) -> list[str]:
    """Lỗi của các cột không bắt buộc cần cả danh sách mới kiểm được (lớp trong Lớp Đang Dạy, thai sản, buổi nghỉ)."""
    errors: list[str] = []
    classes = set(class_list(teachers))
    for t in teachers:
        unknown = sorted(t.history - classes, key=class_sort_key)
        if unknown:
            where = f"không có trong sheet {config.CLASSES_SHEET}" if config.CLASSES else "không có Chủ Nhiệm nào"
            errors.append(f"Dòng {t.row}: cột Lớp Đang Dạy có lớp {where}: {', '.join(unknown)}")
        if t.class_name and t.maternity and not (t.campus2 or class_campus2(t.class_name)):
            errors.append(f"Dòng {t.row}: GVCN thai sản chỉ dạy ở cơ sở 2 nhưng lớp {t.class_name} không đánh dấu "
                          f"Cơ sở 2")
        if t.class_name:
            # Buổi có tiết luôn do GVCN dạy (HOMEROOM_PERIODS) thì GVCN không nghỉ được.
            held = lambda d, name: any(s.name == name and set(s.periods) & config.HOMEROOM_PERIODS  # noqa: E731
                                       for s in config.DAY_SESSIONS.get(d, ()))
            fixed = sorted((d, name) for d, name in t.off_sessions if held(d, name))
            named = [name for name, _ in t.off_any if name and config.DAY_SESSIONS
                     and all(held(d, name) for d in config.DAY_SESSIONS)]
            if fixed or named:
                what = sorted({name for _, name in fixed} | set(named))
                errors.append(f"Dòng {t.row}: GVCN không nghỉ buổi {', '.join(n.lower() for n in what)} được vì tiết "
                              f"{', '.join(map(str, sorted(config.HOMEROOM_PERIODS)))} luôn do GVCN dạy"
                              + (f" ({', '.join(config.DAYS[d] for d, _ in fixed)})" if fixed else ""))
    return errors


def validate(teachers: list[Teacher]) -> None:
    """Báo mọi lỗi trùng lặp cùng lúc (mỗi lỗi một dòng)."""
    errors: list[str] = []
    homeroom: dict[str, Teacher] = {}
    keys: dict[tuple[str, int], Teacher] = {}
    for t in teachers:
        if t.class_name:
            if t.class_name in homeroom:
                errors.append(f"Lớp {t.class_name} có hai Chủ Nhiệm (dòng {homeroom[t.class_name].row} "
                              f"và {t.row})")
            else:
                homeroom[t.class_name] = t
        else:
            key = (t.role, t.index)
            if key in keys:
                errors.append(f"Trùng chức vụ '{t.role} {t.index}' (dòng {keys[key].row} và {t.row})")
            else:
                keys[key] = t
    if not homeroom and not config.CLASSES:
        errors.append(f"Không có Chủ Nhiệm nào nên không xác định được danh sách lớp (hoặc ghi các lớp ở sheet "
                      f"{config.CLASSES_SHEET})")
    if errors:
        raise InputError("\n".join(errors))


def read_staff(path: str | Path, subjects=None) -> list[Teacher]:
    """Đọc sheet nhân sự; `subjects`: các môn của chương trình học để kiểm tra chức vụ (None = không kiểm)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = staff_sheet(wb)
    header_row, cols = _find_columns(ws)

    def cell(r: int, key: str):
        return ws.cell(r, cols[key]).value if key in cols else None

    rows: list[list] = []  # [dòng, tên, chức vụ, stt, lớp, số tiết, chữ chức vụ trong file, cột không bắt buộc]
    errors: list[str] = []
    for r in range(header_row + 1, ws.max_row + 1):
        name, title, lessons, cls = cell(r, "name"), cell(r, "title"), cell(r, "lessons"), cell(r, "class")
        if _blank(title):
            if not all(_blank(v) for v in (name, lessons, cls)):
                errors.append(f"Dòng {r}: thiếu chức vụ")
            continue
        try:
            if re.search(r"\d", str(title)):
                raise InputError(f"Chức Vụ không ghi số thứ tự (chương trình tự đánh số theo thứ tự dòng): {title!r}")
            role, label = normalize(title), clean_name(title)
            index = class_name = None
            if role == config.ROLE_HOMEROOM:
                if _blank(cls):
                    raise InputError(f"Chủ Nhiệm phải ghi Lớp (một lớp của sheet {config.CLASSES_SHEET})"
                                     if config.CLASSES else "Chủ Nhiệm phải ghi Lớp dạng khối/số thứ tự, vd '1/1'")
                class_name = parse_class(cls)
            elif not _blank(cls):
                raise InputError(f"chỉ Chủ Nhiệm mới ghi Lớp (chức vụ đang là {title!r})")
            off_sessions, off_any = parse_off(cell(r, "off"))
            extra = {key: parse_yes(cell(r, key), OPTIONAL_COLUMNS[key]) for key in ("maternity", "contract", "campus2")}
            extra.update(history=parse_classes(cell(r, "history")), off_sessions=off_sessions, off_any=off_any)
            rows.append([r, name, role, index, class_name, lessons, label, extra])
        except InputError as exc:
            errors.append(f"Dòng {r}: {exc}")

    # Đánh số thứ tự trong từng chức vụ theo thứ tự dòng (Chủ Nhiệm dùng lớp thay cho số).
    count: dict[str, int] = defaultdict(int)
    for row in rows:
        if row[2] != config.ROLE_HOMEROOM:
            count[row[2]] += 1
            row[3] = count[row[2]]

    teachers: list[Teacher] = []
    for r, name, role, index, class_name, lessons, label, extra in rows:
        try:
            teachers.append(make_teacher(name, role, index, class_name, lessons, row=r, label=label, **extra))
        except InputError as exc:
            errors.append(f"Dòng {r}: {exc}")
    try:
        validate(teachers)
    except InputError as exc:
        errors.append(str(exc))
    errors += _check_extras(teachers)
    if subjects is not None:
        errors += role_errors(teachers, subjects)
    errors = [line for err in errors for line in err.split("\n")]
    if errors:
        raise InputError(f"File nhân sự có {len(errors)} lỗi:\n  " + "\n  ".join(errors))
    return teachers


def class_list(teachers: list[Teacher]) -> list[str]:
    """Các lớp của trường: sheet LỚP (lớp có thể không có GVCN); file không có sheet LỚP thì lớp của các dòng Chủ
    Nhiệm."""
    names = [c.name for c in config.CLASSES] or [t.class_name for t in teachers if t.class_name]
    return sorted(names, key=class_sort_key)
