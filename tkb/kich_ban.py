"""Kịch bản của một trường cho giao diện (tkb/giao_dien): toàn bộ nội dung file vào V8 dưới dạng dữ liệu JSON.

Kịch bản chỉ là các bảng của file vào (NHÂN SỰ, CHƯƠNG TRÌNH HỌC kèm cột quy định, LỚP, PHÒNG, CHỨC VỤ, QUY ĐỊNH,
LUẬT),
mỗi cột quy định lấy từ tkb/rules.py (`Col`: tiêu đề, khóa, loại ô, ghi chú) nên thêm một quy định vào rules.py là
giao diện tự có ô nhập. Xuất kịch bản ra Excel dùng đúng hàm ghi file mẫu (tkb/template.py); kiểm tra
kịch bản = xuất ra file tạm rồi đọc lại bằng chính các hàm đọc của chương trình, nên giao diện và dòng lệnh luôn hiểu
file vào như nhau.

Kịch bản:
    staff    : [{name, role, class, lessons, maternity, contract, campus2, history, off, extra_roles, tags,
               overtime_order}] theo thứ tự dòng
    grades   : [1, 2, "Lá", ...] các cột Khối <tên> (tên khối toàn chữ số là số)
    subjects : [{name, lessons: {"1": số tiết hoặc None, ...}, rules: {khóa Col: giá trị}}]
    classes  : [{name, grade, campus, tags}] sheet LỚP (trống: lớp là lớp của các dòng Chủ Nhiệm); campus: tên cơ sở
    rooms    : [{name, campus, subjects, grades, capacity}] sheet PHÒNG (phòng học dùng chung; trống: không có), chữ
               như trong ô, capacity số hoặc None
    roles    : [{name, subjects: [tên môn]}] chức vụ GV chuyên biệt (sheet CHỨC VỤ); đọc file thì thêm các chức vụ
               trùng tên môn nhân sự đang dùng (ghi ra sheet CHỨC VỤ thì như không ghi)
    general  : {khóa Col: giá trị} (bảng Quy định | Giá trị)
    days     : [{khóa Col: giá trị}] Thứ 2 … Thứ 7
    periods  : [{khóa Col: giá trị}] tiết 1, 2, ...
    rules    : [{group_label, kind, scope, subject, ...}] mọi luật, kể cả luật có sẵn (sheet LUẬT, các cột
               luat_rieng.COLUMNS): chữ như trong ô, hard true/false, number/level/points số hoặc None
    saved    : các dòng của sheet TKB đã xếp (file vào cập nhật) hoặc None
Giá trị theo loại ô: "yes" → true/false; "int", "order" → số hoặc None; "text" → chữ; "grade" (tên khối) → số nếu ghi
toàn chữ số, không thì chữ, hoặc None.
"""
from __future__ import annotations

import datetime
import hashlib
import json
import re
import tempfile
from contextlib import ExitStack
from pathlib import Path

import openpyxl

from . import bo_ghep, bo_mau, config, khung_gio, luat_co_san, luat_rieng
from .program import grade_columns, read_program
from .rules import (CLASS_CAMPUS, CLASS_CAMPUS2, CLASS_GRADE, CLASS_NAME, CLASS_TAGS, DAY_COLS, DAY_KEY, GENERAL,
                    GENERAL_KEY, NO, PERIOD_COLS, PERIOD_KEY, ROOM_CAMPUS, ROOM_CAPACITY, ROOM_GRADES, ROOM_NAME,
                    ROOM_SUBJECTS, SESSION_PREFIX, VALUE, YES, applied, luat_row, mau as rules_mau, read_rules,
                    rule_tables, subject_columns, visible)
from .rules import SUBJECT_COLS as _ALL_SUBJECT_COLS
from .staff import (CAMPUS1, CAMPUS2, InputError, _find_columns, _fold, _NO, _YES, class_list, class_sort_key,
                    clean_name, find_sheet, grade_key, grade_of, normalize, parse_grade, read_staff, staff_sheet,
                    subject_key)
from .template import NOTES, STAFF_HEADERS, program_rows, write_input

VERSION = 3  # 2: mọi luật ở `rules` (sheet LUẬT); 3: khung giờ là `sessions` + số tiết từng buổi của mỗi ngày
SUBJECT_COLS = visible(_ALL_SUBJECT_COLS)  # cột quy định của môn còn dùng (số của luật nay ở sheet LUẬT)
GENERAL_COLS = visible(GENERAL)
DAY_COLS_V = visible(DAY_COLS)  # các cột HĐTN của bảng Ngày (số tiết từng buổi: day["periods"])
# Tên gợi ý khi thêm ngày học trên trang (đổi được tùy ý).
DAY_SUGGESTIONS = ["Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "Chủ nhật"]
RULE_KEYS = (luat_rieng.GROUP[0], *(k for k, _ in luat_rieng.COLUMNS))
# Các cột của sheet NHÂN SỰ: (khóa trong kịch bản, khóa của staff._find_columns, loại ô); tiêu đề là STAFF_HEADERS.
STAFF_COLS = (("name", "name", "text"), ("role", "title", "role"), ("class", "class", "text"),
              ("lessons", "lessons", "int"), ("maternity", "maternity", "yes"), ("contract", "contract", "yes"),
              ("campus2", "campus2", "yes"), ("history", "history", "text"), ("off", "off", "text"),
              ("extra_roles", "extra_roles", "text"), ("tags", "tags", "text"),
              ("overtime_order", "overtime_order", "int"))
_HEADERS = {key: head for (key, _, _), head in zip(STAFF_COLS, STAFF_HEADERS)}
ROLES = [config.ROLE_LABELS[r] for r in (config.ROLE_HOMEROOM, config.ROLE_GENERAL, config.ROLE_MANAGER)]
# Trang chi tiết môn của giao diện chia các cột quy định của môn thành nhóm; cột chưa có nhóm hiện ở nhóm "Khác", nên
# cột mới thêm ở rules.py vẫn có ô nhập.
SUBJECT_GROUPS = (
    ("Hiển thị", ("DISPLAY_NAMES", "HDTN")),
    ("Giáo viên chủ nhiệm", ("HOMEROOM_PRIORITY", "HOMEROOM_CUT_ORDER", "HOMEROOM_FILL_ORDER", "HOMEROOM_ONLY_SUBJECTS",
                             "HOMEROOM_OVERTIME_SPECIALIST")),
    ("Ai được dạy", ("GENERAL_FORBIDDEN_SUBJECTS", "MANAGER_RULES")),
    ("Nhóm môn, nhãn của luật", ("group", "extra", "PAIR_EXCLUDED")),
    ("Ưu tiên khi xếp (mềm)", ("HEAVY_SUBJECTS", "MORNING_SUBJECTS")),
)
# Các cột của môn mà bước Chức vụ sửa theo chức vụ có sẵn (Chủ Nhiệm, Bộ Môn, Quản Lý).
ROLE_RULES = {"homeroom_take": "HOMEROOM_PRIORITY", "homeroom_only": "HOMEROOM_ONLY_SUBJECTS",
              "homeroom_fill": "HOMEROOM_FILL_ORDER", "general_forbidden": "GENERAL_FORBIDDEN_SUBJECTS",
              "manager_grade": "MANAGER_RULES"}


def _col(c) -> dict:
    return {"key": c.key, "header": c.header, "kind": c.kind, "note": c.note}


def schema() -> dict:
    """Mô tả các bảng, cột cho giao diện (sinh ô nhập theo đây)."""
    return {
        "version": VERSION,
        "staff": [{"key": key, "header": _HEADERS[key], "kind": kind, "note": NOTES[_HEADERS[key]]}
                  for key, _, kind in STAFF_COLS],
        "roles": ROLES,
        "role_rules": ROLE_RULES,
        "subject": [_col(c) for c in SUBJECT_COLS],
        "subject_groups": [{"label": label, "keys": list(keys)} for label, keys in SUBJECT_GROUPS],
        "general": [_col(c) for c in GENERAL_COLS],
        "day": [_col(c) for c in DAY_COLS_V],
        "period": [_col(c) for c in PERIOD_COLS],
        "days": DAY_SUGGESTIONS,
        "presets": [{"key": k, "name": name, "note": note} for k, (name, note) in bo_mau.PRESETS.items()],
        "session_prefix": SESSION_PREFIX,
        "custom": {"columns": [{"key": k, "header": h} for k, h in (luat_rieng.GROUP, *luat_rieng.COLUMNS)],
                   "kinds": [{"key": k.key, "label": k.label, "needs": list(k.needs), "uses": list(k.uses),
                              "note": k.note, "family": k.family} for k in luat_rieng.KINDS],
                   "levels": list(luat_rieng.LEVELS),
                   "sessions": khung_gio.session_names(),  # mặc định; trang dùng các buổi của kịch bản
                   "groups": [*luat_co_san.GROUPS, luat_co_san.CUSTOM_GROUP],
                   "custom_group": luat_co_san.CUSTOM_GROUP,
                   "campus_default": CAMPUS1,  # lớp không ghi cơ sở (tên cơ sở cũng là nhãn lớp, staff.class_tags)
                   "campus2": CAMPUS2,  # cơ sở của lớp ghi Có ở cột Cơ sở 2 (dòng Chủ Nhiệm, kịch bản cũ)
                   "structure": luat_co_san.STRUCTURE,
                   "composer": composer()},
        "sheets": {"staff": config.STAFF_SHEET, "program": config.PROGRAM_SHEET, "classes": config.CLASSES_SHEET,
                   "rooms": config.ROOMS_SHEET, "roles": config.ROLES_SHEET,
                   "rules": config.RULES_SHEET, "luat": luat_rieng.RULES_SHEET, "custom": luat_rieng.SHEET,
                   "saved": config.SAVED_SHEET},
        # Chữ đánh dấu ô của sheet TKB đã xếp (staff.parse_saved_grid): giao diện thêm, bỏ khi khóa ô, đổi ô.
        "saved_marks": {"locked": config.SAVED_LOCKED, "overtime": config.SAVED_OVERTIME,
                        "hire": config.SUPPLEMENT_NAME, "off": config.OFF_LABEL},
    }


def composer() -> dict:
    """Từ vựng của bộ ghép luật (tkb/bo_ghep.py) cho giao diện: thêm một chiều, phép đo, nhãn ở Python là trang có."""
    slot = {c.header for c in (*PERIOD_COLS, *DAY_COLS)}
    tags = bo_ghep.tag_names()
    return {"scopes": [{"key": d.key, "label": d.label} for d in bo_ghep.SCOPES],
            "measures": [{"key": m.key, "label": m.label, "ops": list(m.ops), "number": m.number, "other": m.other,
                          "count_by": m.count_by, "sequence": m.sequence, "note": m.note,
                          "default_op": m.default_op, "family": m.family} for m in bo_ghep.MEASURES],
            "families": list(bo_ghep.FAMILIES),
            "derived": list(bo_ghep.DERIVED.values()),
            "ops": [{"key": k, "label": v} for k, v in bo_ghep.OPS.items()],
            "tags": {"subject": [t for t in tags if t not in slot], "slot": [t for t in tags if t in slot]},
            "compose": list(luat_rieng.COMPOSE)}


# ---- đọc file Excel thành kịch bản ----

def _blank(value) -> bool:
    return value is None or str(value).strip() == ""


def _number(value):
    """Số nguyên nếu ô là số nguyên, ô trống là None, còn lại giữ chữ (kiểm tra sẽ báo lỗi)."""
    if _blank(value) or isinstance(value, bool):
        return None if _blank(value) else str(value)
    if isinstance(value, float) and value == int(value):
        return int(value)
    if isinstance(value, int):
        return value
    text = clean_name(value)
    return int(text) if text.isdigit() else text


def _text(value) -> str:
    if _blank(value):
        return ""
    if isinstance(value, float) and value == int(value):
        value = int(value)
    return clean_name(value)


def _read_staff_rows(wb, warnings: list[str]) -> list[dict]:
    """Các dòng của sheet nhân sự đúng như chữ trong file (không kiểm tra; dòng trống bỏ qua). File không có sheet
    nhân sự (vd file chỉ có sheet LUẬT): không có dòng nào."""
    if find_sheet(wb, config.STAFF_SHEET) is None:
        return []
    ws = staff_sheet(wb)
    header_row, cols = _find_columns(ws)
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        raw = {key: ws.cell(r, cols[col]).value if col in cols else None for key, col, _ in STAFF_COLS}
        if all(_blank(v) for v in raw.values()):
            continue
        row = {}
        for key, _, kind in STAFF_COLS:
            value = raw[key]
            if kind == "yes":
                folded = "" if _blank(value) else _fold(value)
                row[key] = value is True or folded in _YES
                if folded and folded not in _YES | _NO and not isinstance(value, bool):
                    warnings.append(f"{ws.title}, dòng {r}: cột {_HEADERS[key]} ghi {value!r}, đã đọc là không")
            elif kind == "int":
                row[key] = _number(value)
            elif isinstance(value, (datetime.date, datetime.datetime)):
                row[key] = f"{value.day}/{value.month}"
                warnings.append(f"{ws.title}, dòng {r}: ô {value:%d/%m/%Y} là ngày tháng (Excel tự đổi), đã đọc "
                                f"thành lớp {row[key]}; hãy kiểm tra lại")
            else:
                row[key] = _text(value)
        rows.append(row)
    return rows


def _read_program_rows(wb) -> tuple[list[int | str], list[tuple[str, dict]]]:
    """Các khối và các dòng (môn, {khối: số tiết}) của sheet CHƯƠNG TRÌNH HỌC đúng như trong file."""
    ws = find_sheet(wb, config.PROGRAM_SHEET)
    if ws is None:  # vd file chỉ có sheet LUẬT: giao diện chỉ lấy các phần file có
        return [], []
    found = grade_columns(ws)
    if found is None:
        raise InputError(f"Sheet {config.PROGRAM_SHEET} phải có cột 'Môn học' và các cột 'Khối <tên>' (vd 'Khối 1')")
    header_row, subject_col, grades = found
    order = sorted(grades, key=grade_key)
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(r, subject_col).value
        if _blank(name) or normalize(name) in ("tổng", "tổng cộng"):
            continue
        rows.append((clean_name(name), {str(g): _number(ws.cell(r, grades[g]).value) for g in order}))
    return order, rows


def _read_room_rows(wb) -> list[dict]:
    """Các phòng của sheet PHÒNG đúng như chữ trong file (dòng trống giữa bảng giữ lại: phòng i là dòng i + 2 của
    sheet). File không có sheet PHÒNG: không có phòng nào."""
    ws = find_sheet(wb, config.ROOMS_SHEET)
    if ws is None:
        return []
    for top in range(1, min(ws.max_row, 20) + 1):
        heads = {normalize(ws.cell(top, c).value): c for c in range(1, ws.max_column + 1)
                 if not _blank(ws.cell(top, c).value)}
        if normalize(ROOM_NAME) in heads:
            break
    else:
        return []
    get = lambda r, h: ws.cell(r, heads[normalize(h)]).value if normalize(h) in heads else None  # noqa: E731
    rows = []
    for r in range(top + 1, ws.max_row + 1):
        row = {"name": _text(get(r, ROOM_NAME)), "campus": _text(get(r, ROOM_CAMPUS)),
               "subjects": _text(get(r, ROOM_SUBJECTS)), "grades": _text(get(r, ROOM_GRADES)),
               "capacity": _number(get(r, ROOM_CAPACITY))}
        rows.append(row)
    while rows and not any(v not in ("", None) for v in rows[-1].values()):
        rows.pop()  # các dòng kẻ sẵn để trống ở cuối bảng
    return rows


def _read_class_rows(wb) -> list[dict]:
    """Các lớp của sheet LỚP đúng như chữ trong file (dòng trống giữa bảng giữ lại: lớp i là dòng i + 2 của sheet).
    File không có sheet LỚP: không có lớp nào."""
    ws = find_sheet(wb, config.CLASSES_SHEET)
    if ws is None:
        return []
    for top in range(1, min(ws.max_row, 20) + 1):
        heads = {normalize(ws.cell(top, c).value): c for c in range(1, ws.max_column + 1)
                 if not _blank(ws.cell(top, c).value)}
        if normalize(CLASS_NAME) in heads:
            break
    else:
        return []
    get = lambda r, h: ws.cell(r, heads[normalize(h)]).value if normalize(h) in heads else None  # noqa: E731
    rows = []
    for r in range(top + 1, ws.max_row + 1):
        name, grade, campus, tags = get(r, CLASS_NAME), get(r, CLASS_GRADE), get(r, CLASS_CAMPUS), get(r, CLASS_TAGS)
        if isinstance(name, (datetime.date, datetime.datetime)):
            name = f"{name.day}/{name.month}"
        old = get(r, CLASS_CAMPUS2)  # cột Có/Không của bản trước
        if _blank(campus) and (old is True or (not _blank(old) and _fold(old) in _YES)):
            campus = CAMPUS2
        rows.append({"name": _text(name), "grade": _grade(grade), "campus": _text(campus), "tags": _text(tags)})
    while rows and not rows[-1]["name"] and rows[-1]["grade"] is None and not rows[-1]["campus"] \
            and not rows[-1]["tags"]:
        rows.pop()
    return rows


def _grade_name(class_name: str) -> str:
    """Tên khối của một lớp để vẽ trang (lớp không rõ khối: chuỗi trống)."""
    try:
        return str(grade_of(class_name))
    except InputError:
        return ""


def _grade(value):
    """Tên khối trong kịch bản: số nếu ghi toàn chữ số, chữ nếu ghi chữ, None nếu trống."""
    if _blank(value):
        return None
    try:
        return parse_grade(value)
    except InputError:
        return _text(value)


def _from_cell(col, value):
    if col.kind == "yes":
        return value == YES
    if col.kind == "text":
        return _text(value)
    return value


def _json_value(value):
    if value is None or isinstance(value, (bool, int, float, str)):
        return value
    return str(value)


def _read_saved(wb) -> list[list] | None:
    ws = find_sheet(wb, config.SAVED_SHEET)
    if ws is None:
        return None
    rows = [[_json_value(v) for v in row] for row in ws.iter_rows(values_only=True)]
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    return [row[:max((i + 1 for i, v in enumerate(row) if v is not None), default=0)] for row in rows]


def rule_dict(rule) -> dict:
    """Một luật (config.CustomRule) theo dạng kịch bản: các ô như trong sheet LUẬT."""
    cells = {**luat_rieng.cells(rule), luat_rieng.GROUP[0]: rule.group_label}
    return {k: (cells[k] == YES if k in ("hard", luat_rieng.OFF)
                else cells[k] if k in ("number", "level", "points") else (cells[k] or "")) for k in RULE_KEYS}


def _rules_rows() -> list[dict]:
    """Mọi luật đang dùng (sheet LUẬT, hoặc các dòng mặc định cộng luật riêng của file cũ) theo dạng kịch bản."""
    return [rule_dict(r) for r in luat_co_san.rows()]


def _roles_part(staff: list[dict], subject_names: list[str]) -> list[dict]:
    """Các chức vụ GV chuyên biệt: các dòng sheet CHỨC VỤ (config.CUSTOM_ROLES), rồi các chức vụ trùng tên môn mà
    nhân sự đang dùng (cột Chức Vụ và Chức Vụ Thêm), theo thứ tự dòng (để giao diện chọn chức vụ từ danh sách)."""
    from .program import canonical_subject

    roles = [{"name": r.name, "subjects": list(r.subjects)} for r in config.CUSTOM_ROLES]
    known = {subject_key(r["name"]) for r in roles} | {subject_key(label) for label in ROLES}
    by_key = {subject_key(canonical_subject(s)): s for s in subject_names}
    by_key.update({subject_key(s): s for s in subject_names})
    for row in staff:
        for name in (row["role"], *(p.strip() for p in re.split(r"[,;+]", str(row.get("extra_roles") or "")))):
            key = subject_key(name) if name else ""
            if key and key not in known and key in by_key:
                roles.append({"name": name, "subjects": [by_key[key]]})
                known.add(key)
    return roles


def _rules_part(subject_names: list[str]) -> tuple[dict, dict, list[dict], dict[str, dict], list[str]]:
    """Các quy định đang dùng (config) theo dạng kịch bản: (chung, khung giờ {sessions, days}, tiết, {môn: quy định},
    các môn có quy định mà chương trình học không có)."""
    general_t, days_t, periods_t = rule_tables()
    general = {c.key: _from_cell(c, value) for c, (_, value) in zip(GENERAL_COLS, general_t[1])}
    names = [h[len(SESSION_PREFIX):].strip() for h in days_t[0][1:] if h.startswith(SESSION_PREFIX + " ")]
    days = [{"name": row[0], "periods": {n: row[1 + i] or 0 for i, n in enumerate(names)},
             **{c.key: _from_cell(c, v) for c, v in zip(DAY_COLS_V, row[1 + len(names):])}} for row in days_t[1]]
    periods = [{c.key: _from_cell(c, v) for c, v in zip(PERIOD_COLS, row[1:])} for row in periods_t[1]]
    _, values, rest = subject_columns(subject_names)
    subjects = {name: {c.key: _from_cell(c, v) for c, v in zip(SUBJECT_COLS, row)} for name, row in values.items()}
    return general, {"sessions": names, "days": days}, periods, subjects, rest


def upgrade(scenario: dict) -> dict:
    """Kịch bản bản trước (bản nháp còn trong trình duyệt) -> bản hiện tại. Bản 2: khung giờ là số tiết buổi sáng,
    buổi chiều chung mọi ngày (general) và cột Có/Không học sáng, chiều của Thứ 2 … Thứ 7."""
    if not isinstance(scenario, dict) or scenario.get("version") != 2:
        return scenario
    general = dict(scenario.get("general") or {})
    m = general.pop("morning_periods", None) or len(config.MORNING.periods)
    a = general.pop("afternoon_periods", None) or len(config.AFTERNOON.periods)
    morning, afternoon = config.MORNING.name, config.AFTERNOON.name
    days = []
    for d, day in enumerate(scenario.get("days") or []):
        on = bool(day.get("morning_days"))
        days.append({"name": f"Thứ {d + 2}",
                     "periods": {morning: m if on else 0, afternoon: a if on and day.get("afternoon_days") else 0},
                     **{c.key: day.get(c.key) for c in DAY_COLS_V}})
    return {**scenario, "version": VERSION, "general": general, "sessions": [morning, afternoon], "days": days}


def from_excel(path: str | Path) -> tuple[dict, list[str]]:
    """Đọc file vào V8 (cả file vào cập nhật *_cap_nhat.xlsx) thành kịch bản: (kịch bản, các cảnh báo). Nhân sự và số
    tiết đọc đúng như chữ trong file (lỗi sẽ hiện khi kiểm tra); quy định đọc bằng rules.read_rules, cột/bảng không
    có thì lấy giá trị mặc định như khi chạy. Quy định sai cách ghi thì báo lỗi (InputError)."""
    warnings: list[str] = []
    wb = openpyxl.load_workbook(path, data_only=True)
    staff = _read_staff_rows(wb, warnings)
    grades, program = _read_program_rows(wb)
    rules = read_rules(path, warn=warnings.append)
    with applied(rules):
        general, frame, periods, by_subject, rest = _rules_part([name for name, _ in program])
        luat = _rules_rows()
        roles = _roles_part(staff, [name for name, _ in program])
    subjects = [{"name": name, "lessons": lessons, "rules": by_subject[name]} for name, lessons in program]
    subjects += [{"name": name, "lessons": {str(g): None for g in grades}, "rules": by_subject[name]}
                 for name in rest]
    scenario = {"version": VERSION, "staff": staff, "grades": grades, "subjects": subjects,
                "classes": _read_class_rows(wb), "rooms": _read_room_rows(wb), "roles": roles, "general": general,
                **frame, "periods": periods, "rules": luat, "saved": _read_saved(wb)}
    return scenario, warnings


def sheets_in(path: str | Path) -> list[str]:
    """Các sheet của file vào V8 mà file có (tên chuẩn, theo thứ tự trong `schema()["sheets"]`): giao diện cho chọn
    phần nào lấy từ file."""
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return [name for name in schema()["sheets"].values() if find_sheet(wb, name) is not None]
    finally:
        wb.close()


def sample_scenario() -> dict:
    """Kịch bản của trường mẫu tên giả (tkb/truong_mau.py) cho nút "Xem thử với trường mẫu": ghi file vào mẫu rồi đọc
    lại như mọi file vào."""
    from .truong_mau import write_sample_input
    with tempfile.TemporaryDirectory() as tmp:
        scenario, _ = from_excel(write_sample_input(Path(tmp) / "truong_mau.xlsx"))
    return scenario


def default_scenario(mau: str | None = None) -> dict:
    """Kịch bản trống như file mẫu (python -m tkb.template): chưa có nhân sự, các môn có quy định của bộ luật mẫu `mau`
    (tkb/bo_mau.py; không ghi: giá trị mặc định, tức bộ Tiểu học Việt Nam) với số tiết để trống."""
    with applied(rules_mau(mau) if mau else None):
        grades, _, rows = program_rows(None)
        names = [row[0] for row in rows]
        general, frame, periods, by_subject, _ = _rules_part(names)
        luat = [rule_dict(r) for r in (luat_co_san.rows() if mau else luat_co_san.default_rows())]
    return {"version": VERSION, "staff": [], "grades": grades,
            "subjects": [{"name": n, "lessons": {str(g): None for g in grades}, "rules": by_subject[n]} for n in names],
            "classes": [], "rooms": [], "roles": [], "general": general, **frame, "periods": periods, "rules": luat,
            "saved": None}


# ---- ghi kịch bản ra file Excel ----

def _to_cell(col, value):
    if col.kind == "yes":
        return YES if value is True or (isinstance(value, str) and _fold(value) in _YES) else NO
    if col.kind == "text":
        return _text(value) or None
    if col.kind == "grade":
        return _grade(value)
    return _number(value)


def _staff_row(row: dict) -> list:
    out = []
    for key, _, kind in STAFF_COLS:
        value = row.get(key)
        if kind == "yes":
            out.append(YES if value is True or (isinstance(value, str) and _fold(value) in _YES) else None)
        elif kind == "int":
            out.append(_number(value))
        else:
            out.append(_text(value) or None)
    return out


def to_excel(scenario: dict, path: str | Path) -> None:
    """Ghi kịch bản ra file vào V8 (NHÂN SỰ, CHƯƠNG TRÌNH HỌC kèm cột quy định, LỚP, CHỨC VỤ, QUY ĐỊNH, LUẬT, HƯỚNG
    DẪN và, nếu có, sheet TKB đã xếp), cùng cách ghi với file mẫu."""
    staff = [_staff_row(row) for row in scenario.get("staff", [])]  # cả dòng trống: dòng i là dòng i + 2 của sheet
    grades = [g for g in map(_grade, scenario.get("grades", [])) if g is not None]
    rows = []
    for subject in scenario.get("subjects", []):
        lessons, rules = subject.get("lessons") or {}, subject.get("rules") or {}
        row = [_text(subject.get("name")) or None, *[_number(lessons.get(str(g))) for g in grades],
               *[_to_cell(c, rules.get(c.key)) for c in SUBJECT_COLS]]
        if row[0] is not None or any(v is not None for v in row[1:1 + len(grades)]):
            rows.append(row)
    scenario = upgrade(scenario)
    general = scenario.get("general") or {}
    sessions = [_text(n) for n in scenario.get("sessions") or []]
    days = scenario.get("days") or []
    periods = scenario.get("periods") or []
    tables = [([GENERAL_KEY, VALUE], [[c.header, _to_cell(c, general.get(c.key))] for c in GENERAL_COLS]),
              ([DAY_KEY, *(f"{SESSION_PREFIX} {n}" for n in sessions), *(c.header for c in DAY_COLS_V)],
               [[_text(day.get("name")) or None,
                 *((_number((day.get("periods") or {}).get(n)) or None) for n in sessions),
                 *(_to_cell(c, day.get(c.key)) for c in DAY_COLS_V)] for day in days]),
              ([PERIOD_KEY, *(c.header for c in PERIOD_COLS)],
               [[p, *(_to_cell(c, row.get(c.key)) for c in PERIOD_COLS)] for p, row in enumerate(periods, start=1)])]
    saved = scenario.get("saved")

    def write_saved(wb):
        ws = wb.create_sheet(config.SAVED_SHEET)
        for row in saved:
            ws.append(row)

    # Cả dòng trống: dòng i của bảng là dòng i + 2 của sheet CHỨC VỤ.
    roles = [[_text(role.get("name")) or None,
              ", ".join(t for t in map(_text, role.get("subjects") or []) if t) or None]
             for role in scenario.get("roles") or []]
    # Cả dòng trống: lớp i là dòng i + 2 của sheet LỚP.
    classes = [[_text(c.get("name")) or None, _grade(c.get("grade")),
                _text(c.get("campus")) or (CAMPUS2 if c.get("campus2") is True else None),  # campus2: kịch bản cũ
                _text(c.get("tags")) or None] for c in scenario.get("classes") or []]
    # Cả dòng trống: phòng i là dòng i + 2 của sheet PHÒNG.
    rooms = [[_text(r.get("name")) or None, _text(r.get("campus")) or None, _text(r.get("subjects")) or None,
              _text(r.get("grades")) or None, _number(r.get("capacity"))] for r in scenario.get("rooms") or []]
    write_input(path, staff, (grades, [c.header for c in SUBJECT_COLS], rows), tables,
                extra=write_saved if saved else None, rules=luat_sheet_rows(scenario.get("rules") or []),
                roles=roles, classes=classes, rooms=rooms)


def luat_sheet_rows(rows: list[dict]) -> list[list]:
    """Các dòng của sheet LUẬT từ các luật của kịch bản (cả dòng trống: dòng i là dòng i + 2 của sheet); cột Luật đọc
    là ghi câu đọc lại của dòng (dòng ghi sai thì để trống)."""
    out = []
    for i, row in enumerate(rows):
        cells = {k: _custom_cell(k, row.get(k)) for k, _ in luat_rieng.COLUMNS}
        rule = luat_rieng.parse(cells, i + 2, lambda _: None)
        group = _text(row.get(luat_rieng.GROUP[0])) or None
        out.append(luat_row(rule) if rule else [group, *cells.values(), None])
        out[-1][0] = group
    return out


def rules_to_excel(scenario: dict, path: str | Path) -> None:
    """File Excel chỉ có các luật (sheet LUẬT) và cách ghi (sheet HƯỚNG DẪN): xuất luật để sửa trong Excel, chép
    sang trường khác, hoặc làm mẫu luật; nhập lại bằng Nhập từ Excel."""
    from .template import write_guide, write_luat_sheet
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_luat_sheet(wb, luat_sheet_rows(scenario.get("rules") or []))
    write_guide(wb)
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _custom_cell(key: str, value):
    """Ô của sheet LUẬT RIÊNG (cả dòng trống: dòng i của bảng là dòng i + 2 của sheet)."""
    if key == "hard":
        return YES if value is True or (isinstance(value, str) and _fold(value) in _YES) else NO
    if key in ("group", luat_rieng.OFF):
        return YES if value is True or (isinstance(value, str) and _fold(value) in _YES) else None
    if key in ("number", "level", "points"):
        return _number(value)
    return _text(value) or None


# ---- kiểm tra ----

def _lines(exc: Exception, sheet: str = "") -> list[str]:
    """Các dòng lỗi của một InputError (dòng tiêu đề "... có n lỗi:" bỏ đi), thêm tên sheet nếu lỗi chưa ghi."""
    lines = [line.strip() for line in str(exc).splitlines() if line.strip()]
    lines = lines[1:] if len(lines) > 1 and lines[0].endswith(":") else lines
    return [line if not sheet or line.startswith(sheet) else f"{sheet}: {line}" for line in lines]


def describe(scenario: dict, rows: list[dict] | None = None, student_rules: bool = True) -> dict:
    """Câu đọc lại của từng luật (rows; mặc định các luật của kịch bản), theo quy định của kịch bản: {rules: [{text,
    errors, native, level}]}; native là khóa luật có sẵn nếu dòng ở dạng gốc (chương trình xếp bằng mã hóa riêng),
    level là mức ưu tiên bằng chữ (dòng ghi Điểm: mức gần nhất). Lỗi các sheet khác không chặn (nút Kiểm tra báo)."""
    from .config import Role
    rows = scenario.get("rules") or [] if rows is None else rows
    values: dict = {}
    with tempfile.TemporaryDirectory() as tmp, ExitStack() as stack:
        path = Path(tmp) / "kich_ban.xlsx"
        to_excel({**scenario, "rules": []}, path)
        try:
            values = read_rules(path) or {}
        except InputError:
            values = {"CUSTOM_ROLES": [Role(clean_name(r.get("name") or ""), ()) for r in scenario.get("roles") or []
                                       if r.get("name")]}
        stack.enter_context(applied(values))
        try:  # câu đọc lại ghi Mã GV thay cho họ tên; nhân sự còn lỗi thì đọc như chức vụ
            bo_ghep.know_staff(read_staff(path))
        except InputError:
            pass
        out = []
        for i, row in enumerate(rows):
            errors: list[str] = []
            cells = {k: _custom_cell(k, row.get(k)) for k, _ in luat_rieng.COLUMNS}
            rule = luat_rieng.parse(cells, i + 2, errors.append)
            native = luat_co_san.native_of(rule) if rule else None
            out.append({"text": luat_rieng.describe(rule) if rule else None, "errors": errors,
                        "native": native.key if native else None,
                        "level": luat_rieng.level_label(rule) if rule and not rule.hard else None})
    return {"rules": out}


def check(scenario: dict, mode: str = config.MODE_OVERTIME, overtime_max: int = config.OVERTIME_MAX,
          student_rules: bool = True, quick: bool = False) -> dict:
    """Kiểm tra kịch bản như khi chạy: ghi ra file tạm, đọc lại bằng các hàm đọc của chương trình, đếm tìm các quy
    định mâu thuẫn (chan_doan.precheck), rồi dự toán (phân công, không xếp giờ, vài giây). Trả về {errors, warnings,
    info}: lỗi chặn việc xếp TKB; info là số liệu và dự toán. quick: dừng trước dự toán (dưới 1 giây; giao diện kiểm
    tra tự động mỗi lần sửa), cùng các lỗi. Đọc theo thứ tự (quy định và luật, chương trình học, nhân sự, rồi các phép
    đếm cần cả ba): gặp lỗi ở bước trước thì dừng, nên kết quả có thêm unchecked: các sheet chưa kiểm được ("" là các
    phép đếm chung), để giao diện không báo "không lỗi" cho phần chưa kiểm. Đổi tạm config (rules.applied) nên không gọi
    song song."""
    from .allocation import build_problem
    from .chan_doan import precheck
    from .phan_cong import phan_cong
    from .solver import ShortageError, SolveError, du_toan_lines

    errors: list[str] = []
    warnings: list[str] = []
    info: list[str] = []

    def result(*unchecked: str) -> dict:
        return {"errors": errors, "warnings": warnings, "info": info, "unchecked": list(unchecked)}

    with tempfile.TemporaryDirectory() as tmp:
        path = Path(tmp) / "kich_ban.xlsx"
        to_excel(scenario, path)
        try:
            rules = read_rules(path, warn=warnings.append)
        except InputError as exc:
            errors += _lines(exc)
            return result(config.PROGRAM_SHEET, config.STAFF_SHEET, luat_rieng.RULES_SHEET, "")
        with applied(rules):
            try:
                curriculum = read_program(path)
            except InputError as exc:
                errors += _lines(exc, config.PROGRAM_SHEET)
                curriculum = None
            try:
                subjects = [s for req in curriculum.values() for s in req] if curriculum else None
                staff = read_staff(path, subjects=subjects)
                bo_ghep.know_staff(staff)  # câu đọc lại của luật ghi Mã GV thay cho họ tên
            except InputError as exc:
                errors += _lines(exc, config.STAFF_SHEET)
                staff = None
            if errors or not curriculum:
                return result(*([] if curriculum else [config.STAFF_SHEET]), luat_rieng.RULES_SHEET, "")
            classes = class_list(staff)
            total = sum(sum(curriculum.get(grade_of(c), {}).values()) for c in classes)
            quota = sum(t.max_lessons for t in staff)
            info.append(f"{len(staff)} nhân sự, {len(classes)} lớp, "
                        f"{len({s for req in curriculum.values() for s in req})} môn; cần dạy {total} tiết/tuần, "
                        f"tổng định mức {quota} tiết/tuần" + (f" (thiếu {total - quota})" if total > quota else ""))
            try:
                base = build_problem(staff, curriculum, {}, overtime_max=overtime_max)
                warnings += base.warnings
                errors += luat_rieng.validate(base)
                conflicts = precheck(base, student_rules)
                if conflicts:
                    errors += [f"Quy định mâu thuẫn, không có TKB nào thỏa: {line}" for line in conflicts]
                    return result()
                if quick:
                    return result()
                plan = phan_cong(base, config.Weights())
            except (InputError, SolveError) as exc:
                errors += _lines(exc)
                return result(luat_rieng.RULES_SHEET)
            info += [line.strip() for line in du_toan_lines(base, plan)]
            if plan.missing:
                rows = ShortageError(base, plan).rows()
                head = (f"Bù tối đa +{overtime_max} tiết/người vẫn thiếu {plan.missing_total()} tiết: chế độ bù giờ sẽ "
                        f"báo lỗi; tăng số tiết bù tối đa, sửa nhân sự, hoặc chọn chế độ tuyển thêm"
                        if mode == config.MODE_OVERTIME else
                        f"Thiếu {plan.missing_total()} tiết dù đã bù tối đa +{overtime_max}: chế độ tuyển thêm sẽ thêm "
                        f"người \"{config.SUPPLEMENT_NAME}\"")
                info.append(head)
                info += [f"Lớp {cls}: {subject} thiếu {n} tiết ({reason})" for cls, subject, n, reason in rows[:30]]
                if len(rows) > 30:
                    info.append(f"... và {len(rows) - 30} dòng khác")
    return result()


# ---- TKB đã xếp trên giao diện: xem, đổi ô, kiểm luật tức thì ----

def grid_key(scenario: dict, mode: str, overtime_max: int, student_rules: bool) -> str:
    """Khóa của một Grid: mọi thứ trừ TKB đã xếp (đổi ô trên TKB thì dùng lại bài toán đã dựng)."""
    text = json.dumps([{k: v for k, v in scenario.items() if k != "saved"}, mode, overtime_max, student_rules],
                      ensure_ascii=False, sort_keys=True, default=str)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


class Grid:
    """Kịch bản đã đọc để xem và đổi ô của TKB đã xếp (scenario["saved"]) trên giao diện: đọc file vào và dựng bài
    toán một lần (như solver.reuse: build_problem, không xếp), mỗi lần xem hay thử đổi ô chỉ dựng lại các tiết và
    chạy checker (vài chục ms). Giao diện giữ lại theo grid_key. Đổi tạm config (rules.applied) nên không gọi song
    song."""

    def __init__(self, scenario: dict, mode: str = config.MODE_OVERTIME, overtime_max: int = config.OVERTIME_MAX,
                 student_rules: bool = True):
        from .allocation import build_problem
        from .rules import code

        self.student_rules = student_rules
        self.errors: list[str] = []  # lỗi đọc kịch bản: chưa kiểm được TKB
        self.values: dict = {}
        self.staff: list = []
        self.problem = None
        self.rules_code = None
        self.choices: dict[str, dict] = {}
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "kich_ban.xlsx"
            to_excel({**scenario, "saved": None}, path)
            try:
                self.values = read_rules(path)
            except InputError as exc:
                self.errors = _lines(exc)
                return
            with applied(self.values):
                try:
                    curriculum = read_program(path)
                    self.staff = read_staff(path, subjects=[s for req in curriculum.values() for s in req])
                    bo_ghep.know_staff(self.staff)
                    over = overtime_max if mode == config.MODE_OVERTIME else 0
                    self.problem = build_problem(self.staff, curriculum, {}, overtime_max=over)
                    # Người có thể dạy từng môn của lớp khi bỏ các luật "Chỉ giáo viên dạy" (đổi người dạy ở bước 7
                    # là ghi / sửa đúng luật đó).
                    with applied({"CUSTOM_RULES": [r for r in config.CUSTOM_RULES if r.kind != "chi_gv"]}):
                        self.choices = _choices(build_problem(self.staff, curriculum, {}, overtime_max=over))
                except InputError as exc:
                    self.errors = _lines(exc)
                    return
                self.rules_code = code()

    def _lessons(self, saved) -> tuple[list, list[str]]:
        from .checker import check
        from .solver import saved_lessons

        lessons, errors = saved_lessons(self.problem, saved.rows)
        if not errors:
            errors = check(self.problem, lessons, self.student_rules)
        return lessons, errors

    def view(self, grid: list[list] | None) -> dict:
        """TKB đã xếp để vẽ trên trang: {days, sessions, classes, teachers, cells, errors, ok, code, rules_changed}.
        cells: mọi ô của lưới, cả ô trống ({cls, d, p, r, c, subject, code, name, overtime, locked}; r, c là dòng,
        cột trong scenario["saved"] đếm từ 1); errors: [{text, cells: [[lớp, ngày, tiết]], cls, teacher}], câu lỗi
        ghi Mã GV."""
        from .staff import parse_saved_grid

        saved = parse_saved_grid(grid or [])
        with applied(self.values):
            bo_ghep.know_staff(self.staff)
            days = khung_gio.days()  # các ngày học theo thứ tự (như sheet TKB đã xếp)
            periods = {d: set(khung_gio.periods(d)) for d in days}
            # Hàng của lưới: tiết 1 … tiết của ngày dài nhất (như sheet TKB đã xếp); nhãn buổi của mỗi tiết theo ngày
            # đầu tiên có tiết đó; ngày không có tiết đó: ô nghỉ.
            sessions: list[dict] = []
            for p in range(1, khung_gio.max_periods() + 1):
                name = khung_gio.session_at(next(d for d in days if p in periods[d]), p).name
                if sessions and sessions[-1]["name"] == name:
                    sessions[-1]["periods"].append(p)
                else:
                    sessions.append({"name": name, "periods": [p]})
            out = {"days": [config.DAYS[d] for d in days], "sessions": sessions,
                   "off": [[d, p] for d in days for p in range(1, khung_gio.max_periods() + 1) if p not in periods[d]],
                   "code": saved.result_code, "cells": _cells(saved, self.staff), "teachers": _teachers(self.staff),
                   "choices": self.choices}
            out["classes"] = sorted({c["cls"] for c in out["cells"]}, key=class_sort_key)
            out["grades"] = {c: _grade_name(c) for c in out["classes"]}  # khối của từng lớp (sheet LỚP)
            out["rules_changed"] = saved.rules_code not in (None, self.rules_code) and self.problem is not None
            if self.problem is None:
                return {**out, "ok": False, "input_errors": self.errors, "errors": []}
            from .phong_hoc import labels, placed
            names = labels(self.problem)
            out["rooms"] = names  # sheet PHÒNG: xem TKB theo phòng
            if not saved.rows:
                return {**out, "ok": False, "input_errors": [], "errors": []}
            lessons, errors = self._lessons(saved)
            if self.problem.room_fit and lessons:  # tiết học ở phòng dùng chung: ô ghi tên phòng
                at = {(normalize(les.class_name), les.day, les.period): names[r]
                      for les, r in zip(lessons, placed(self.problem, lessons)) if r is not None}
                for cell in out["cells"]:
                    cell["room"] = at.get((normalize(cell["cls"]), cell["d"], cell["p"]), "")
            return {**out, "ok": not errors, "input_errors": [],
                    "errors": [_marked(text, self.problem, out["cells"]) for text in errors]}

    def swaps(self, grid: list[list] | None, cls: str, day: int, period: int) -> list[dict]:
        """Thử đổi ô (cls, day, period) với từng ô khác của lớp (trừ ô cùng môn, cùng người dạy): [{d, p, new,
        fixed}], new là số lỗi mới (0: đổi được), fixed là số lỗi cũ hết đi. Lưới không đọc được (lỗi đọc, sai số
        tiết) thì không thử."""
        from dataclasses import replace

        from .checker import check
        from .staff import parse_saved_grid

        saved = parse_saved_grid(grid or [])
        if self.problem is None or not saved.rows:
            return []
        with applied(self.values):
            bo_ghep.know_staff(self.staff)
            lessons, base = self._lessons(saved)
            if not lessons:
                return []
            base = set(base)
            slots = sorted({(d, p) for d, sessions in config.DAY_SESSIONS.items() for s in sessions
                            for p in s.periods})
            mine = {(les.day, les.period): i for i, les in enumerate(lessons) if les.class_name == cls}
            out = []
            here = mine.get((day, period))
            same = (lambda i: i is not None and here is not None and
                    (lessons[i].subject, lessons[i].teacher, lessons[i].overtime)
                    == (lessons[here].subject, lessons[here].teacher, lessons[here].overtime))
            for slot in slots:
                if slot == (day, period) or (slot not in mine and here is None) or same(mine.get(slot)):
                    continue  # đổi với chính nó, hai ô trống, hai tiết như nhau: không đổi gì
                trial = list(lessons)
                for i, to in ((mine.get((day, period)), slot), (mine.get(slot), (day, period))):
                    if i is not None:
                        trial[i] = replace(trial[i], day=to[0], period=to[1])
                errors = set(check(self.problem, trial, self.student_rules))
                out.append({"d": slot[0], "p": slot[1], "new": len(errors - base), "fixed": len(base - errors)})
            return out


def timetable(scenario: dict, mode: str = config.MODE_OVERTIME, overtime_max: int = config.OVERTIME_MAX,
              student_rules: bool = True) -> dict:
    """Grid(...).view của TKB đã xếp trong kịch bản (một lần, không giữ lại)."""
    return Grid(scenario, mode, overtime_max, student_rules).view(scenario.get("saved"))


def _choices(problem) -> dict[str, dict]:
    """"lớp|môn như ghi trong TKB" -> {subject: tên môn để ghi luật, codes: Mã GV các người có thể dạy}, cho các môn
    không có phần GVCN dạy (phần đó cố định theo quy định GVCN) và có ít nhất hai người (kể cả người cần tuyển ghi ở
    nhân sự, trừ người tuyển dự kiến của chương trình)."""
    homeroom = {(c.class_name, c.subject) for c in problem.courses if c.homeroom}
    found: dict[tuple[str, str], list[str]] = {}
    for c in problem.courses:
        if (c.class_name, c.subject) in homeroom:
            continue
        names = found.setdefault((c.class_name, c.subject), [])
        names += [t for t in c.teachers if t not in names and not problem.teachers[t].supplementary]
    return {f"{cls}|{problem.subject_label(subject)}": {"subject": subject,
                                                        "codes": [problem.teachers[t].code for t in titles]}
            for (cls, subject), titles in found.items() if len(titles) > 1}


def _cells(saved, staff) -> list[dict]:
    name_of = {normalize(t.code): t.name for t in staff}
    day_of = {normalize(d): i for i, d in enumerate(config.DAYS)}
    out = []
    for cls, day, period, r, c, i in saved.places:
        d = day_of.get(normalize(day))
        try:
            p = int(period)
        except (TypeError, ValueError):
            continue
        if d is None:
            continue
        cell = {"cls": cls, "d": d, "p": p, "r": r, "c": c, "subject": None, "code": None, "name": None,
                "overtime": False, "locked": False}
        if i is not None:
            _, _, _, subject, code, overtime, locked, _ = saved.rows[i]
            cell.update(subject=subject, code=code, name=name_of.get(normalize(code)), overtime=overtime,
                        locked=locked)
        out.append(cell)
    return out


def _teachers(staff) -> list[dict]:
    return [{"code": t.code, "name": t.name, "cls": t.class_name} for t in staff]


def _marked(text: str, problem, cells: list[dict]) -> dict:
    """Một lỗi của checker kèm các ô nó nói tới (tô viền đỏ trên trang), tìm theo chữ: lớp, Mã GV, Thứ, tiết, buổi,
    môn trong câu. Câu của checker ghi chức vụ chuẩn hóa (Teacher.title): đổi ra Mã GV."""
    import re

    def find(names, within: str):
        found = []
        for name in sorted(names, key=len, reverse=True):
            pattern = rf"(?<![\w/]){re.escape(name)}(?![\w/])"
            if re.search(pattern, within):
                found.append(name)
                within = re.sub(pattern, " ", within)
        return found, within

    teachers = sorted(problem.teachers.values(), key=lambda t: len(t.title), reverse=True)
    for t in teachers:
        text = re.sub(rf"(?<![\w/]){re.escape(t.title)}(?![\w/])", t.code, text)
    body = text
    for r in config.CUSTOM_RULES:  # "LUẬT dòng n: <câu của luật>: <chỗ sai>": chỉ đọc chỗ sai
        head = luat_rieng.label(r) + ": "
        if body.startswith(head):
            body = body[len(head):]
            break
    codes, rest = find({t.code for t in problem.teachers.values()}, body)
    classes, rest = find({c["cls"] for c in cells}, rest)
    labels = {problem.subject_label(s): s for s in {c.subject for c in problem.courses}}
    subjects = set()
    for name in sorted(labels, key=len, reverse=True):
        pattern = rf"(?<!\w){re.escape(name)}(?!\w)"
        if re.search(pattern, rest):
            subjects.add(name)
            rest = re.sub(pattern, " ", rest)
    subjects |= {s for s in labels if config.SUBJECT_GROUPS.get(labels[s]) in {labels.get(x) for x in subjects}}
    days = [i for i, d in enumerate(config.DAYS) if re.search(rf"(?<!\w){re.escape(d)}(?!\d)", rest, re.IGNORECASE)]
    periods: set[int] = set()
    for m in re.finditer(r"tiết (\d+(?:\s*(?:,|–|-|và)\s*\d+)*)", rest):
        numbers = [int(x) for x in re.findall(r"\d+", m.group(1))]
        periods |= set(range(numbers[0], numbers[-1] + 1)) if "–" in m.group(1) else set(numbers)
    if not periods:
        for s in {s for sessions in config.DAY_SESSIONS.values() for s in sessions}:
            if re.search(rf"buổi {re.escape(s.name)}(?!\w)", rest, re.IGNORECASE):
                periods |= set(s.periods)
    marks = []
    if days:
        chosen = [c for c in cells if c["subject"] and c["d"] in days and (not periods or c["p"] in periods)
                  and (c["cls"] in classes if classes else c["code"] in codes if codes else False)]
        narrowed = [c for c in chosen if c["subject"] in subjects]
        marks = [[c["cls"], c["d"], c["p"]] for c in (narrowed or chosen)]
    return {"text": text, "cells": marks, "cls": classes[0] if classes else None,
            "teacher": codes[0] if codes else None}
