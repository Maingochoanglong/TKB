from pathlib import Path

import openpyxl

from tkb import config
from tkb.program import read_program
from tkb.solver import solve
from tkb.staff import read_staff
from tkb.rules import DEFAULTS, read_rules
from tkb.template import main as template_main, write_staff_template
from tkb.writer import write_updated_staff

from .conftest import CURRICULUM, small_staff

ROOT = Path(__file__).resolve().parent.parent
HEADER = ("Họ và Tên", "Chức Vụ", "Lớp", "Số Tiết/Tuần", "Thai Sản", "Hợp Đồng", "Cơ sở 2", "Lớp Đang Dạy", "Buổi Nghỉ")
NONE5 = (None,) * 5  # 5 cột không bắt buộc để trống


def _key(teachers):
    return [(t.name, t.title, t.max_lessons) for t in teachers]


def _rows(ws):
    return [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]


def test_template_is_plain(tmp_path, sample_staff):
    path = tmp_path / "mau.xlsx"
    write_staff_template(path, sample_staff, CURRICULUM)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["NHÂN SỰ", "CHƯƠNG TRÌNH HỌC", "QUY ĐỊNH", "HƯỚNG DẪN"]
    ws = wb["NHÂN SỰ"]
    # Style như file của nhà trường: Times New Roman 14 chữ đen, tiêu đề đậm, không tô nền, viền mảnh, dòng cao 25.
    assert ws["A1"].font.name == "Times New Roman" and ws["A1"].font.sz == 14 and ws["A1"].font.b
    assert ws["B2"].font.sz == 14 and not ws["B2"].font.b and ws["B2"].border.left.style == "thin"
    assert ws["B2"].alignment.horizontal == "center" and ws.row_dimensions[2].height == 25
    assert ws.max_column == 9
    # Đơn giản: không tô nền, chữ đen, không cố định dòng/cột, không danh sách thả xuống, không định dạng theo điều
    # kiện, không ghi chú trong ô, không sheet ẩn.
    for sheet in wb.worksheets:
        assert sheet.sheet_state == "visible" and sheet.freeze_panes is None
        assert not sheet.data_validations.dataValidation and not len(sheet.conditional_formatting)
        for row in sheet.iter_rows():
            for c in row:
                assert c.comment is None and c.fill.fill_type is None
                assert c.value is None or c.font.color is None or c.font.color.rgb == "FF000000"
    assert ws["C2"].number_format == ws["H2"].number_format == "@"  # chữ để Excel không đổi thành ngày tháng
    program = _rows(wb["CHƯƠNG TRÌNH HỌC"])
    assert program[0][:7] == ("Môn học", "Khối 1", "Khối 2", "Khối 3", "Khối 4", "Khối 5", "Tên trong TKB")
    assert len(program[0]) == 6 + 15 and len(program) == 17
    rules = _rows(wb["QUY ĐỊNH"])  # ba bảng xếp chồng (bỏ dòng trống)
    assert rules[0][:2] == ("Quy định", "Giá trị") and rules[7][0] == "Ngày" and rules[14][0] == "Tiết"
    assert rules[8][:5] == ("Thứ 2", "Có", "Có", 1, "Không") and len(rules) == 7 + 7 + 8
    # Mỗi ô quy định chỉ là Có, Không hoặc số nguyên dương (trừ cột đầu, cột Khối và cột Tên trong TKB).
    blocks = [program, rules[:7], rules[7:14], rules[14:]]
    for header, *rows in blocks:
        for row in rows:
            for head, value in list(zip(header, row))[1:]:
                if head and head != "Tên trong TKB" and not head.startswith("Khối"):
                    assert value in ("Có", "Không", None) or (isinstance(value, int) and value > 0), (head, value)
    assert read_rules(path) == DEFAULTS  # các quy định điền sẵn đúng giá trị mặc định
    assert _rows(wb["HƯỚNG DẪN"])[0] == ("Mục", "Cách ghi")
    assert _key(read_staff(path)) == _key(sample_staff)
    assert read_program(path) == CURRICULUM


def _values(path):
    return {ws.title: _rows(ws) for ws in openpyxl.load_workbook(path).worksheets}


def test_saved_input_template_is_up_to_date(tmp_path):
    """data/Input_Template_V8.xlsx là file mẫu trống (python -m tkb.template data/Input_Template_V8.xlsx)."""
    write_staff_template(tmp_path / "moi.xlsx")
    assert _values(ROOT / "data" / "Input_Template_V8.xlsx") == _values(tmp_path / "moi.xlsx"), \
        "data/Input_Template_V8.xlsx đã cũ: chạy  python -m tkb.template data/Input_Template_V8.xlsx"


def test_blank_template(tmp_path):
    """File mẫu trống: nhân sự chỉ có tiêu đề; chương trình học điền sẵn các môn có quy định, số tiết để trống."""
    path = tmp_path / "trong.xlsx"
    write_staff_template(path)
    wb = openpyxl.load_workbook(path)
    assert _rows(wb["NHÂN SỰ"]) == [HEADER]
    program = _rows(wb["CHƯƠNG TRÌNH HỌC"])
    assert [r[0] for r in program[1:4]] == ["Tiếng Việt", "Toán", "Hoạt động trải nghiệm"]
    assert all(r[1:6] == (None,) * 5 for r in program[1:])
    assert read_rules(path) == DEFAULTS


def test_updated_staff_keeps_template_and_style(tmp_path):
    src = tmp_path / "ns.xlsx"
    write_staff_template(src, small_staff(general=False), CURRICULUM)
    sol = solve(read_staff(src), CURRICULUM, config.Settings(time_limit=20, workers=4), log=lambda *_: None)
    dst = tmp_path / "ns_cap_nhat.xlsx"
    _old_template_extras(src)
    write_updated_staff(sol, src, dst)
    wb = openpyxl.load_workbook(dst)
    ws = wb["NHÂN SỰ"]
    # File vào theo mẫu cũ có danh sách thả xuống, định dạng theo điều kiện, sheet danh mục ẩn, ghi chú ở tiêu đề
    # cột và cố định dòng; file cập nhật (file kết quả) chỉ còn chữ, số và màu.
    assert not ws.data_validations.dataValidation and not len(ws.conditional_formatting) and ws.freeze_panes is None
    assert wb.sheetnames == ["NHÂN SỰ", "CHƯƠNG TRÌNH HỌC", "QUY ĐỊNH", "HƯỚNG DẪN", "TKB đã xếp"]
    assert not any(c.comment for sheet in wb.worksheets for row in sheet.iter_rows() for c in row)
    rows = _rows(ws)
    assert rows[0] == HEADER + ("Mã GV", "Số Tiết Thực Dạy", "Số Tiết Dư")
    assert rows[1] == ("CN A", "Chủ Nhiệm", "3/1", 19, *NONE5, "Chủ Nhiệm 3/1", 19, None)
    assert rows[-1] == ("chưa có", "Bộ Môn", None, 23, *NONE5, "Bộ Môn 1", 8, 15)
    # Dòng mới và cột mới cùng style với file vào.
    last = len(rows)
    for cell in (ws.cell(last, 2), ws.cell(last, 10), ws.cell(1, 11)):
        assert cell.font.name == "Times New Roman" and cell.font.sz == 14 and cell.border.left.style == "thin"
    assert ws.cell(1, 11).font.b and ws.row_dimensions[last].height == 25
    assert _rows(wb["CHƯƠNG TRÌNH HỌC"])[1:] == _rows(openpyxl.load_workbook(src)["CHƯƠNG TRÌNH HỌC"])[1:]
    again = read_staff(dst)
    assert again[-1].title == "bộ môn 1"  # tự đánh số khi đọc lại
    # Chạy lại trên file cập nhật: không tuyển thêm nữa, các cột kết quả được ghi đè chứ không thêm mới.
    sol2 = solve(again, CURRICULUM, config.Settings(time_limit=20, workers=4), log=lambda *_: None)
    assert sol2.used_supplements() == []
    write_updated_staff(sol2, dst, tmp_path / "lan2.xlsx")
    rows2 = _rows(openpyxl.load_workbook(tmp_path / "lan2.xlsx")["NHÂN SỰ"])
    assert rows2[0] == rows[0] and len(rows2) == len(rows)


def _old_template_extras(path):
    """Thêm vào file vào những thứ của file mẫu cũ: danh sách thả xuống, tô đỏ theo điều kiện, ghi chú ở tiêu đề,
    cố định dòng, sheet danh mục ẩn."""
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = openpyxl.load_workbook(path)
    ws = wb["NHÂN SỰ"]
    ws.add_data_validation(DataValidation(type="list", formula1="'Danh mục'!$A$1:$A$3", sqref="B2:B300"))
    ws.conditional_formatting.add("C2:C300", FormulaRule(formula=['$C2=""'], fill=PatternFill(bgColor="FFC7CE")))
    ws["B1"].comment = Comment("Chủ Nhiệm, Bộ Môn...", "TKB")
    ws.freeze_panes = "A2"
    lists = wb.create_sheet("Danh mục")
    lists["A1"] = "Chủ Nhiệm"
    lists.sheet_state = "hidden"
    wb.save(path)


def test_cli_writes_blank_template(tmp_path):
    out = tmp_path / "moi.xlsx"
    assert template_main([str(out)]) == 0
    assert _rows(openpyxl.load_workbook(out)["NHÂN SỰ"]) == [HEADER]


def test_optional_columns_round_trip(tmp_path):
    """Các cột Thai Sản, Hợp Đồng, Cơ sở 2, Lớp Đang Dạy, Buổi Nghỉ ghi ra file mẫu rồi đọc lại đúng."""
    from dataclasses import replace
    staff = small_staff()
    staff = [replace(t, campus2=True, maternity=True, off_any=(("Chiều", 2),)) if t.class_name == "3/1" else
             replace(t, contract=True, history=frozenset({"3/1"}), off_sessions=frozenset({(3, "Chiều"), (4, "Sáng")}))
             if not t.class_name else t for t in staff]
    path = tmp_path / "mau.xlsx"
    write_staff_template(path, staff, CURRICULUM)
    ws = openpyxl.load_workbook(path)["NHÂN SỰ"]
    assert [c.value for c in ws[4]][4:] == [None, "Có", None, "3/1", "Chiều T5, Sáng T6"]  # dòng TA
    fields = lambda ts: [(t.maternity, t.contract, t.campus2, t.history, t.off_sessions, t.off_any) for t in ts]
    assert fields(read_staff(path)) == fields(staff)


def test_updated_staff_turns_formulas_into_values(tmp_path):
    """Cột STT ghi công thức =ROW()-1: file cập nhật ghi số, người cần tuyển nhận số tiếp theo."""
    src = tmp_path / "stt.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NHÂN SỰ"
    ws.append(["STT", *HEADER[:4]])
    for t in small_staff(general=False):
        ws.append(["=ROW()-1", t.name, t.role.title(), t.class_name, t.max_lessons])
    prog = wb.create_sheet("CHƯƠNG TRÌNH HỌC")
    prog.append(["Môn học", "Khối 3"])
    for subject, n in CURRICULUM[3].items():
        prog.append([subject, n])
    wb.save(src)
    sol = solve(read_staff(src), CURRICULUM, config.Settings(time_limit=20, workers=4), log=lambda *_: None)
    dst = tmp_path / "stt_cap_nhat.xlsx"
    write_updated_staff(sol, src, dst)
    out = openpyxl.load_workbook(dst)["NHÂN SỰ"]
    stt = [out.cell(r, 1).value for r in range(2, out.max_row + 1) if out.cell(r, 2).value]
    assert stt == list(range(1, len(stt) + 1)) and len(stt) == len(small_staff(general=False)) + 1
    assert not any(c.data_type == "f" for row in out.iter_rows() for c in row)
    # File vào không ghi quy định: file cập nhật ghi các quy định đã dùng (cột quy định của môn thêm vào sheet chương
    # trình học, sheet QUY ĐỊNH sau sheet đó) và sheet HƯỚNG DẪN.
    wb = openpyxl.load_workbook(dst)
    assert wb.sheetnames == ["NHÂN SỰ", "CHƯƠNG TRÌNH HỌC", "QUY ĐỊNH", "HƯỚNG DẪN", "TKB đã xếp"]
    program = _rows(wb["CHƯƠNG TRÌNH HỌC"])
    assert program[0][:3] == ("Môn học", "Khối 3", "Tên trong TKB") and program[1][:2] == ("Tiếng Việt", 7)
    assert read_rules(dst) == DEFAULTS
