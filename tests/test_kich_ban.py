"""Kịch bản của giao diện (tkb/kich_ban.py): file Excel -> kịch bản -> file Excel không mất gì, và kiểm tra báo lỗi
như khi chạy."""
import datetime
import json

import openpyxl

from tkb import config, kich_ban, luat_co_san, luat_rieng
from tkb.program import read_program
from tkb.rules import DEFAULTS, LEGACY, SUBJECT_COLS, applied, code, read_rules
from tkb.staff import read_saved_timetable, read_staff
from tkb.template import write_staff_template

from .conftest import CURRICULUM, INPUT_FILE, small_staff, plain_rules


DEFAULT_CODE = code()


def _values(path):
    return {ws.title: [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]
            for ws in openpyxl.load_workbook(path).worksheets}


def _key(teachers):
    return [(t.name, t.title, t.max_lessons, t.class_name) for t in teachers]


def test_round_trip_keeps_the_file(tmp_path):
    scenario, warnings = kich_ban.from_excel(INPUT_FILE)
    assert warnings == []
    json.dumps(scenario)  # gửi được cho trang web
    assert len(scenario["staff"]) == 45 and scenario["grades"] == [1, 2, 3, 4, 5]
    assert scenario["staff"][0] == {"name": "Giáo viên CN 1", "role": "Chủ Nhiệm", "class": "1/1", "lessons": 19,
                                    "maternity": False, "contract": False, "campus2": False, "history": "", "off": ""}
    # Chức vụ trùng tên môn nhân sự đang dùng thành các dòng của sheet CHỨC VỤ (để giao diện chọn từ danh sách).
    assert [(r["name"], r["subjects"]) for r in scenario["roles"]] == [
        ("Tiếng Anh", ["Tiếng Anh"]), ("Thể Dục", ["Thể dục"]), ("Âm Nhạc", ["Âm nhạc"]), ("Mỹ Thuật", ["Mỹ thuật"]),
        ("Tin Học", ["Tin học"])]
    out = tmp_path / "ra.xlsx"
    kich_ban.to_excel(scenario, out)
    written, source = _values(out), _values(INPUT_FILE)
    assert written.pop("CHỨC VỤ")[1:] == [(r["name"], r["subjects"][0]) for r in scenario["roles"]]
    assert source.pop("CHỨC VỤ") == [("Chức vụ", "Môn được dạy")]
    assert written == source  # các sheet khác ghi lại đúng từng ô
    rules = read_rules(out)
    assert read_program(out) == CURRICULUM and {**plain_rules(rules), "CUSTOM_ROLES": []} == DEFAULTS
    with applied(rules):  # các dòng đó như không ghi: mã quy định không đổi
        assert code() == DEFAULT_CODE
    assert _key(read_staff(out)) == _key(read_staff(INPUT_FILE))


def test_new_scenario_is_the_blank_template(tmp_path):
    write_staff_template(tmp_path / "mau.xlsx")
    kich_ban.to_excel(kich_ban.default_scenario(), tmp_path / "moi.xlsx")
    assert _values(tmp_path / "moi.xlsx") == _values(tmp_path / "mau.xlsx")


def _row(scenario, native):
    """Dòng luật của kịch bản là dạng gốc của luật có sẵn `native`."""
    keys = [r["native"] for r in kich_ban.describe(scenario)["rules"]]
    return scenario["rules"][keys.index(native)]


def test_schema_follows_rules_columns():
    s = kich_ban.schema()
    assert [c["key"] for c in s["subject"]] == [c.key for c in SUBJECT_COLS if c.key not in LEGACY]
    assert [c["header"] for c in s["staff"]][:4] == ["Họ và Tên", "Chức Vụ", "Lớp", "Số Tiết/Tuần"]
    assert s["days"][0] == "Thứ 2" and s["days"][-1] == "Thứ 7"
    keys = {c.key for c in SUBJECT_COLS}  # nhóm cột của trang chi tiết môn, cột theo chức vụ: đúng khóa của rules.py
    assert all(k in keys for g in s["subject_groups"] for k in g["keys"])
    assert set(s["role_rules"].values()) <= keys and s["sheets"]["roles"] == "CHỨC VỤ"
    # Loại luật theo bốn họ: mọi mẫu luật (trừ Tự ghép) và mọi phép đo đều thuộc một họ; mức ưu tiên bằng chữ.
    families = s["custom"]["composer"]["families"]
    assert all(k["family"] in families for k in s["custom"]["kinds"] if k["key"] != "tu_ghep")
    assert all(m["family"] in families for m in s["custom"]["composer"]["measures"])
    assert s["custom"]["levels"] == ["Thấp", "Vừa", "Cao", "Rất cao"]


def test_rules_edited_in_the_scenario_reach_the_file(tmp_path):
    scenario, _ = kich_ban.from_excel(INPUT_FILE)
    _row(scenario, "nhom_buoi")["number"] = 3  # số của luật có sẵn: ở dòng luật
    scenario["days"][5]["morning_days"] = True  # học sáng Thứ 7
    english = next(s for s in scenario["subjects"] if s["name"] == "Tiếng Anh")
    english["rules"]["MORNING_SUBJECTS"] = True
    english["rules"]["DISPLAY_NAMES"] = "TA"
    out = tmp_path / "ra.xlsx"
    kich_ban.to_excel(scenario, out)
    rules = read_rules(out)
    assert rules["SESSION_GROUP_LIMIT"] == 3
    assert rules["DAYS"][-1] == "Thứ 7" and len(rules["DAY_SESSIONS"]) == 6
    assert config.TIENG_ANH in rules["MORNING_SUBJECTS"] and rules["DISPLAY_NAMES"][config.TIENG_ANH] == "TA"


def test_check_reads_back_like_a_run():
    scenario, _ = kich_ban.from_excel(INPUT_FILE)
    ok = kich_ban.check(scenario, config.MODE_OVERTIME, 2)
    assert ok["errors"] == [] and any(line.startswith("Dự toán:") for line in ok["info"])
    scenario["staff"][20]["class"] = "1/1"  # hai Chủ Nhiệm lớp 1/1
    scenario["subjects"][0]["lessons"]["1"] = "mười"
    res = kich_ban.check(scenario, config.MODE_OVERTIME, 2)
    assert "NHÂN SỰ: Lớp 1/1 có hai Chủ Nhiệm (dòng 2 và 22)" in res["errors"]
    assert "CHƯƠNG TRÌNH HỌC: Dòng 2, Khối 1: số tiết không hợp lệ 'mười'" in res["errors"]


def test_check_estimates_shortage(tmp_path):
    """Trường nhỏ không có bộ môn: bù +1 vẫn thiếu tiết (như khi chạy, chế độ bù giờ sẽ báo lỗi mã 3); không ai
    được bù thì có môn không ai dạy được (chạy cũng báo lỗi đó)."""
    write_staff_template(tmp_path / "vao.xlsx", small_staff(general=False), CURRICULUM)
    scenario, _ = kich_ban.from_excel(tmp_path / "vao.xlsx")
    res = kich_ban.check(scenario, config.MODE_OVERTIME, 1)
    assert res["errors"] == []
    assert any("chế độ bù giờ sẽ báo lỗi" in line for line in res["info"])
    assert any(line.startswith("Lớp 3/1: TNXH thiếu 2 tiết") for line in res["info"])
    res = kich_ban.check(scenario, config.MODE_HIRE, 1)
    assert any("chế độ tuyển thêm sẽ thêm người" in line for line in res["info"])
    assert kich_ban.check(scenario, config.MODE_OVERTIME, 0)["errors"] == [
        "Lớp 3/1: không có GV nào được phép dạy môn Công nghệ"]


def test_blank_staff_rows_keep_row_numbers(tmp_path):
    scenario, _ = kich_ban.from_excel(INPUT_FILE)
    scenario["staff"].insert(0, {"name": "", "role": "", "class": "", "lessons": None})
    kich_ban.to_excel(scenario, tmp_path / "ra.xlsx")
    assert [t.row for t in read_staff(tmp_path / "ra.xlsx")][:2] == [3, 4]  # dòng i của bảng là dòng i + 2


def test_excel_date_in_class_column_is_read_back_with_a_warning(tmp_path):
    path = tmp_path / "vao.xlsx"
    write_staff_template(path, small_staff(), CURRICULUM)
    wb = openpyxl.load_workbook(path)
    wb["NHÂN SỰ"]["C2"] = datetime.datetime(2025, 1, 3)  # Excel đổi "3/1" thành ngày 3 tháng 1
    wb.save(path)
    scenario, warnings = kich_ban.from_excel(path)
    assert scenario["staff"][0]["class"] == "3/1" and "ngày tháng" in warnings[0]


def test_saved_timetable_sheet_is_kept(tmp_path):
    path = tmp_path / "vao_cap_nhat.xlsx"
    write_staff_template(path, small_staff(), CURRICULUM)
    wb = openpyxl.load_workbook(path)
    ws = wb.create_sheet(config.SAVED_SHEET)
    ws.append(["Mã kết quả", "AAAA-BBBB-CCCC", "Mã quy định", "123456789ABC"])
    ws.append([])
    ws.append(["Lớp", "Tiết", "Thứ 2"])
    ws.append(["3/1", 1, "HĐTN\nChủ Nhiệm 3/1"])
    ws.append([None, 2, "Tiếng Việt\nBộ Môn 1 (bù)"])
    wb.save(path)
    scenario, _ = kich_ban.from_excel(path)
    assert scenario["saved"][0] == ["Mã kết quả", "AAAA-BBBB-CCCC", "Mã quy định", "123456789ABC"]
    kich_ban.to_excel(scenario, tmp_path / "ra.xlsx")
    assert read_saved_timetable(tmp_path / "ra.xlsx") == read_saved_timetable(path)


def test_check_finds_rules_in_conflict():
    scenario, _ = kich_ban.from_excel(INPUT_FILE)
    _row(scenario, "nhom_buoi")["number"] = 1
    res = kich_ban.check(scenario, config.MODE_OVERTIME, 2)
    assert res["errors"][0].startswith("Quy định mâu thuẫn, không có TKB nào thỏa: Khối 1: Tiếng Việt có 14 tiết/tuần")
    assert res["info"] == [] or not any(line.startswith("Dự toán:") for line in res["info"])
    assert kich_ban.check(scenario, config.MODE_OVERTIME, 2, student_rules=False)["errors"] == []


def test_rules_round_trip(tmp_path):
    """Mọi luật ở `rules` của kịch bản (sheet LUẬT): luật có sẵn ở dạng gốc và luật thêm vào; ghi ra Excel rồi đọc
    lại như cũ."""
    from tkb.config import CustomRule

    scenario, _ = kich_ban.from_excel(INPUT_FILE)
    n = len(luat_co_san.default_rows())
    assert len(scenario["rules"]) == n and all(r["group_label"] for r in scenario["rules"])
    assert all(r["native"] for r in kich_ban.describe(scenario)["rules"])
    assert len(kich_ban.schema()["custom"]["kinds"]) == len(luat_rieng.KINDS)
    scenario["rules"] += [{"kind": "Chỉ xếp vào", "subject": "Thể dục", "sessions": "Chiều", "hard": False, "level": "Cao"},
                          {"kind": "Học trước", "subject": "Tiếng Việt", "other": "Toán", "hard": True},
                          {"kind": "Tự ghép", "scope": "Giáo viên, Ngày", "measure": "Số khác nhau", "op": "Tối đa",
                           "number": 2, "count_by": "Lớp", "role": "Tiếng Anh", "hard": True}]
    out = tmp_path / "ra.xlsx"
    kich_ban.to_excel(scenario, out)
    assert read_rules(out)["CUSTOM_RULES"] == [
        CustomRule("chi_xep", "Thể dục", sessions=("Chiều",), level=3, row=n + 2),
        CustomRule("truoc", "Tiếng Việt", other="Toán", hard=True, row=n + 3),
        CustomRule("tu_ghep", role="tiếng anh", number=2, hard=True, row=n + 4, scope=("gv", "ngay"),
                   measure="so_khac", op="<=", count_by="lop")]
    back, _ = kich_ban.from_excel(out)
    assert back["rules"][:n] == scenario["rules"][:n]
    assert back["rules"][n]["level"] == "Cao"  # cột Mức ghi bằng chữ
    assert back["rules"][n + 1] == {"group_label": "", "kind": "Học trước", "subject": "Tiếng Việt", "other": "Toán",
                                    "grades": "", "days": "", "periods": "", "sessions": "", "role": "",
                                    "number": None, "hard": True, "level": None, "scope": "", "group": "", "tags": "",
                                    "classes": "", "measure": "", "op": "", "count_by": "", "when": "",
                                    "exclude": "", "points": None}
    assert kich_ban.check(back, config.MODE_OVERTIME, 2)["errors"] == []
    back["rules"].append({"kind": "Học trước", "subject": "Toán"})
    assert f"LUẬT, dòng {n + 5}: kiểu luật Học trước phải ghi cột Môn thứ hai" in \
        kich_ban.check(back, config.MODE_OVERTIME, 2)["errors"]


def test_teacher_rules_read_with_codes():
    """Luật ghi họ tên ở cột Giáo viên: câu đọc lại ghi Mã GV, không ghi họ tên; tên không có thì Kiểm tra báo lỗi."""
    scenario, _ = kich_ban.from_excel(INPUT_FILE)
    name = next(t["name"] for t in scenario["staff"] if t["role"] == "Tiếng Anh")
    busy = {"kind": "Không xếp vào", "role": name, "days": "Thứ 2", "periods": "2", "hard": True}
    found = kich_ban.describe(scenario, [busy])["rules"][0]
    assert found["errors"] == [] and found["text"] == "Tiếng Anh 1 không dạy vào Thứ 2 tiết 2 (bắt buộc)"
    scenario["rules"].append({**busy, "role": "Không Có Ai"})
    assert any("không có giáo viên nào có chức vụ, Mã GV hay họ tên 'Không Có Ai'" in e
               for e in kich_ban.check(scenario, config.MODE_OVERTIME, 2)["errors"])
    assert config.PEOPLE == {}


def test_rules_only_file(tmp_path):
    """Xuất luật ra Excel (chỉ sheet LUẬT, HƯỚNG DẪN) rồi nhập lại: chỉ có phần luật; mẫu luật là các luật có sẵn."""
    scenario = kich_ban.default_scenario()
    _row(scenario, "nhom_buoi")["number"] = 3
    del scenario["rules"][-1]
    scenario["rules"].append({"kind": "Không xếp vào", "subject": "Tin học", "days": "Thứ 2", "hard": True})
    path = tmp_path / "luat.xlsx"
    kich_ban.rules_to_excel(scenario, path)
    assert openpyxl.load_workbook(path).sheetnames == ["LUẬT", "HƯỚNG DẪN"]
    assert kich_ban.sheets_in(path) == ["LUẬT"]
    back, _ = kich_ban.from_excel(path)
    assert back["staff"] == [] and back["grades"] == [] and back["rules"] == scenario["rules"][:-1] + [
        {**{k: "" for k in kich_ban.RULE_KEYS}, "kind": "Không xếp vào", "subject": "Tin học", "days": "Thứ 2",
         "hard": True, "number": None, "level": None, "points": None}]
    rules = read_rules(path)
    assert rules["SESSION_GROUP_LIMIT"] == 3 and rules["OFF"] == frozenset({"tiet_trong"})

def test_history_and_leave_as_the_page_writes_them(tmp_path):
    """Trang Giáo viên ghi Lớp Đang Dạy, Buổi Nghỉ bằng ô đánh dấu, ra chữ dạng "3/1, 4/2" và "Chiều T5, 2 buổi
    chiều, 1 buổi sáng, 1 buổi": chương trình đọc đúng như khi ghi tay."""
    scenario, _ = kich_ban.from_excel(INPUT_FILE)
    i = next(i for i, row in enumerate(scenario["staff"]) if row["role"] == "Tiếng Anh")
    scenario["staff"][i] |= {"history": "3/1, 3/2, 4/2", "off": "Chiều T5, 2 buổi chiều, 1 buổi sáng, 1 buổi"}
    kich_ban.to_excel(scenario, tmp_path / "ra.xlsx")
    t = next(t for t in read_staff(tmp_path / "ra.xlsx") if t.row == i + 2)
    assert t.history == {"3/1", "3/2", "4/2"} and t.off_sessions == {(3, "Chiều")}
    assert dict(t.off_any) == {"Chiều": 2, "Sáng": 1, None: 1}
    assert kich_ban.check(scenario, config.MODE_OVERTIME, 2)["errors"] == []


def test_quick_check_and_sample():
    """Kiểm tra nhanh (giao diện gọi mỗi lần sửa) ra cùng lỗi với Kiểm tra đầy đủ, chỉ bỏ dự toán; trường mẫu của trang
    bắt đầu là trường mẫu tên giả, kiểm tra không lỗi."""
    sample = kich_ban.sample_scenario()
    assert len(sample["staff"]) == 45 and sample["staff"][0]["name"] == "Giáo viên CN 1"
    full, quick = kich_ban.check(sample, config.MODE_OVERTIME, 3), kich_ban.check(sample, config.MODE_OVERTIME, 3,
                                                                                    quick=True)
    assert full["errors"] == quick["errors"] == [] and len(quick["info"]) == 1 < len(full["info"])
    assert full["unchecked"] == quick["unchecked"] == []
    sample["staff"][2]["lessons"] = None
    assert kich_ban.check(sample, quick=True)["errors"] == kich_ban.check(sample)["errors"] == [
        "NHÂN SỰ: Dòng 4: Số tiết của 'Chủ Nhiệm 1/3' bị trống hoặc không hợp lệ"]
    assert kich_ban.check(sample, quick=True)["unchecked"] == ["LUẬT", ""]  # chưa đếm chéo được
    sample["rules"].append({"kind": "Học trước", "subject": "Toán"})  # luật ghi sai: báo trước khi đọc nhân sự
    res = kich_ban.check(sample, quick=True)
    assert res["errors"] == ["LUẬT, dòng 26: kiểu luật Học trước phải ghi cột Môn thứ hai"]
    assert res["unchecked"] == ["CHƯƠNG TRÌNH HỌC", "NHÂN SỰ", "LUẬT", ""]  # trang hiện "?", không hiện ✓
