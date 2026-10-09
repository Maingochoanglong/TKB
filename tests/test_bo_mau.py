"""Bộ luật mẫu (tkb/bo_mau.py): giá trị mặc định của chương trình là bộ Tiểu học Việt Nam (dữ liệu, không phải code);
bộ Trống không có quy định nào gắn tên môn, các luật có sẵn theo quy ước để Tạm tắt."""
import openpyxl

from tkb import bo_mau, config, kich_ban, luat_co_san
from tkb.__main__ import main
from tkb.rules import applied, code, mau, read_rules
from tkb.staff import read_saved_timetable
from tkb.template import write_staff_template

OFF = {n.key for n in luat_co_san.NATIVES} - bo_mau.TRONG_ON - {"toi_da_ngay"}  # toi_da_ngay: bộ Trống không có môn


def test_defaults_are_the_vietnamese_preset():
    """config chỉ lấy giá trị từ bộ mẫu Tiểu học Việt Nam; áp bộ mẫu đó là như không áp gì (cùng mã quy định)."""
    for attr, value in bo_mau.TIEU_HOC_VN.items():
        got = getattr(config, attr)
        assert ([(r.subject, r.grade) for r in got] if attr == "MANAGER_RULES" else got) == value, attr
    with applied(None):
        default = code()
    with applied(mau("tieu_hoc_vn")):
        assert code() == default and config.RULES is None


def test_empty_preset_template(tmp_path):
    path = tmp_path / "trong.xlsx"
    write_staff_template(path, mau="trong")
    wb = openpyxl.load_workbook(path)
    program = list(wb["CHƯƠNG TRÌNH HỌC"].iter_rows(values_only=True))
    assert program[0][:6] == ("Môn học", "Khối 1", "Khối 2", "Khối 3", "Khối 4", "Khối 5")
    assert not any(row[0] for row in program[1:])  # chưa có môn nào
    rules = read_rules(path)
    with applied(rules):
        assert config.OFF == OFF and not config.CUSTOM_RULES
        assert config.HDTN == "" and not config.HOMEROOM_PERIODS and not config.HEAVY_SUBJECTS
        assert not config.HOMEROOM_PRIORITY and not config.GENERAL_FORBIDDEN_SUBJECTS and not config.MANAGER_RULES
        assert all([s.name for s in config.DAY_SESSIONS[d]] == ["Sáng", "Chiều"] for d in range(5))
        assert {luat_co_san.native_of(r).key for r in config.RULES if r.off} == OFF
    # Kịch bản Soạn mới bộ Trống trên giao diện: cùng nội dung với file mẫu đó.
    sc = kich_ban.default_scenario("trong")
    assert sc["subjects"] == [] and sum(r["off"] for r in sc["rules"]) == len(sc["rules"]) - len(bo_mau.TRONG_ON)
    assert kich_ban.from_excel(path)[0]["rules"] == sc["rules"]


def test_school_from_the_empty_preset(tmp_path, capsys):
    """Một trường soạn từ bộ Trống, không theo quy ước Việt Nam: môn tên tiếng Anh, khối "Year 1", hai lớp không có
    GVCN, chỉ có giáo viên bộ môn; luật riêng: Math không xếp buổi chiều. Xếp đạt mọi luật."""
    path = tmp_path / "school.xlsx"
    sc = kich_ban.default_scenario("trong")
    lessons = {"Reading": 5, "Math": 5, "English": 4, "Science": 3, "Art": 2, "PE": 2}
    sc["grades"] = ["Year 1"]
    sc["subjects"] = [{"name": s, "lessons": {"Year 1": n}, "rules": {}} for s, n in lessons.items()]
    sc["classes"] = [{"name": c, "grade": "Year 1", "campus2": False} for c in ("1 Blue", "1 Green")]
    sc["staff"] = [{"name": f"Teacher {i}", "role": "Bộ Môn", "lessons": 20} for i in range(1, 4)]
    blank = {k: "" for k in kich_ban.RULE_KEYS}
    sc["rules"].append({**blank, "kind": "Không xếp vào", "subject": "Math", "sessions": "Chiều", "hard": True})
    kich_ban.to_excel(sc, path)
    argv = [str(path), "-o", str(tmp_path / "TKB.xlsx"), "--time-limit", "10", "--workers", "4", "--mode", "tuyen_them"]
    assert main(argv) == 0
    assert "kiểm tra luật bắt buộc: ĐẠT" in capsys.readouterr().out
    with applied(read_rules(path)):
        saved = read_saved_timetable(tmp_path / "school_cap_nhat.xlsx")
    rows = saved.rows
    assert len(rows) == 2 * sum(lessons.values()) and {r[0] for r in rows} == {"1 Blue", "1 Green"}
    assert not [r for r in rows if r[3] == "Math" and r[2] >= 5]  # tiết 5–7 là buổi chiều
    assert openpyxl.load_workbook(tmp_path / "TKB.xlsx").sheetnames == ["Khối Year 1"]
