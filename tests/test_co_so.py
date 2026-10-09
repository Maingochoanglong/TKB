"""Nhiều cơ sở: cột Cơ sở của sheet LỚP ghi tên cơ sở (điểm trường) tùy ý, bao nhiêu cơ sở cũng được (trống: Cơ sở 1;
cột Có/Không Cơ sở 2 của bản trước vẫn đọc được). Luật mỗi buổi một cơ sở (cứng) và hạn chế đổi cơ sở trong ngày
(mềm) áp cho mọi số cơ sở; hai cơ sở Cơ sở 1, Cơ sở 2 giữ cách mã hóa cũ (cùng mã kết quả, tests/test_reproducible.py
và tests/test_teacher_rules.py). Tên cơ sở cũng là nhãn lớp ở cột Lớp của sheet LUẬT."""
import dataclasses
from collections import defaultdict

import openpyxl
import pytest

from tkb import bo_mau, config, kich_ban
from tkb.__main__ import main
from tkb.checker import check
from tkb.program import read_program
from tkb.rules import applied, code, read_rules
from tkb.solver import solve
from tkb.staff import InputError, read_saved_timetable, read_staff
from tkb.template import write_staff_template

from .conftest import CURRICULUM, small_staff

LA = {**{s: 0 for s in CURRICULUM[3]}, bo_mau.TV: 6, bo_mau.TOAN: 4, bo_mau.TIENG_ANH: 2, "Thể dục": 2,
      bo_mau.AM_NHAC: 2, bo_mau.MY_THUAT: 2}
THREE = [("3/1", 3, ""), ("3/2", 3, "Điểm Hòa Bình"), ("Lá 1", "Lá", "Điểm Tân Phú")]
BLANK = {k: "" for k in kich_ban.RULE_KEYS}
SESSION = {p: s.name for ss in config.DAY_SESSIONS.values() for s in ss for p in s.periods}


def _file(tmp_path, classes=THREE, rules=(), name="vao"):
    """Trường nhỏ (2 lớp Khối 3 có Chủ Nhiệm) thêm lớp Lá 1 không có Chủ Nhiệm; sheet LỚP ghi `classes` [(lớp, khối,
    cơ sở)]."""
    path = tmp_path / f"{name}.xlsx"
    write_staff_template(path, small_staff(), {3: CURRICULUM[3], "Lá": LA})
    sc, _ = kich_ban.from_excel(path)
    sc["classes"] = [{"name": n, "grade": g, "campus": c} for n, g, c in classes]
    sc["rules"] += [{**BLANK, "hard": True, **r} for r in rules]
    kich_ban.to_excel(sc, path)
    return path


def test_three_campuses_end_to_end(tmp_path, capsys):
    """Ba cơ sở: mỗi buổi mỗi GV chỉ ở một cơ sở; mỗi cơ sở một file TKB; luật theo tên cơ sở (Tiếng Anh ở Điểm Tân
    Phú không xếp Thứ 2); thống kê ghi các buổi ở cơ sở khác."""
    path = _file(tmp_path, rules=[{"kind": "Không xếp vào", "subject": bo_mau.TIENG_ANH, "classes": "điểm tân phú",
                                   "days": "Thứ 2"}])
    argv = [str(path), "-o", str(tmp_path / "TKB.xlsx"), "--time-limit", "10", "--workers", "4"]
    assert main(argv) == 0
    out = capsys.readouterr().out
    assert "kiểm tra luật bắt buộc: ĐẠT" in out
    assert "Cơ sở 1: 1 lớp, Điểm Hòa Bình: 1 lớp, Điểm Tân Phú: 1 lớp" in out
    with applied(read_rules(path)):
        rows = read_saved_timetable(tmp_path / "vao_cap_nhat.xlsx").rows
    campus = {n: c or "Cơ sở 1" for n, _, c in THREE}
    where = defaultdict(set)
    for cls, day, period, subject, who, *_ in rows:
        where[who, day, SESSION[period]].add(campus[cls])
    assert all(len(cs) == 1 for cs in where.values())
    assert len({c for (who, *_), cs in where.items() if who == "Thể Dục 1" for c in cs}) == 3  # dạy ở cả ba cơ sở
    assert not [r for r in rows if r[0] == "Lá 1" and r[3] == bo_mau.TIENG_ANH and r[1] == "Thứ 2"]
    for stem, cls in (("diem_chinh", "3/1"), ("diem_hoa_binh", "3/2"), ("diem_tan_phu", "Lá 1")):
        ws = openpyxl.load_workbook(tmp_path / f"TKB_{stem}.xlsx").worksheets[0]
        assert {c.value for row in ws.iter_rows() for c in row if c.value in ("3/1", "3/2", "Lá 1")} == {cls}
    stats = openpyxl.load_workbook(tmp_path / "Thong_Ke.xlsx")["Thống kê"]
    assert [c.value for c in stats[1]][-2:] == ["Buổi Ở Cơ Sở Khác", "Đổi Cơ Sở Trong Ngày"]
    teachers = openpyxl.load_workbook(tmp_path / "TKB_giao_vien.xlsx").worksheets[0]
    assert any("Lá 1 (Điểm Tân Phú)" in str(c.value) for row in teachers.iter_rows() for c in row)


def test_checker_and_campus_rule_with_three_campuses(tmp_path):
    """Bộ kiểm tra độc lập: một GV dạy ba cơ sở trong một buổi là sai luật mỗi buổi một cơ sở."""
    path = _file(tmp_path)
    with applied(read_rules(path)):
        sol = solve(read_staff(path), read_program(path), config.Settings(time_limit=10, workers=4),
                    log=lambda *_: None)
        assert sol.problem.campuses == ("Cơ sở 1", "Điểm Hòa Bình", "Điểm Tân Phú")
        assert check(sol.problem, sol.lessons) == []
        # Ba tiết khác giờ của ba lớp trong cùng một buổi giao cho GV thể dục.
        day, session = 0, SESSION[1]
        picked, periods = [], set()
        for cls, _, _ in THREE:
            les = next(l for l in sol.lessons if l.class_name == cls and l.day == day and SESSION[l.period] == session
                       and l.period not in periods)
            picked.append(les)
            periods.add(les.period)
        moved = [dataclasses.replace(l, teacher="thể dục 1") if l in picked else l for l in sol.lessons]
        errors = check(sol.problem, moved)
    assert any("thể dục 1 dạy 3 cơ sở (cơ sở 1, Điểm Hòa Bình, Điểm Tân Phú) trong buổi" in e for e in errors), errors


def test_old_campus_column_and_errors(tmp_path):
    """Cột Có/Không Cơ sở 2 của bản trước: lớp ghi Có ở Cơ sở 2, cùng mã quy định với cột Cơ sở ghi "Cơ sở 2". Tên cơ
    sở trùng tên lớp thì báo lỗi (tên cơ sở cũng là nhãn lớp)."""
    new = _file(tmp_path, classes=[("3/1", 3, ""), ("3/2", 3, "cơ sở 2"), ("Lá 1", "Lá", "")])
    old = tmp_path / "cu.xlsx"
    wb = openpyxl.load_workbook(new)
    ws = wb["LỚP"]
    ws.cell(1, 3).value = "Cơ sở 2"
    ws.cell(3, 3).value = "Có"
    wb.save(old)
    codes = []
    for path in (new, old):
        values = read_rules(path)
        assert [c.campus.lower() for c in values["CLASSES"]] == ["", "cơ sở 2", ""]
        with applied(values):
            codes.append(code())
    assert codes[0] == codes[1]
    bad = _file(tmp_path, classes=[("3/1", 3, ""), ("3/2", 3, "3/1"), ("Lá 1", "Lá", "")], name="sai")
    with pytest.raises(InputError, match="LỚP, dòng 3: tên cơ sở '3/1' trùng tên lớp ở dòng 2"):
        read_rules(bad)
