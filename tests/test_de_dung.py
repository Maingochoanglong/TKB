"""File Excel do người không rành IT điền tay (trường mẫu tên giả): tên sheet, tiêu đề cột viết không dấu hay thừa dấu
cách, ô Có/Không ghi ✓, chức vụ "GVCN", ô gộp, dòng ẩn, dòng trùng, định mức gõ nhầm, dòng Tổng, tên gõ khác dấu,
file không phải .xlsx: đọc đúng hoặc báo rõ chỗ sửa, không lặng lẽ bỏ qua."""
import shutil
import subprocess
import sys
from dataclasses import replace

import openpyxl
import pytest

from tkb import config, kich_ban
from tkb.allocation import build_problem
from tkb.program import read_program
from tkb.rules import applied, read_rules
from tkb.staff import (InputError, near_sheets, parse_class, parse_classes, parse_off, read_staff, role_text,
                       suggest, to_number)

from .conftest import INPUT_FILE


def _copy(tmp_path, change=None):
    """Bản sao file trường mẫu, sửa bằng `change(wb)`."""
    path = tmp_path / "Truong.xlsx"
    shutil.copy(INPUT_FILE, path)
    if change is not None:
        wb = openpyxl.load_workbook(path)
        change(wb)
        wb.save(path)
    return path


def _header(ws, name):
    return next(c for c in ws[1] if c.value == name)


def test_sheet_and_column_names_without_accents(tmp_path):
    """Sheet QUY ĐỊNH đổi tên "Quy dinh" vẫn được đọc (trước: lặng lẽ dùng khung giờ mặc định); tiêu đề cột nhân sự
    viết không dấu, thừa dấu cách, kiểu khác ("Ho va Ten", "So tiet/ tuan", "Hop dong")."""
    def change(wb):
        ws = wb["QUY ĐỊNH"]
        for row in ws.iter_rows():
            if row[0].value == "Thứ 6":
                row[2].value = 3  # thêm buổi chiều Thứ 6
        ws.title = "Quy dinh"
        staff = wb["NHÂN SỰ"]
        _header(staff, "Họ và Tên").value = "Ho va Ten"
        _header(staff, "Số Tiết/Tuần").value = "So tiet/ tuan"
        _header(staff, "Hợp Đồng").value = "Hop dong"
        staff.cell(3, _header(staff, "Hop dong").column).value = "✓"
    path = _copy(tmp_path, change)
    values = read_rules(path)
    assert len(values["DAY_SESSIONS"][4]) == 2  # Thứ 6 có buổi chiều: sheet đã được đọc
    with applied(values):
        staff = read_staff(path)
    assert staff[1].contract and not staff[0].contract  # ✓ là Có


def test_yes_no_words_roles_and_classes(tmp_path):
    def change(wb):
        ws = wb["NHÂN SỰ"]
        col = _header(ws, "Hợp Đồng").column
        ws.cell(2, col).value, ws.cell(3, col).value, ws.cell(4, col).value = "Đúng", "ko", "v"
        _header(ws, "Chức Vụ")
        ws.cell(2, _header(ws, "Chức Vụ").column).value = "GVCN"
    path = _copy(tmp_path, change)
    staff = read_staff(path)
    assert [t.contract for t in staff[:3]] == [True, False, True]
    assert staff[0].code.startswith("Chủ Nhiệm ")
    assert role_text("gv bộ môn") == "Bộ Môn" and role_text("Tiếng Anh") == "Tiếng Anh"
    from tkb.staff import parse_classes
    assert parse_classes("3/1 và 3/2; 3/3") == {"3/1", "3/2", "3/3"}
    with pytest.raises(InputError, match="cột Lớp Đang Dạy"):
        parse_classes("3/1 va ba")


def test_merged_cells_hidden_rows_and_unread_columns(tmp_path):
    """Ô gộp ở cột Chức Vụ: mọi dòng của vùng gộp ghi chức vụ đó (trước: "thiếu chức vụ"); dòng ẩn và cột tiêu đề gần
    giống vẫn đọc như cũ nhưng có cảnh báo."""
    rows = {}

    def change(wb):
        ws = wb["NHÂN SỰ"]
        col = _header(ws, "Chức Vụ").column
        last = [r for r in range(2, ws.max_row + 1) if ws.cell(r, col).value == "Bộ Môn"][-3:]
        rows["merged"] = last
        for r in last[1:]:
            ws.cell(r, col).value = None
        ws.merge_cells(start_row=last[0], start_column=col, end_row=last[-1], end_column=col)
        ws.row_dimensions[2].hidden = True
        _header(ws, "Thứ Tự Bù").value = "Thứ tự bù (GV)"
    path = _copy(tmp_path, change)
    notes = []
    staff = read_staff(path, warn=notes.append)
    by_row = {t.row: t for t in staff}
    assert all(by_row[r].role == config.ROLE_GENERAL for r in rows["merged"])
    assert any("ô gộp" in n for n in notes)
    assert any("dòng 2: dòng đang ẩn" in n for n in notes) and 2 in by_row
    assert any("'Thứ tự bù (GV)' không được đọc" in n and "Thứ Tự Bù" in n for n in notes)
    # Trang nhập liệu đọc như dòng lệnh và hiện cùng các cảnh báo.
    scenario, warnings = kich_ban.from_excel(path)
    assert len(scenario["staff"]) == len(staff) and any("ô gộp" in w for w in warnings)


def test_two_row_header(tmp_path):
    """Tiêu đề hai tầng: các ô tiêu đề gộp dọc (A1:A2...), dòng 2 ghi "Số Tiết/Tuần" dưới ô "Định mức" (trước: lỗi
    không tìm thấy tiêu đề; rồi dòng 2 bị đọc như một người)."""
    count = {}

    def change(wb):
        ws = wb["NHÂN SỰ"]
        count["rows"] = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value)
        lessons = _header(ws, "Số Tiết/Tuần").column
        ws.insert_rows(2)
        for c in range(1, ws.max_column + 1):
            if c == lessons:
                ws.cell(2, c).value, ws.cell(1, c).value = ws.cell(1, c).value, "Định mức"
            elif ws.cell(1, c).value is not None:
                ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)
    path = _copy(tmp_path, change)
    notes = []
    staff = read_staff(path, warn=notes.append)
    assert len(staff) == count["rows"] and min(t.row for t in staff) == 3 and notes == []
    scenario, warnings = kich_ban.from_excel(path)
    assert len(scenario["staff"]) == len(staff) and not any("NHÂN SỰ" in w for w in warnings)


def test_near_sheet_names(tmp_path):
    def change(wb):
        wb.copy_worksheet(wb["NHÂN SỰ"]).title = "NHÂN SỰ (2)"
    path = _copy(tmp_path, change)
    wb = openpyxl.load_workbook(path)
    notes = near_sheets(wb, (config.STAFF_SHEET, config.PROGRAM_SHEET))
    assert notes == ["Sheet 'NHÂN SỰ (2)' không được đọc (chương trình đọc sheet 'NHÂN SỰ')"]


def test_same_names_quota_and_total_row(tmp_path):
    """Dòng chép trùng: cảnh báo ghi Mã GV, không ghi họ tên; định mức lớn hơn số tiết cả tuần: lỗi ngay khi kiểm
    tra (trước: chạy hết thời gian mới báo không xếp được); dòng "Tổng số tiết" của chương trình học bỏ qua."""
    def change(wb):
        ws = wb["NHÂN SỰ"]
        name = _header(ws, "Họ và Tên").column
        ws.cell(41, name).value = ws.cell(40, name).value
        ws.cell(40, _header(ws, "Số Tiết/Tuần").column).value = 230
        prog = wb["CHƯƠNG TRÌNH HỌC"]
        prog.cell(prog.max_row + 1, 1).value = "Tổng số tiết"
        prog.cell(prog.max_row, 2).value = 999
    path = _copy(tmp_path, change)
    curriculum = read_program(path)
    assert not any("Tổng" in s for req in curriculum.values() for s in req)
    staff = read_staff(path)
    with pytest.raises(InputError, match="Dòng 40: .* ghi Số Tiết/Tuần 230, nhiều hơn số tiết cả tuần"):
        build_problem(staff, curriculum, {})
    staff = [replace(t, max_lessons=23) if t.row == 40 else t for t in staff]
    problem = build_problem(staff, curriculum, {})
    same = [w for w in problem.warnings if "cùng họ tên" in w]
    name = next(t.name for t in staff if t.row == 40)
    assert len(same) == 1 and "Các dòng 40, 41" in same[0] and name not in same[0]


def test_suggestions():
    assert suggest("Toan", ["Toán", "Tiếng Việt"]) == " (ý là 'Toán'?)"
    assert suggest("Tieng Anh", ["Tiếng Anh"]) == " (ý là 'Tiếng Anh'?)"
    assert suggest("Toán", ["Toán"]) == "" and suggest("Âm nhạc", ["Toán"]) == ""


def test_not_xlsx(tmp_path):
    path = tmp_path / "nhan_su.csv"
    path.write_text("Họ và Tên,Chức Vụ\n", encoding="utf-8")
    run = subprocess.run([sys.executable, "-m", "tkb", str(path), "-o", str(tmp_path / "TKB.xlsx")],
                         capture_output=True, text=True, encoding="utf-8")
    assert run.returncode == 1
    assert "chỉ đọc file Excel .xlsx" in run.stderr and "Traceback" not in run.stderr


def test_numbers_written_by_hand(tmp_path):
    """Số tiết ghi "23 tiết", "23,0"; chương trình học ghi "14 tiết", "-" (là 0): đọc được (trước: lỗi)."""
    assert to_number("23 tiết") == to_number("23,0") == to_number(" 23t ") == 23
    with pytest.raises(ValueError):
        to_number("nghỉ ốm")
    found = {}

    def change(wb):
        ws = wb["NHÂN SỰ"]
        col = _header(ws, "Số Tiết/Tuần").column
        found["row"], found["lessons"] = 31, ws.cell(31, col).value
        ws.cell(31, col).value = f"{ws.cell(31, col).value} tiết"
        ws.cell(32, col).value = f"{ws.cell(32, col).value},0"
        prog = wb["CHƯƠNG TRÌNH HỌC"]
        found["subject"], found["first"] = prog.cell(2, 1).value, prog.cell(2, 2).value
        prog.cell(2, 2).value = f"{prog.cell(2, 2).value} tiết"
        found["zero"] = next(prog.cell(r, 1).value for r in range(2, prog.max_row + 1) if prog.cell(r, 2).value == 0)
        r = next(r for r in range(2, prog.max_row + 1) if prog.cell(r, 2).value == 0)
        prog.cell(r, 2).value = "-"
    path = _copy(tmp_path, change)
    staff = {t.row: t for t in read_staff(path)}
    assert staff[found["row"]].max_lessons == found["lessons"]
    curriculum = read_program(path)
    assert curriculum[1][found["subject"]] == found["first"] and curriculum[1][found["zero"]] == 0
    scenario, _ = kich_ban.from_excel(path)
    assert scenario["staff"][found["row"] - 2]["lessons"] == found["lessons"]


def test_footer_rows_and_sheet_names(tmp_path):
    """Dòng "Tổng cộng", "Ghi chú: ..." cuối bảng nhân sự bỏ qua; sheet nhân sự đổi tên "Nhân sự 2026" và chuyển ra
    sau vẫn được đọc; tiêu đề chương trình học "Tên môn", "Khối1", hoặc "Lớp 1"... (khối của tiểu học)."""
    base = len(read_staff(_copy(tmp_path)))

    def change(wb):
        ws = wb["NHÂN SỰ"]
        last = ws.max_row
        ws.cell(last + 1, 1).value, ws.cell(last + 1, _header(ws, "Số Tiết/Tuần").column).value = "Tổng cộng", 893
        ws.cell(last + 3, 1).value = "Ghi chú: danh sách cập nhật 01/9"
        ws.title = "Nhân sự 2026"
        wb.move_sheet(ws, offset=1)
        prog = wb["CHƯƠNG TRÌNH HỌC"]
        prog.cell(1, 1).value = "Tên môn"
        for c in range(2, prog.max_column + 1):
            if str(prog.cell(1, c).value).startswith("Khối "):
                prog.cell(1, c).value = prog.cell(1, c).value.replace("Khối ", "Lớp ")
    (tmp_path / "x").mkdir()
    path = _copy(tmp_path / "x", change)
    notes = []
    values = read_rules(path, warn=notes.append)
    with applied(values):
        assert len(read_staff(path)) == base
        assert sorted(read_program(path)) == [1, 2, 3, 4, 5]
    assert not any("Lớp 1" in n or "không được đọc" in n for n in notes)
    scenario, warnings = kich_ban.from_excel(path)
    assert len(scenario["staff"]) == base and len(scenario["grades"]) == 5


def test_class_names_and_leave_written_by_hand():
    assert parse_class("Lớp 3/2") == "3/2" and parse_class("3.1") == "3/1" and parse_class("3-4") == "3/4"
    with pytest.raises(InputError, match="Excel bỏ số 0 cuối"):
        parse_class(3.1)  # ô số: "3.10" đã thành 3.1
    assert parse_classes("3/1 3/2") == parse_classes("3/1 - 3/2") == {"3/1", "3/2"}
    afternoon = {(3, "Chiều")}
    assert parse_off("T5 chiều")[0] == parse_off("chiều thứ năm")[0] == parse_off("Chiều T5")[0] == afternoon
    assert parse_off("Thứ 6")[0] == {(4, "Sáng")}  # cả ngày: chiều Thứ 6 vốn nghỉ
    assert parse_off("2 buổi sáng") == (frozenset(), (("Sáng", 2),))  # không phải sáng Thứ 2
    with pytest.raises(InputError, match="Buổi Nghỉ"):
        parse_off("2")


def test_empty_staff_and_renamed_xlsx(tmp_path):
    def change(wb):
        ws = wb["NHÂN SỰ"]
        ws.delete_rows(2, ws.max_row)
    with pytest.raises(InputError, match="chưa có giáo viên nào"):
        read_staff(_copy(tmp_path, change))
    path = tmp_path / "Truong.xls"
    shutil.copy(INPUT_FILE, path)
    run = subprocess.run([sys.executable, "-m", "tkb", str(path), "-o", str(tmp_path / "TKB.xlsx")],
                         capture_output=True, text=True, encoding="utf-8")
    assert run.returncode == 1 and "đổi tên file thành Truong.xlsx" in run.stderr
