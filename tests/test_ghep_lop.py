"""Kiểu luật Ghép lớp (phép đo Cùng giờ, so sánh Cùng một người): các lớp học chung một môn, cùng giờ, một người dạy;
người đó tính một giờ dạy cho cả nhóm (Course.lead, Problem.merge_groups, Problem.taught). Không có luật này: như
trước."""
import dataclasses
from collections import defaultdict

import openpyxl
import pytest

from tkb import bo_mau, config, kich_ban
from tkb.__main__ import main
from tkb.allocation import build_problem
from tkb.checker import check
from tkb.program import read_program
from tkb.rules import applied, read_rules
from tkb.solver import solve, timetable
from tkb.staff import InputError, read_saved_timetable, read_staff
from tkb.template import write_staff_template

from .conftest import CURRICULUM, INPUT_FILE, small_staff, teacher

BLANK = {k: "" for k in kich_ban.RULE_KEYS}
TD = "Thể dục"
MERGE = {"kind": "Ghép lớp", "subject": TD, "classes": "3/1, 3/2"}
SETTINGS = config.Settings(time_limit=10, workers=4)


def _file(tmp_path, rules=(), staff=None, rooms=(), name="vao", source=None):
    """Trường nhỏ (2 lớp Khối 3; `source`: file vào khác), thêm các dòng luật `rules` (bắt buộc) và sheet PHÒNG
    `rooms`. Ghi qua kịch bản như giao diện."""
    path = tmp_path / f"{name}.xlsx"
    if source is None:
        write_staff_template(path, small_staff() if staff is None else staff, {3: CURRICULUM[3]})
    sc, _ = kich_ban.from_excel(source or path)
    sc["rules"] += [{**BLANK, "hard": True, **r} for r in rules]
    sc["rooms"] = [{"campus": "", "grades": "", "capacity": None, **r} for r in rooms]
    kich_ban.to_excel(sc, path)
    return path


def _read(path):
    curriculum = read_program(path)
    return read_staff(path, subjects=[s for req in curriculum.values() for s in req]), curriculum


def _merged(lessons, classes, subject):
    """(lớp -> các (giờ, người dạy)) của môn ở các lớp."""
    at = defaultdict(set)
    for les in lessons:
        if les.class_name in classes and les.subject == subject:
            at[les.class_name].add((les.day, les.period, les.teacher))
    return at


@pytest.mark.parametrize("mode", [config.MODE_HIRE, config.MODE_OVERTIME])
def test_merge_end_to_end(tmp_path, capsys, mode):
    """Thể dục 3/1, 3/2 ghép lớp: cùng giờ, cùng người; người dạy tính 2 giờ (không phải 4), dự toán bớt 2 tiết; ô TKB
    ghi "(ghép …)", TKB giáo viên ghi "3/1, 3/2"; nạp lại dùng lại TKB; trang web đánh dấu ô ghép, không đổi tay."""
    path = _file(tmp_path, [MERGE])
    argv = [str(path), "-o", str(tmp_path / "TKB.xlsx"), "--time-limit", "10", "--workers", "4", "--mode", mode,
            "--max-overtime", "2"]
    assert main(argv) == 0
    out = capsys.readouterr().out
    total = 2 * sum(CURRICULUM[3].values())
    assert "kiểm tra luật bắt buộc: ĐẠT" in out and f"Dự toán: cần {total - 2} tiết" in out
    updated = tmp_path / "vao_cap_nhat.xlsx"
    with applied(read_rules(updated)):
        rows = read_saved_timetable(updated).rows
    at = defaultdict(set)
    for cls, day, period, subject, code, *_ in rows:
        if subject == TD:
            at[cls].add((day, period, code))
    assert at["3/1"] == at["3/2"] and len(at["3/1"]) == CURRICULUM[3][TD]
    cells = [str(c.value) for ws in openpyxl.load_workbook(tmp_path / "TKB.xlsx") for row in ws.iter_rows()
             for c in row if c.value]
    assert sum(v.startswith(f"{TD} (ghép 3/2)\n") for v in cells) == 2
    assert sum(v.startswith(f"{TD} (ghép 3/1)\n") for v in cells) == 2
    teachers = [str(c.value) for row in openpyxl.load_workbook(tmp_path / "TKB_giao_vien.xlsx")["Giáo viên"]
                .iter_rows() for c in row if c.value]
    assert sum(v == f"3/1, 3/2\n{TD}" for v in teachers) == 2
    assert any(v.endswith(": 2 tiết") and "Thể Dục 1" in v for v in teachers), teachers
    stats = openpyxl.load_workbook(tmp_path / "Thong_Ke.xlsx")["Thống kê"]
    head = next(stats.iter_rows(values_only=True))
    line = next(r for r in stats.iter_rows(values_only=True) if r[1] == "Thể Dục 1")
    assert line[head.index(TD)] == 2 and line[head.index("Tổng Tiết")] == 2
    again = [str(updated), "-o", str(tmp_path / "lai" / "TKB.xlsx"), "--time-limit", "10", "--workers", "4", "--mode",
             mode, "--max-overtime", "2"]
    assert main(again) == 0 and "Dùng lại TKB đã xếp" in capsys.readouterr().out
    sc = kich_ban.from_excel(updated)[0]
    grid = kich_ban.Grid(sc, mode=mode, overtime_max=2)
    view = grid.view(sc["saved"])
    assert view["ok"], view["errors"]
    merged = [c for c in view["cells"] if c.get("merged")]
    assert len(merged) == 4 and {tuple(c["merged"]) for c in merged} == {("3/1",), ("3/2",)}
    one = merged[0]
    assert grid.swaps(sc["saved"], one["cls"], one["d"], one["p"]) == []
    other = next(c for c in view["cells"] if c["cls"] == "3/1" and c["subject"] and not c.get("merged"))
    assert all((t["d"], t["p"]) not in {(c["d"], c["p"]) for c in merged if c["cls"] == "3/1"}
               for t in grid.swaps(sc["saved"], "3/1", other["d"], other["p"]))


def test_rule_errors(tmp_path):
    """Ghi sai: Tự ghép Với mỗi Lớp, ưu tiên, có Ngày; các lớp khác số tiết, một lớp, lớp có hai môn của nhóm; môn chỉ
    GVCN dạy; không ai dạy được cả nhóm."""
    rules = [{**MERGE, "kind": "Tự ghép", "measure": "Cùng giờ", "op": "Cùng một người", "scope": "Lớp"},
             {**MERGE, "hard": False}, {**MERGE, "days": "Thứ 2"}]
    with pytest.raises(InputError) as err:
        read_rules(_file(tmp_path, rules))
    for text in ("dòng 26: phép đo Cùng giờ, so sánh Cùng một người (ghép lớp): cột Với mỗi để trống hoặc ghi Khối",
                 "dòng 27: ghép lớp chỉ ghi Bắt buộc", "dòng 28: cột Ngày không dùng cho kiểu luật Ghép lớp"):
        assert text in str(err.value), str(err.value)
    rules = [{**MERGE, "other": bo_mau.TIN_HOC}, {**MERGE, "classes": "3/1"},
             {"kind": "Ghép lớp", "subject": bo_mau.AM_NHAC, "other": bo_mau.TIN_HOC, "classes": "3/1, 3/2"}]
    errors = kich_ban.check(kich_ban.from_excel(_file(tmp_path, rules, name="sai"))[0], quick=True)["errors"]
    for text in (f"dòng 26: ghép lớp {TD}: lớp 3/1, 3/2 có nhiều môn của nhóm",
                 f"dòng 27: ghép lớp {TD}: chỉ có lớp 3/1: ghép lớp cần ít nhất hai lớp",
                 f"dòng 28: ghép lớp {bo_mau.AM_NHAC}: lớp 3/1, 3/2 có nhiều môn"):
        assert any(text in e for e in errors), errors
    path = _file(tmp_path, [{"kind": "Ghép lớp", "subject": bo_mau.HDTN, "classes": "3/1, 3/2"}], name="gvcn")
    with applied(read_rules(path)), pytest.raises(InputError, match="chỉ GVCN được dạy .* nên không ghép lớp được"):
        build_problem(*_read(path))
    # Hai GV Thể dục, mỗi lớp chỉ một người được dạy: không ai dạy được cả nhóm.
    staff = small_staff() + [teacher("TD2", "thể dục 2", 23, row=10)]
    who = [{"kind": "Chỉ giáo viên dạy", "subject": TD, "classes": cls, "role": code}
           for cls, code in (("3/1", "Thể Dục 1"), ("3/2", "Thể Dục 2"))]
    path = _file(tmp_path, [MERGE, *who], staff=staff, name="khong_ai")
    with applied(read_rules(path)), pytest.raises(InputError, match="không GV nào được dạy môn này ở mọi lớp"):
        build_problem(*_read(path))


def test_checker_and_integrated_model(tmp_path):
    """Kiểm tra độc lập: hai lớp của nhóm cùng người cùng giờ không phải trùng giờ, tải tính một lần; dời tiết của
    một lớp khỏi giờ chung thì báo. Mô hình tích hợp (vừa phân công vừa xếp) cũng xếp được nhóm ghép lớp."""
    path = _file(tmp_path, [MERGE, {"kind": "Ghép lớp", "subject": bo_mau.TIENG_ANH, "classes": "3/1, 3/2"}])
    with applied(read_rules(path)):
        staff, curriculum = _read(path)
        sol = solve(staff, curriculum, SETTINGS, log=lambda *_: None)
        assert check(sol.problem, sol.lessons) == []
        load = {sol.problem.teachers[g].code: n for g, n in sol.teacher_load().items()}
        assert load["Thể Dục 1"] == 2 and load["Tiếng Anh 1"] == 4
        groups = sol.problem.merge_groups()
        assert len(groups) == 2 and all(len(m) == 2 for m in groups.values())
        at = _merged(sol.lessons, ("3/1", "3/2"), TD)
        assert at["3/1"] == at["3/2"]
        moved_one = next(l for l in sol.lessons if l.class_name == "3/2" and l.subject == TD)
        other = next(l for l in sol.lessons if l.class_name == "3/2" and l.subject == bo_mau.TV)
        moved = [dataclasses.replace(l, day=other.day, period=other.period) if l is moved_one else
                 dataclasses.replace(l, day=moved_one.day, period=moved_one.period) if l is other else l
                 for l in sol.lessons]
        errors = check(sol.problem, moved)
        assert any("ghép lớp: lớp 3/1 có tiết" in e and "mà lớp 3/2 không học cùng giờ" in e for e in errors), errors
        problem = build_problem(staff, curriculum, {config.ROLE_GENERAL: 1})
        res = timetable(problem, SETTINGS)
        assert res is not None and check(problem, res.lessons) == []
        at = _merged(res.lessons, ("3/1", "3/2"), bo_mau.TIENG_ANH)
        assert at["3/1"] == at["3/2"] and len(at["3/1"]) == CURRICULUM[3][bo_mau.TIENG_ANH]


def test_merge_in_one_room(tmp_path):
    """Sheet PHÒNG: cả nhóm ghép lớp học ở một phòng, chiếm một chỗ (sân 1 lớp vẫn chứa được nhóm 2 lớp)."""
    path = _file(tmp_path, [MERGE], rooms=[{"name": "Sân", "subjects": TD, "capacity": 1}])
    with applied(read_rules(path)):
        sol = solve(*_read(path), SETTINGS, log=lambda *_: None)
        assert check(sol.problem, sol.lessons) == []
        rooms = sol.rooms()
    used = defaultdict(set)
    for (cls, d, p, subject), name in rooms.items():
        assert subject == TD and name == "Sân"
        used[d, p].add(cls)
    assert len(used) == CURRICULUM[3][TD] and all(v == {"3/1", "3/2"} for v in used.values())


def test_merge_per_grade_on_sample_school(tmp_path):
    """Tự ghép Với mỗi Khối: Thể dục mỗi khối của trường mẫu là một nhóm; xếp đạt mọi luật bắt buộc (có các vòng xếp
    lại từng vùng), mỗi nhóm cùng giờ cùng người."""
    rule = {"kind": "Tự ghép", "measure": "Cùng giờ", "op": "Cùng một người", "scope": "Khối", "subject": TD,
            "grades": "1, 2"}
    path = _file(tmp_path, [rule], source=INPUT_FILE)
    with applied(read_rules(path)):
        staff, curriculum = _read(path)
        sol = solve(staff, curriculum, config.Settings(time_limit=15, workers=4), log=lambda *_: None)
        assert check(sol.problem, sol.lessons) == []
        assert sorted(len(m) for m in sol.problem.merge_groups().values()) == [6, 6]
        for g in (1, 2):
            classes = [c for c in sol.problem.classes if c.startswith(f"{g}/")]
            at = _merged(sol.lessons, classes, TD)
            assert len({frozenset(v) for v in at.values()}) == 1 and len(at) == 6
