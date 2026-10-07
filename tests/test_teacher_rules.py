"""Luật về GV theo các cột không bắt buộc của file vào: hai cơ sở (Cơ sở 2), thai sản, hợp đồng (thứ tự bù),
giữ phân công của TKB cũ (Lớp Đang Dạy), buổi nghỉ (Buổi Nghỉ)."""
import dataclasses
from collections import Counter, defaultdict

import pytest

from tkb import config
from tkb.allocation import build_problem
from tkb.checker import check
from tkb.phan_cong import phan_cong
from tkb.solver import solve

from .conftest import CURRICULUM, small_staff, teacher

W = config.Weights()
FAST = config.Settings(time_limit=20, workers=4)
SESSION = {p: s.name for ss in config.DAY_SESSIONS.values() for s in ss for p in s.periods}


def _set(staff, title, **fields):
    return [dataclasses.replace(t, **fields) if t.title == title else t for t in staff]


def _general(bm_lessons: list[int], homeroom=("3/1", "3/2"), **extra):
    """GVCN các lớp `homeroom`, đủ GV chuyên biệt, và các bộ môn với định mức cho trước."""
    rows = [(f"CN {c}", f"chủ nhiệm {c}", 19) for c in homeroom]
    rows += [("TA", "tiếng anh 1", 23), ("TD", "thể dục 1", 23), ("AN", "âm nhạc 1", 23), ("MT", "mỹ thuật 1", 23),
             ("TH", "tin học 1", 23)]
    rows += [(f"BM{i}", f"bộ môn {i}", n) for i, n in enumerate(bm_lessons, start=1)]
    return [teacher(n, t, s, row=i + 2) for i, (n, t, s) in enumerate(rows)]


def _classes_of(plan, problem, g):
    return {problem.courses[cid].class_name for (cid, h), n in plan.lessons.items() if h == g and n}


# --------------------------------------------------------------------------
# Thai sản, hợp đồng: ai dạy bù
# --------------------------------------------------------------------------
def test_maternity_homeroom_takes_no_overtime():
    # Mỗi lớp cần 4 tiết ngoài phần GVCN; bộ môn 2 tiết (+2 bù); GVCN 3/1 thai sản không bù: thiếu 8 - 2 - 2 - 2 = 2.
    staff = _set(_general([2]), "chủ nhiệm 3/1", maternity=True, campus2=True)
    problem = build_problem(staff, CURRICULUM, {}, overtime_max=2)
    assert "chủ nhiệm 3/1" not in problem.overtime
    plan = phan_cong(problem, W)
    assert plan.overtime == {"chủ nhiệm 3/2": 2, "bộ môn 1": 2} and plan.missing_total() == 2


def test_maternity_general_teaches_only_campus_two():
    staff = _set(_general([23]), "chủ nhiệm 3/2", campus2=True)
    staff = _set(staff, "bộ môn 1", maternity=True)
    problem = build_problem(staff, CURRICULUM, {}, overtime_max=2)
    assert problem.campus2 == {"3/2"} and "bộ môn 1" not in problem.overtime
    for c in problem.courses:
        assert ("bộ môn 1" in c.teachers) <= (c.class_name == "3/2")


def test_contract_homeroom_takes_overtime_first():
    # Bộ môn chỉ 4 tiết, còn thiếu 4 tiết (bù tối đa +4): GVCN hợp đồng bù hết trước khi GVCN khác bù.
    staff = _general([4])
    assert phan_cong(build_problem(staff, CURRICULUM, {}, overtime_max=4), W).overtime == \
        {"chủ nhiệm 3/1": 2, "chủ nhiệm 3/2": 2}
    staff = _set(staff, "chủ nhiệm 3/2", contract=True)
    assert phan_cong(build_problem(staff, CURRICULUM, {}, overtime_max=4), W).overtime == {"chủ nhiệm 3/2": 4}


def test_contract_general_takes_overtime_first():
    # GVCN thai sản không bù: phần thiếu do bộ môn bù, bộ môn hợp đồng trước.
    staff = _general([4, 2])
    for c in ("3/1", "3/2"):
        staff = _set(staff, f"chủ nhiệm {c}", maternity=True, campus2=True)
    assert phan_cong(build_problem(staff, CURRICULUM, {}, overtime_max=2), W).overtime == \
        {"bộ môn 1": 1, "bộ môn 2": 1}
    staff = _set(staff, "bộ môn 2", contract=True)
    assert phan_cong(build_problem(staff, CURRICULUM, {}, overtime_max=2), W).overtime == {"bộ môn 2": 2}


# --------------------------------------------------------------------------
# Giữ phân công của TKB cũ: khối trước, lớp sau
# --------------------------------------------------------------------------
@pytest.mark.parametrize("h1, h2", [({"4/1"}, {"3/1", "3/2"}), ({"3/1", "3/2"}, {"4/1"})])
def test_general_teachers_keep_their_old_grade(h1, h2):
    staff = _general([23, 23], homeroom=("3/1", "3/2", "4/1"))
    staff = _set(_set(staff, "bộ môn 1", history=frozenset(h1)), "bộ môn 2", history=frozenset(h2))
    problem = build_problem(staff, CURRICULUM, {}, overtime_max=0)
    plan = phan_cong(problem, W)
    grades = lambda g: {int(c[0]) for c in _classes_of(plan, problem, g)}  # noqa: E731
    assert grades("bộ môn 1") == {int(c[0]) for c in h1} and grades("bộ môn 2") == {int(c[0]) for c in h2}


def test_general_teachers_keep_their_old_class():
    staff = _general([23, 23])
    staff = _set(_set(staff, "bộ môn 1", history=frozenset({"3/2"})), "bộ môn 2", history=frozenset({"3/1"}))
    problem = build_problem(staff, CURRICULUM, {}, overtime_max=0)
    plan = phan_cong(problem, W)
    assert _classes_of(plan, problem, "bộ môn 1") == {"3/2"} and _classes_of(plan, problem, "bộ môn 2") == {"3/1"}


# --------------------------------------------------------------------------
# Xếp giờ: mỗi buổi một cơ sở, buổi nghỉ
# --------------------------------------------------------------------------
def _campus_staff():
    staff = _set(small_staff(), "chủ nhiệm 3/2", campus2=True)
    staff = _set(staff, "tiếng anh 1", off_sessions=frozenset({(3, "Chiều"), (4, "Sáng")}))
    return _set(staff, "chủ nhiệm 3/1", off_any=(("Chiều", 2),))


@pytest.fixture(scope="module")
def campus_solution():
    return solve(_campus_staff(), CURRICULUM, FAST, log=lambda *_: None)


def test_one_campus_per_session_and_leave_are_kept(campus_solution):
    sol = campus_solution
    assert check(sol.problem, sol.lessons) == []
    campuses, busy = defaultdict(set), defaultdict(set)
    for les in sol.lessons:
        key = (les.day, SESSION[les.period])
        campuses[les.teacher, key].add(les.class_name in sol.problem.campus2)
        busy[les.teacher].add(key)
    shared = {g for (g, _), cs in campuses.items() if cs}
    assert all(len(cs) == 1 for cs in campuses.values())
    assert {"tiếng anh 1", "thể dục 1", "bộ môn 1"} <= shared  # dạy cả hai lớp nhưng mỗi buổi một cơ sở
    assert not busy["tiếng anh 1"] & {(3, "Chiều"), (4, "Sáng")}
    afternoons = {(d, "Chiều") for d, ss in config.DAY_SESSIONS.items() if config.AFTERNOON in ss}
    assert len(afternoons - busy["chủ nhiệm 3/1"]) >= 2


def _campus_day_switches(sol) -> list[tuple[str, int]]:
    days = defaultdict(set)
    for les in sol.lessons:
        days[les.teacher, les.day].add(les.class_name in sol.problem.campus2)
    return sorted(k for k, cs in days.items() if len(cs) > 1)


def test_whole_day_at_one_campus_is_preferred(campus_solution):
    """Mục tiêu mềm campus_day_switch: không phạt thì GV tiếng anh có ngày sáng một cơ sở, chiều cơ sở kia."""
    assert _campus_day_switches(campus_solution) == []
    settings = dataclasses.replace(FAST, weights=config.Weights(campus_day_switch=0))
    free = solve(_campus_staff(), CURRICULUM, settings, log=lambda *_: None)
    assert check(free.problem, free.lessons) == [] and _campus_day_switches(free)


def test_checker_flags_campus_and_leave_violations(campus_solution):
    sol = campus_solution
    first = next(l for l in sol.lessons if l.teacher == "thể dục 1")
    key = (first.day, SESSION[first.period])
    counts = Counter(l.class_name for l in sol.lessons if l.teacher == "thể dục 1")
    teachers = dict(sol.problem.teachers)
    teachers["thể dục 1"] = dataclasses.replace(teachers["thể dục 1"], off_sessions=frozenset({key}),
                                               off_any=((None, 9),))
    teachers["tiếng anh 1"] = dataclasses.replace(teachers["tiếng anh 1"], maternity=True)
    errors = check(dataclasses.replace(sol.problem, teachers=teachers), sol.lessons)
    assert any("thể dục 1 có tiết buổi nghỉ" in e for e in errors)
    assert any("thể dục 1 xin nghỉ 9 buổi" in e for e in errors)
    assert any("tiếng anh 1 chỉ dạy ở cơ sở 2 nhưng lớp ở cơ sở 1" in e for e in errors)
    assert counts["3/1"] and counts["3/2"]
    # Chuyển một tiết cơ sở 1 của GV thể dục sang buổi GV đó đang dạy ở cơ sở 2 (đổi chỗ với tiết cùng ô).
    at2 = next(l for l in sol.lessons if l.teacher == "thể dục 1" and l.class_name == "3/2")
    same = [l for l in sol.lessons if l.class_name == "3/1" and l.day == at2.day
            and SESSION[l.period] == SESSION[at2.period] and l.teacher != "thể dục 1" and l.period != at2.period]
    other = same[0]
    moved = [dataclasses.replace(l, teacher="thể dục 1") if l is other else l for l in sol.lessons]
    assert any("thể dục 1 dạy cả hai cơ sở" in e for e in check(sol.problem, moved))


def test_timetable_class_column_is_plain_name(campus_solution, tmp_path):
    import openpyxl

    from tkb.writer import write_timetable
    out = tmp_path / "TKB.xlsx"
    write_timetable(campus_solution, out)
    titles = [c.value for c in openpyxl.load_workbook(out)["Khối 3"]["A"] if c.value not in (None, "LỚP")]
    assert titles == ["3/1", "3/2"]  # chỉ tên lớp, không thêm chữ LỚP hay CƠ SỞ 2


def test_statistics_show_campus_moves(campus_solution, tmp_path):
    """File thống kê: hai cột cho người dạy ở cả hai cơ sở (các buổi ở cơ sở 2, đổi cơ sở trong ngày)."""
    import openpyxl

    from tkb.writer import MOVE_HEADERS, write_statistics
    out = tmp_path / "Thong_Ke.xlsx"
    write_statistics(campus_solution, out)
    rows = list(openpyxl.load_workbook(out)["Thống kê"].iter_rows(values_only=True))
    assert rows[0][-2:] == MOVE_HEADERS
    both = {g for g in {les.teacher for les in campus_solution.lessons}
            if len({les.class_name in campus_solution.problem.campus2
                    for les in campus_solution.lessons if les.teacher == g}) > 1}
    codes = {campus_solution.problem.teachers[g].code for g in both}
    table = rows[1:[r[0] for r in rows].index("Tổng")]
    assert {r[1] for r in table if r[-2]} == codes and {"Tiếng Anh 1", "Thể dục 1"} <= codes
    for r in table:
        assert (r[-2] is None) == (r[1] not in codes)
        assert all(part.split()[0] in ("Sáng", "Chiều") for part in (r[-2] or "Sáng T2").split(", "))
        assert r[-1] is None  # có phạt: không ai sáng một cơ sở, chiều cơ sở kia
    total = rows[[r[0] for r in rows].index("Tổng")]
    assert total[-2:] == (f"{len(codes)} người", "0 lần")
    assert any(r[1] and str(r[1]).startswith("Dạy ở cả hai cơ sở:") for r in rows)


def test_cli_splits_timetables_by_campus(tmp_path):
    """Có lớp ở cơ sở 2: mỗi file TKB tách thành file điểm chính (cơ sở 1) và file điểm phụ (cơ sở 2)."""
    import openpyxl

    from tkb.__main__ import main
    from tkb.template import write_staff_template
    src = tmp_path / "vao.xlsx"
    write_staff_template(src, _set(small_staff(), "chủ nhiệm 3/2", campus2=True), CURRICULUM)
    out = tmp_path / "out"
    assert main([str(src), "-o", str(out / "TKB.xlsx"), "--time-limit", "10", "--workers", "4"]) == 0
    assert sorted(p.name for p in out.glob("TKB*.xlsx")) == [  # TKB giáo viên không tách theo cơ sở
        "TKB_chuc_vu_diem_chinh.xlsx", "TKB_chuc_vu_diem_phu.xlsx", "TKB_diem_chinh.xlsx", "TKB_diem_phu.xlsx",
        "TKB_giao_vien.xlsx"]
    ws = openpyxl.load_workbook(out / "TKB_giao_vien.xlsx")["Giáo viên"]
    assert any("3/2 (CS2)" in str(c.value) for row in ws.iter_rows() for c in row)  # lớp ở cơ sở 2 ghi rõ
    titles = lambda name: [c.value for c in openpyxl.load_workbook(out / name)["Khối 3"]["A"]  # noqa: E731
                           if c.value not in (None, "LỚP")]
    assert titles("TKB_diem_chinh.xlsx") == titles("TKB_chuc_vu_diem_chinh.xlsx") == ["3/1"]
    assert titles("TKB_diem_phu.xlsx") == titles("TKB_chuc_vu_diem_phu.xlsx") == ["3/2"]
