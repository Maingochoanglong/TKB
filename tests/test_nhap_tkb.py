"""Nhập TKB đang dùng dạng bảng tin (tkb/nhap_tkb.py): file TKB toàn trường của chương trình đọc lại đúng TKB đó;
TKB điền tay kiểu nhà trường (dòng tựa phía trên, tên lớp "Lớp 3/1", chữ ngày ở giữa khối, môn viết tắt, môn và tên
một dòng, tiết ngoài khung giờ) đọc được, chữ lạ ghép tay; dòng lệnh ghi file vào kèm TKB đã xếp để xếp tiếp."""
from types import SimpleNamespace

import openpyxl
import pytest

from tkb import config, khung_gio, nhap_tkb
from tkb.allocation import build_problem
from tkb.checker import check
from tkb.program import read_program
from tkb.rules import applied, read_rules
from tkb.solver import saved_lessons
from tkb.staff import InputError, parse_saved_grid, read_saved_timetable, read_staff

from .conftest import teacher


def _problem(path):
    """Bài toán của file vào `path` (gọi trong rules.applied của file đó)."""
    curriculum = read_program(path)
    staff = read_staff(path, subjects=[s for req in curriculum.values() for s in req])
    return build_problem(staff, curriculum, {}, overtime_max=4)


def _cells(rows):
    return sorted((r[0], r[1], int(r[2]), r[3], r[4]) for r in rows)


def test_school_timetable_reads_back_the_same(small_updated):
    """TKB toàn trường của chương trình (writer.write_school_timetable) nhập lại ra đúng TKB đã xếp, đúng mọi luật."""
    board = small_updated.parent / "TKB_toan_truong.xlsx"
    with applied(read_rules(small_updated)):
        problem = _problem(small_updated)
        grid, report = nhap_tkb.read_board(board, problem)
        saved = parse_saved_grid(grid)
        lessons, errors = saved_lessons(problem, saved.rows)
        assert errors == [] and check(problem, lessons) == []
    want = read_saved_timetable(small_updated).rows
    assert _cells(saved.rows) == _cells(want)
    assert report["cells"] == len(want) and report["sessions"][0] == report["sessions"][1]
    assert not report["unknown_subjects"] and not report["unknown_teachers"] and not report["outside"]


def _hand_made(path, rows, problem):
    """TKB kiểu nhà trường tự kẻ: hai dòng tựa, tên lớp "Lớp 3/1", dòng GVCN, mỗi tiết hai dòng không gộp, số tiết ở
    dòng trên, chữ ngày ("2") và buổi ("S", "C") ở giữa khối; Toán viết "T"; một ô "TA TA" (môn rồi tên một dòng), một ô
    chỉ ghi môn ở dòng dưới; thêm tiết 4 chiều Thứ 5 "Stem+" (khung giờ chỉ 3 tiết chiều)."""
    homeroom = {t.class_name: t for t in problem.teachers.values() if t.class_name}
    name = {t.code: t.name for t in problem.teachers.values()}
    at = {(r[0], r[1], int(r[2])): (r[3], r[4]) for r in rows}
    classes = sorted(problem.classes)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "THỜI KHÓA BIỂU (bản tạm)"
    ws.append([])
    ws.append(["THỨ", "BUỔI", "TIẾT", *(f"Lớp {c}" for c in classes)])
    ws.append(["GVCN", None, None, *(homeroom[c].name for c in classes)])
    odd = {"one_line": None, "below": None}
    for d in khung_gio.days():
        day = config.DAYS[d]
        for s_i, session in enumerate(khung_gio.sessions(d)):
            periods = list(session.periods) + ([None] if d == 3 and s_i == 1 else [])
            for k, p in enumerate(periods, start=1):
                top, below = [None, None, k], [None, None, None]
                if k == 2:
                    top[0] = day.split()[-1] if s_i == 0 else None
                    below[1] = session.name[0]
                for c in classes:
                    if p is None:
                        top.append("Stem+")
                        below.append(None)
                        continue
                    subject, code = at.get((c, day, p), (None, None))
                    text = "T" if subject == "Toán" else subject
                    who = name[code] if code and homeroom[c].code != code else None
                    if who == "TA" and odd["one_line"] is None:
                        odd["one_line"] = (c, d, p)
                        top.append(f"{text} {who}")
                        below.append(None)
                    elif who is None and subject and odd["below"] is None and k > 1:
                        odd["below"] = (c, d, p)
                        top.append(None)
                        below.append(text)
                    else:
                        top.append(text)
                        below.append(who)
                ws.append(top)
                ws.append(below)
    wb.save(path)
    return odd


def test_hand_made_board(small_updated, tmp_path):
    want = read_saved_timetable(small_updated).rows
    with applied(read_rules(small_updated)):
        problem = _problem(small_updated)
        board = tmp_path / "tkb_tay.xlsx"
        odd = _hand_made(board, want, problem)
        assert None not in odd.values()  # có đủ các ô lạ để thử
        grid, report = nhap_tkb.read_board(board, problem)
        # "T" chưa khớp môn nào (tên GV trường nhỏ "TA", "TD"… trùng chữ viết tắt môn: dòng dưới môn vẫn là tên GV).
        assert [u["token"] for u in report["unknown_subjects"]] == ["T"] and not report["unknown_teachers"]
        assert report["outside"] == len(problem.classes)  # tiết 4 chiều Thứ 5 (Stem+) ngoài khung giờ
        assert report["classes"] == sorted(problem.classes) and report["sessions"] == [9, 9]
        toan = sum(1 for r in want if r[3] == "Toán")
        assert report["cells"] == len(want) - toan
        grid, report = nhap_tkb.read_board(board, problem, {"subjects": {"T": "Toán"}})
        assert not report["unknown_subjects"] and report["cells"] == len(want)
        saved = parse_saved_grid(grid)
        assert _cells(saved.rows) == _cells(want)
        lessons, errors = saved_lessons(problem, saved.rows)
        assert errors == [] and check(problem, lessons) == []


def test_command_line_writes_the_saved_sheet(small_updated, tmp_path, capsys):
    """python -m tkb.nhap_tkb: ghi file vào kèm sheet TKB đã xếp; xếp tiếp trên file đó dùng lại đúng TKB."""
    from tkb.__main__ import main as tkb_main
    source = small_updated.parent / "nho.xlsx"
    with applied(read_rules(small_updated)):
        _hand_made(tmp_path / "tkb_tay.xlsx", read_saved_timetable(small_updated).rows, _problem(small_updated))
    out = tmp_path / "nho_tkb.xlsx"
    assert nhap_tkb.main([str(tmp_path / "tkb_tay.xlsx"), str(source), "-o", str(out), "--ghep", "T=Toán",
                          "--max-overtime", "4"]) == 0
    assert "Đã ghi" in capsys.readouterr().out
    assert config.SAVED_SHEET in openpyxl.load_workbook(out).sheetnames
    argv = [str(out), "-o", str(tmp_path / "ra" / "TKB.xlsx"), "--mode", "bu_gio", "--max-overtime", "4",
            "--time-limit", "20", "--workers", "4"]
    assert tkb_main(argv) == 0
    assert "Dùng lại TKB đã xếp" in capsys.readouterr().out


def test_board_errors(small_updated, tmp_path):
    with applied(read_rules(small_updated)):
        problem = _problem(small_updated)
        wb = openpyxl.Workbook()
        wb.active.append(["không có tên lớp"])
        wb.save(tmp_path / "trong.xlsx")
        with pytest.raises(InputError, match="dòng tên lớp"):
            nhap_tkb.read_board(tmp_path / "trong.xlsx", problem)
        wb = openpyxl.Workbook()
        wb.active.append(["Ghi chú", *sorted(problem.classes)])
        wb.active.append(["sáng", "Toán", "Toán"])
        wb.save(tmp_path / "khong_tiet.xlsx")
        with pytest.raises(InputError, match="cột tiết"):
            nhap_tkb.read_board(tmp_path / "khong_tiet.xlsx", problem)


def test_teacher_names_with_initials():
    """Tên viết tắt kèm chữ đầu: "N.Ly" là người họ Nguyễn tên Ly, "Võ Ly" người có chữ Võ trong họ tên, "Ly" (hai
    người) không khớp đúng một người; cột Tên Viết Tắt khớp đúng."""
    from dataclasses import replace
    teachers = {f"bộ môn {i}": teacher(n, f"bộ môn {i}", 23) for i, n in
                enumerate(["Nguyễn Thị Ly", "Võ Thị Ly", "Trần Văn Khoa"], start=1)}
    teachers["bộ môn 3"] = replace(teachers["bộ môn 3"], short_name="Thầy Khoa")
    problem = SimpleNamespace(courses=[], subject_label=lambda s: s, teachers=teachers)
    names = nhap_tkb._Names(problem, {}, None)
    assert names.teacher("N.Ly") == "bộ môn 1" and names.teacher("Võ Ly") == "bộ môn 2"
    assert names.teacher("N. Ly") == "bộ môn 1" and names.teacher("Ly") is None
    assert sorted(names.candidates("Ly")) == ["bộ môn 1", "bộ môn 2"]
    assert names.teacher("thầy khoa") == "bộ môn 3" and names.teacher("Khoa") == "bộ môn 3"
