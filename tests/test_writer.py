import openpyxl

from tkb import config
from tkb.checker import check
from tkb.solver import solve
from tkb.staff import read_staff
from tkb import writer
from tkb.style import Style
from tkb.writer import write_timetable, write_updated_staff

from .conftest import CURRICULUM, INPUT_FILE, small_staff, teacher

STYLE = Style.from_file(INPUT_FILE)  # style của file vào: Times New Roman 14, viền mảnh, dòng cao 25


def _solve_small(general):
    return solve(small_staff(general), CURRICULUM, config.Settings(time_limit=20, workers=4), log=lambda *_: None)


def test_style_is_read_from_input_file():
    assert STYLE.header.font.name == "Times New Roman" and STYLE.header.font.sz == 14 and STYLE.header.font.b
    assert STYLE.header.fill is None and STYLE.body.border.left.style == "thin"
    assert STYLE.row_height == 25 and STYLE.line_height == 21


def test_timetable_layout(tmp_path):
    sol = _solve_small(general=True)
    out = tmp_path / "TKB.xlsx"
    write_timetable(sol, out, STYLE)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Khối 3"]  # file TKB chỉ có thời khóa biểu
    ws = wb["Khối 3"]
    # Bố cục như data/Output_Template_TKB_V8.xlsx: mỗi lớp 1 dòng tiêu đề + 7 tiết + 2 dòng trống.
    assert [ws.cell(1, c).value for c in range(1, 9)] == [
        "LỚP", "BUỔI", "TIẾT", "THỨ 2", "THỨ 3", "THỨ 4", "THỨ 5", "THỨ 6"]
    assert {str(r) for r in ws.merged_cells.ranges} == {
        "A2:A8", "B2:B5", "B6:B8", "A12:A18", "B12:B15", "B16:B18"}
    assert ws["A2"].value == "LỚP 3/1" and ws["A12"].value == "LỚP 3/2"
    assert ws["A11"].value == "LỚP"
    assert ws["B2"].value == "SÁNG" and ws["B6"].value == "CHIỀU"
    assert [ws.cell(r, 3).value for r in range(2, 9)] == [1, 2, 3, 4, 5, 6, 7]  # tiết trong ngày: chiều là 5, 6, 7
    assert ws["D2"].value == "HĐTN\nCN A"  # môn, xuống dòng tên giáo viên
    assert ws["H5"].value == "HĐTN\nCN A"
    assert [ws.cell(r, 8).value for r in (6, 7, 8)] == ["Nghỉ"] * 3
    cells = [ws.cell(r, c).value for r in range(2, 9) for c in range(4, 9)]
    assert all(v for v in cells) and not any("thiếu" in v for v in cells)
    # Style chép từ file vào.
    for cell in (ws["A1"], ws["A2"], ws["C2"], ws["D2"]):
        assert cell.font.name == "Times New Roman" and cell.font.sz == 14
    assert ws["A1"].font.b and ws["A1"].fill.fill_type is None and not ws["D2"].font.b
    assert ws.row_dimensions[1].height == 25
    assert ws.row_dimensions[2].height == 2 * STYLE.line_height  # môn + tên giáo viên
    assert ws["A5"].border.left.style == "thin"  # viền cả ô nằm trong vùng gộp
    assert ws.column_dimensions["D"].width >= STYLE.text_width("TV tăng cường")
    # Cột BUỔI, TIẾT rộng hơn chữ dài nhất (CHIỀU, TIẾT) để dễ nhìn.
    assert ws.column_dimensions["B"].width >= STYLE.text_width("CHIỀU") + writer.LABEL_PAD - 0.1
    assert ws.column_dimensions["C"].width >= STYLE.text_width("TIẾT") + writer.LABEL_PAD - 0.1


def _values(ws):
    return [v for row in ws.iter_rows(values_only=True) for v in row if v is not None]


def _stats(path):
    """Các dòng của bảng thống kê, đến dòng Tổng (bỏ phần chú thích màu bên dưới)."""
    rows = list(openpyxl.load_workbook(path)["Thống kê"].iter_rows(values_only=True))
    return rows[:[r[0] for r in rows].index("Tổng") + 1]


def _fills(path):
    """Màu nền cột A của từng dòng bảng thống kê (None: không tô) và các dòng chú thích dưới bảng."""
    ws = openpyxl.load_workbook(path)["Thống kê"]
    n = len(_stats(path))
    color = lambda c: c.fill.start_color.rgb[-6:] if c.fill.fill_type else None
    legend = [(color(ws.cell(r, 1)), ws.cell(r, 2).value) for r in range(n + 1, ws.max_row + 1) if ws.cell(r, 2).value]
    return [color(ws.cell(r, 1)) for r in range(2, n + 1)], legend


def test_statistics_file_is_one_table(tmp_path):
    sol = _solve_small(general=True)
    out = tmp_path / "Thong_Ke.xlsx"
    writer.write_statistics(sol, out, STYLE)
    assert openpyxl.load_workbook(out).sheetnames == ["Thống kê"]
    rows = _stats(out)
    header = rows[0]
    # Tên, chức vụ (Mã GV), số tiết từng môn (chỉ các môn có người dạy), tổng tiết.
    assert header[:2] == ("Họ và Tên", "Chức Vụ") and header[-1] == "Tổng Tiết"
    assert "HĐTN" in header and "Tiếng Việt" in header and "Mã GV" not in header
    load = sol.teacher_load()
    for r in rows[1:-1]:
        assert sum(v or 0 for v in r[2:-1]) == r[-1]
    assert rows[1][:2] == ("CN A", "Chủ Nhiệm 3/1") and rows[1][-1] == load["chủ nhiệm 3/1"]
    assert rows[1][header.index("HĐTN")] == 3 and rows[1][header.index("Tiếng Anh")] is None  # ô trống: không dạy
    assert rows[-1][0] == "Tổng" and rows[-1][-1] == 64  # 2 lớp × 32 tiết


def test_supplement_in_statistics(tmp_path):
    sol = _solve_small(general=False)
    out = tmp_path / "TKB.xlsx"
    stats = tmp_path / "Thong_Ke.xlsx"
    write_timetable(sol, out, STYLE)
    writer.write_statistics(sol, stats, STYLE)
    rows = _stats(stats)
    # Người cần tuyển: tên "tuyển thêm", chức vụ là Mã GV, thực dạy 8 tiết.
    assert rows[-2][:2] == ("tuyển thêm", "Bộ Môn 1") and rows[-2][-1] == 8
    ws = openpyxl.load_workbook(stats)["Thống kê"]
    assert ws["A1"].font.b and ws["A2"].font.sz == 14 and ws.row_dimensions[2].height == 25
    assert ws["A2"].font.name == "Times New Roman" and ws["A2"].border.left.style == "thin"
    assert ws.cell(len(rows), 1).font.b and ws.freeze_panes == "C2"
    grid = [v for row in openpyxl.load_workbook(out)["Khối 3"].iter_rows(values_only=True) for v in row if v]
    assert any(isinstance(v, str) and v.endswith("\nBộ Môn 1") for v in grid)
    # Dòng người cần tuyển tô xanh, chú thích dưới bảng; không ai dạy bù.
    fills, legend = _fills(stats)
    assert fills == [None] * (len(rows) - 3) + [writer.HIRE_FILL, None]
    assert legend == [(writer.HIRE_FILL, "Cần tuyển thêm: 1 người, 8 tiết")]


def test_updated_staff_file_is_reusable(tmp_path):
    src = tmp_path / "staff.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Họ và Tên", "Chức Vụ", "Lớp", "Số Tiết/Tuần"])  # sheet nhân sự mẫu V8, không có style
    for t in small_staff(general=False):
        ws.append([t.name, t.role.title(), t.class_name, t.max_lessons])
    wb.save(src)
    sol = solve(read_staff(src), CURRICULUM, config.Settings(time_limit=20, workers=4), log=lambda *_: None)
    dst = tmp_path / "staff_cap_nhat.xlsx"
    write_updated_staff(sol, src, dst)
    rows = list(openpyxl.load_workbook(dst).active.iter_rows(values_only=True))
    assert rows[0] == ("Họ và Tên", "Chức Vụ", "Lớp", "Số Tiết/Tuần", "Mã GV", "Số Tiết Thực Dạy")
    assert rows[-1] == ("chưa có", "Bộ Môn", None, 23, "Bộ Môn 1", 8)
    again = read_staff(dst)
    assert again[-1].name == "chưa có" and again[-1].title == "bộ môn 1"
    assert again[-1].max_lessons == 23
    sol2 = solve(again, CURRICULUM, config.Settings(time_limit=20, workers=4), log=lambda *_: None)
    assert sol2.used_supplements() == []
    assert check(sol2.problem, sol2.lessons) == []


def test_teacher_labels():
    teachers = {t.title: t for t in [teacher("Lan", "chủ nhiệm 1/1", 19), teacher("Lan", "bộ môn 1", 23),
                                     teacher("", "bộ môn 2", 23), teacher("Hoa", "bộ môn 3", 23)]}
    teachers["bộ môn 2"].label = "Bộ Môn"
    labels = writer.teacher_labels(teachers)
    assert labels == {"chủ nhiệm 1/1": "Lan (chủ nhiệm 1/1)", "bộ môn 1": "Lan (bộ môn 1)",
                      "bộ môn 2": "Bộ Môn 2", "bộ môn 3": "Hoa"}


def test_blank_names_show_teacher_code(tmp_path):
    staff = small_staff()
    for t in staff:
        t.name = ""  # file vào để trống tên (bảo mật)
    sol = solve(staff, CURRICULUM, config.Settings(time_limit=20, workers=4), log=lambda *_: None)
    out = tmp_path / "TKB.xlsx"
    write_timetable(sol, out, STYLE)
    ws = openpyxl.load_workbook(out)["Khối 3"]
    assert ws["D2"].value == "HĐTN\nChủ Nhiệm 3/1"
    grid = {v.split("\n")[1] for row in ws.iter_rows(min_row=2, max_row=8, min_col=4, values_only=True)
            for v in row if v and "\n" in v}
    assert {"Tiếng Anh 1", "Thể dục 1"} <= grid  # chức vụ chuyên biệt ghi theo tên môn


def test_timetable_with_codes(tmp_path):
    sol = solve(small_staff(), CURRICULUM, config.Settings(time_limit=20, workers=4), log=lambda *_: None)
    plain, coded = tmp_path / "TKB.xlsx", tmp_path / "TKB_chuc_vu.xlsx"
    write_timetable(sol, plain, STYLE)
    write_timetable(sol, coded, STYLE, with_codes=True)
    a, b = openpyxl.load_workbook(plain)["Khối 3"], openpyxl.load_workbook(coded)["Khối 3"]
    codes = {t.title: t.code for t in sol.problem.teachers.values()}
    names = writer.teacher_labels(sol.problem.teachers)
    by_label = {names[g]: codes[g] for g in names}
    for row_a, row_b in zip(a.iter_rows(min_col=4, values_only=True), b.iter_rows(min_col=4, values_only=True)):
        for va, vb in zip(row_a, row_b):
            if va and "\n" in str(va):
                subject, label = va.split("\n")
                assert vb == f"{va}\n{by_label[label]}"
    assert b.row_dimensions[2].height > a.row_dimensions[2].height  # 3 dòng chữ


def test_shortage_file(tmp_path):
    out = tmp_path / "Thong_Ke.xlsx"
    writer.write_shortage([("3/1", "KNS", 1, "GVCN và bộ môn đã bù tối đa +2 tiết"),
                           ("3/2", "KNS", 1, "GVCN và bộ môn đã bù tối đa +2 tiết")], out, STYLE)
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Thiếu tiết"]
    rows = [r for r in wb["Thiếu tiết"].iter_rows(values_only=True) if any(r)]
    assert rows[0] == ("Lớp", "Môn", "Số Tiết Thiếu", "Lý Do") and rows[-1][:3] == ("Tổng", None, 2)


def test_long_names_widen_columns_and_rows(tmp_path):
    staff = small_staff()
    staff[0].name = "Nguyễn Thị Thanh Hương Giang Mai"  # GVCN 3/1, tên rất dài
    sol = solve(staff, CURRICULUM, config.Settings(time_limit=20, workers=4), log=lambda *_: None)
    out = tmp_path / "TKB.xlsx"
    write_timetable(sol, out, STYLE)
    ws = openpyxl.load_workbook(out)["Khối 3"]
    assert ws["D2"].value == "HĐTN\nNguyễn Thị Thanh Hương Giang Mai"
    assert ws.column_dimensions["D"].width == writer.MAX_DAY_WIDTH  # nới hết mức
    assert ws.row_dimensions[2].height > 2 * STYLE.line_height  # tên xuống dòng, hàng cao thêm


def test_statistics_file(tmp_path):
    out = tmp_path / "Thong_Ke.xlsx"
    writer.write_statistics(_solve_small(general=False), out, STYLE)
    rows = _stats(out)
    header = rows[0]
    assert rows[1][:2] == ("CN A", "Chủ Nhiệm 3/1") and rows[1][-1] == 19
    assert [r[1] for r in rows[1:-1]] == ["Chủ Nhiệm 3/1", "Chủ Nhiệm 3/2", "Tiếng Anh 1", "Thể dục 1", "Âm nhạc 1",
                                          "Mỹ thuật 1", "Tin học 1", "Bộ Môn 1"]
    ta = rows[3]  # GV Tiếng Anh chỉ dạy Tiếng Anh, 4 tiết mỗi lớp
    assert ta[header.index("Tiếng Anh")] == 8 and ta[-1] == 8
    assert rows[-1][0] == "Tổng" and rows[-1][header.index("Tiếng Anh")] == 8 and rows[-1][-1] == 64


def test_statistics_file_overtime(tmp_path):
    settings = config.Settings(time_limit=20, workers=4, mode=config.MODE_OVERTIME, overtime_max=4)
    sol = solve(small_staff(general=False), CURRICULUM, settings, log=lambda *_: None)
    out = tmp_path / "Thong_Ke.xlsx"
    writer.write_statistics(sol, out, STYLE)
    rows = _stats(out)
    assert rows[1][:2] == ("CN A", "Chủ Nhiệm 3/1") and rows[1][-1] == 23  # 19 tiết + 4 tiết bù
    assert all(r[0] != "tuyển thêm" for r in rows)
    assert rows[-1][0] == "Tổng" and rows[-1][-1] == 64
    # Dòng GVCN dạy bù tô vàng (cả dòng), chú thích dưới bảng.
    fills, legend = _fills(out)
    assert fills == [writer.OVERTIME_FILL] * 2 + [None] * (len(rows) - 3)
    ws = openpyxl.load_workbook(out)["Thống kê"]
    assert {ws.cell(2, c).fill.start_color.rgb[-6:] for c in range(1, len(rows[0]) + 1)} == {writer.OVERTIME_FILL}
    assert legend == [(writer.OVERTIME_FILL, "Dạy bù (vượt định mức): 2 người, 8 tiết")]
