import openpyxl

import main
from tkb import writer
from tkb.template import write_staff_template

from .conftest import CURRICULUM, small_staff


def _write_staff(path, general=True):
    """File vào mẫu V8: sheet NHÂN SỰ + CHƯƠNG TRÌNH HỌC."""
    write_staff_template(path, small_staff(general), CURRICULUM)


def _rows(path, sheet="NHÂN SỰ"):
    ws = openpyxl.load_workbook(path)[sheet]
    return [r for r in ws.iter_rows(values_only=True) if any(v is not None for v in r)]


def test_run_writes_outputs(tmp_path):
    in_dir, out_dir = tmp_path / "in", tmp_path / "out"
    in_dir.mkdir()
    _write_staff(in_dir / "nhan_su.xlsx", general=False)
    code = main.run(in_dir / "nhan_su.xlsx", out_dir, "TKB.xlsx", thoi_gian_toi_da=20,
                    che_do="tuyen_them")
    assert code == 0
    # TKB.xlsx chỉ có thời khóa biểu; nhân sự và thống kê nằm ở Thong_Ke.xlsx.
    assert openpyxl.load_workbook(out_dir / "TKB.xlsx").sheetnames == ["Khối 3"]
    assert openpyxl.load_workbook(out_dir / "Thong_Ke.xlsx").sheetnames == ["Thống kê"]
    rows = _rows(out_dir / "nhan_su_cap_nhat.xlsx")
    assert rows[-1] == ("chưa có", "Bộ Môn", None, 23, *(None,) * 5, "Bộ Môn 1", 8, 15)  # 5 cột không bắt buộc trống
    stats = _rows(out_dir / "Thong_Ke.xlsx", "Thống kê")
    total = [r[0] for r in stats].index("Tổng")  # dưới dòng Tổng là chú thích màu
    assert stats[total - 1][:2] == ("tuyển thêm", "Bộ Môn 1") and stats[total - 1][-4:] == (8, 23, None, 15)
    assert stats[total + 1][1] == "Cần tuyển thêm: 1 người, 8 tiết"
    # Các file ra dùng style của file vào.
    tkb = openpyxl.load_workbook(out_dir / "TKB.xlsx")["Khối 3"]
    assert tkb["D2"].font.name == "Times New Roman" and tkb["D2"].font.sz == 14
    # TKB có chức vụ: cùng TKB, mỗi ô thêm dòng Mã GV.
    roles = openpyxl.load_workbook(out_dir / "TKB_chuc_vu.xlsx")["Khối 3"]
    assert roles["D2"].value == tkb["D2"].value + "\nChủ Nhiệm 3/1"


def test_overtime_shortage_is_an_error_with_a_table(tmp_path, capsys):
    in_dir, out_dir = tmp_path / "in", tmp_path / "out"
    in_dir.mkdir()
    _write_staff(in_dir / "nhan_su.xlsx", general=False)
    code = main.run(in_dir / "nhan_su.xlsx", out_dir, "TKB.xlsx", thoi_gian_toi_da=20, che_do="bu_gio",
                    so_tiet_bu_toi_da=2)
    assert code == 3
    assert "thiếu 4 tiết" in capsys.readouterr().err
    assert not (out_dir / "TKB.xlsx").exists()
    rows = _rows(out_dir / "Thong_Ke.xlsx", "Thiếu tiết")
    assert rows[0] == ("Lớp", "Môn", "Số Tiết Thiếu", "Lý Do")
    assert rows[-1][:3] == ("Tổng", None, 4)


def test_run_reports_missing_file(tmp_path, capsys):
    (tmp_path / "in").mkdir()
    (tmp_path / "in" / "khac.xlsx").write_bytes(b"")
    assert main.run(tmp_path / "in" / "khong_co.xlsx", tmp_path / "out", "TKB.xlsx") == 1
    out = capsys.readouterr().out
    assert "không tìm thấy file vào" in out and "khac.xlsx" in out


def test_relative_paths_resolve_from_script_dir():
    assert main._resolve("data") == main.BASE_DIR / "data"
    assert main._resolve(main.FILE_VAO).is_file()


def test_run_rejects_bad_thread_count(tmp_path, capsys):
    in_dir = tmp_path / "in"
    in_dir.mkdir()
    _write_staff(in_dir / "nhan_su.xlsx")
    assert main.run(in_dir / "nhan_su.xlsx", tmp_path / "out", "TKB.xlsx", so_luong=0) == 1
    assert "SO_LUONG" in capsys.readouterr().out


def test_run_passes_thread_count(tmp_path, monkeypatch):
    in_dir = tmp_path / "in"
    in_dir.mkdir()
    _write_staff(in_dir / "nhan_su.xlsx")
    seen = {}
    import tkb.__main__ as cli
    monkeypatch.setattr(cli, "main", lambda argv: seen.setdefault("argv", argv) and 0)
    main.run(in_dir / "nhan_su.xlsx", tmp_path / "out", "TKB.xlsx", so_luong=4)
    argv = seen["argv"]
    assert argv[argv.index("--workers") + 1] == "4"


def test_run_overtime_mode_needs_no_hire(tmp_path):
    in_dir, out_dir = tmp_path / "in", tmp_path / "out"
    in_dir.mkdir()
    _write_staff(in_dir / "nhan_su.xlsx", general=False)
    code = main.run(in_dir / "nhan_su.xlsx", out_dir, "TKB.xlsx", thoi_gian_toi_da=20,
                    che_do="bu_gio", so_tiet_bu_toi_da=4)
    assert code == 0
    rows = _rows(out_dir / "nhan_su_cap_nhat.xlsx")
    assert all(r[0] != "chưa có" for r in rows)
    assert rows[0][-4:] == ("Mã GV", "Số Tiết Thực Dạy", "Số Tiết Bù", "Số Tiết Dư")
    assert rows[1][-4:] == ("Chủ Nhiệm 3/1", 23, 4, None)
    # Bản thống kê gọn theo mẫu file vào: tô cả dòng người dạy bù (vàng) và người còn dư tiết (xanh dương).
    ws = openpyxl.load_workbook(out_dir / "nhan_su_cap_nhat.xlsx").active
    color = lambda cell: cell.fill.start_color.rgb[-6:] if cell.fill.fill_type else None  # noqa: E731
    for r, row in enumerate(rows[1:], start=2):
        colors = {color(ws.cell(r, c)) for c in range(1, len(row) + 1)}
        assert colors == {writer.OVERTIME_FILL if row[-2] else writer.SPARE_FILL if row[-1] else None}
    assert any(row[-1] for row in rows[1:])
    # Giải thích cột kết quả và màu bằng chữ thường ở sheet HƯỚNG DẪN; mọi file ra chỉ có chữ, số và màu: không ghi
    # chú (comment), công thức, cố định dòng/cột, danh sách thả xuống hay định dạng theo điều kiện.
    wb = openpyxl.load_workbook(out_dir / "nhan_su_cap_nhat.xlsx")
    assert wb.sheetnames == ["NHÂN SỰ", "CHƯƠNG TRÌNH HỌC", "QUY ĐỊNH", "LUẬT RIÊNG", "HƯỚNG DẪN", "TKB đã xếp"]
    notes = [c.value or "" for c in wb["HƯỚNG DẪN"]["B"]]
    assert any("xanh dương" in n for n in notes) and any("số tiết dạy vượt định mức" in n for n in notes)
    for path in out_dir.glob("*.xlsx"):
        for sheet in openpyxl.load_workbook(path).worksheets:
            cells = [c for row in sheet.iter_rows() for c in row]
            assert not any(c.comment for c in cells) and not any(c.data_type == "f" for c in cells), path.name
            assert sheet.freeze_panes is None and not sheet.data_validations.dataValidation, path.name
            assert not len(sheet.conditional_formatting), path.name


def test_run_rejects_bad_mode(tmp_path, capsys):
    in_dir = tmp_path / "in"
    in_dir.mkdir()
    _write_staff(in_dir / "nhan_su.xlsx")
    assert main.run(in_dir / "nhan_su.xlsx", tmp_path / "out", "TKB.xlsx", che_do="khac") == 1
    assert "CHE_DO" in capsys.readouterr().out
    assert main.run(in_dir / "nhan_su.xlsx", tmp_path / "out", "TKB.xlsx", so_tiet_bu_toi_da=-1) == 1
    assert "SO_TIET_BU_TOI_DA" in capsys.readouterr().out


def test_run_single_input_file_with_program_sheet(tmp_path, capsys):
    src = tmp_path / "input.xlsx"
    _write_staff(src, general=False)  # có sheet NHÂN SỰ và CHƯƠNG TRÌNH HỌC
    code = main.run(src, tmp_path / "out", "TKB.xlsx", thoi_gian_toi_da=20, che_do="tuyen_them")
    assert code == 0
    assert "chương trình học: 16 môn (sheet CHƯƠNG TRÌNH HỌC)" in capsys.readouterr().out
    wb = openpyxl.load_workbook(tmp_path / "out" / "input_cap_nhat.xlsx")
    assert wb.sheetnames[:2] == ["NHÂN SỰ", "CHƯƠNG TRÌNH HỌC"]  # file cập nhật vẫn là file vào đầy đủ


def test_run_without_program_sheet_fails(tmp_path, capsys):
    src = tmp_path / "nhan_su.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "NHÂN SỰ"
    wb.active.append(["Họ và Tên", "Chức Vụ", "Lớp", "Số Tiết/Tuần"])
    wb.active.append(["A", "Chủ Nhiệm", "3/1", 19])
    wb.save(src)
    assert main.run(src, tmp_path / "out", "TKB.xlsx", thoi_gian_toi_da=5) == 1
    assert "thiếu sheet CHƯƠNG TRÌNH HỌC" in capsys.readouterr().err


def _argv(monkeypatch, **kwargs):
    """Tham số main.run truyền cho tkb (không giải thật)."""
    seen = {}
    import tkb.__main__ as cli
    monkeypatch.setattr(cli, "main", lambda argv: seen.setdefault("argv", argv) and 0)
    main.run(**kwargs)
    return seen.get("argv")


def test_defaults_are_overtime_student_rules_1200s_reproducible():
    assert main.CHE_DO == "bu_gio" and main.LUAT_HOC_SINH is True
    assert main.THOI_GIAN_TOI_DA == 1200 and main.CHAY_TAI_LAP_DUOC is True


def test_blank_output_folder_means_project_folder(tmp_path, monkeypatch):
    _write_staff(tmp_path / "in.xlsx")
    for blank in ("", None, "  "):
        argv = _argv(monkeypatch, file_vao=tmp_path / "in.xlsx", thu_muc_out=blank)
        assert argv[argv.index("-o") + 1] == str(main.BASE_DIR / "TKB.xlsx")
        assert argv[argv.index("--stats-out") + 1] == str(main.BASE_DIR / "Thong_Ke.xlsx")


def test_blank_time_limit_means_unlimited(tmp_path, monkeypatch):
    _write_staff(tmp_path / "in.xlsx")
    for blank in (None, "", 0):
        argv = _argv(monkeypatch, file_vao=tmp_path / "in.xlsx", thu_muc_out=tmp_path, thoi_gian_toi_da=blank)
        assert argv[argv.index("--time-limit") + 1] == "0"  # 0 = không giới hạn
    argv = _argv(monkeypatch, file_vao=tmp_path / "in.xlsx", thu_muc_out=tmp_path)
    assert argv[argv.index("--time-limit") + 1] == "1200" and "--non-reproducible" not in argv
    assert "--no-student-rules" not in argv and argv[argv.index("--mode") + 1] == "bu_gio"


def test_bad_time_limit_and_blank_input_are_rejected(tmp_path, capsys):
    _write_staff(tmp_path / "in.xlsx")
    assert main.run(tmp_path / "in.xlsx", tmp_path / "out", thoi_gian_toi_da=-5) == 1
    assert "THOI_GIAN_TOI_DA" in capsys.readouterr().out
    assert main.run("", tmp_path / "out") == 1
    assert "chưa điền FILE_VAO" in capsys.readouterr().out


def test_cli_time_limit_zero_is_unlimited():
    from tkb.solver import _configure
    from ortools.sat.python import cp_model
    from tkb import config
    solver = cp_model.CpSolver()
    _configure(solver, config.Settings(time_limit=None), None)
    p = solver.parameters
    assert p.interleave_search and p.max_deterministic_time > 1e10 and p.max_time_in_seconds > 1e10


def _code(out: str) -> str:
    return next(line.split()[3] for line in out.splitlines() if line.startswith("Mã kết quả:"))


def _subjects(path):
    """Môn ở từng ô TKB (dòng đầu của ô), không tính tên người dạy."""
    ws = openpyxl.load_workbook(path)["Khối 3"]
    return [[str(v).split("\n")[0] if v else v for v in row] for row in ws.iter_rows(values_only=True)]


def test_reloading_updated_file_keeps_timetable(tmp_path, capsys):
    # Chạy chế độ tuyển; đổi "chưa có" thành tên người mới trong file vào cập nhật rồi chạy lại: TKB giữ nguyên.
    in_dir, out_dir = tmp_path / "in", tmp_path / "out"
    in_dir.mkdir()
    _write_staff(in_dir / "nhan_su.xlsx", general=False)
    assert main.run(in_dir / "nhan_su.xlsx", out_dir, "TKB.xlsx", thoi_gian_toi_da=20, che_do="tuyen_them") == 0
    first = _code(capsys.readouterr().out)
    # TKB đã xếp dạng lưới như TKB: dòng đầu ghi mã kết quả, mỗi ô "môn" xuống dòng "Mã GV".
    saved = _rows(out_dir / "nhan_su_cap_nhat.xlsx", "TKB đã xếp")
    assert saved[0][:2] == ("Mã kết quả", first) and saved[0][2] == "Mã quy định"
    assert saved[1] == ("Lớp", "Tiết", "Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6")
    assert len(saved) == 2 + 2 * 7 and saved[2][:3] == ("3/1", 1, "HĐTN\nChủ Nhiệm 3/1")
    assert saved[6][-1] == "Nghỉ"  # chiều Thứ 6
    wb = openpyxl.load_workbook(out_dir / "nhan_su_cap_nhat.xlsx")
    ws = wb["NHÂN SỰ"]
    hire = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "chưa có")
    ws.cell(hire, 1, "Cô Mới")
    wb.save(in_dir / "da_tuyen.xlsx")
    for mode, folder in (("tuyen_them", "lai"), ("bu_gio", "lai_bu")):
        assert main.run(in_dir / "da_tuyen.xlsx", tmp_path / folder, "TKB.xlsx", thoi_gian_toi_da=20,
                        che_do=mode, so_tiet_bu_toi_da=4) == 0
        out = capsys.readouterr().out
        assert "Dùng lại TKB đã xếp" in out and _code(out) == first, mode
        assert _subjects(tmp_path / folder / "TKB.xlsx") == _subjects(out_dir / "TKB.xlsx")
        cells = [v for row in openpyxl.load_workbook(tmp_path / folder / "TKB.xlsx")["Khối 3"].iter_rows(values_only=True)
                 for v in row if v]
        assert any(str(v).endswith("\nCô Mới") for v in cells)
        stats = _rows(tmp_path / folder / "Thong_Ke.xlsx", "Thống kê")
        assert any(r[:2] == ("Cô Mới", "Bộ Môn 1") for r in stats) and all(r[0] != "tuyển thêm" for r in stats)


def test_reloading_falls_back_when_saved_timetable_breaks_rules(tmp_path, capsys):
    in_dir, out_dir = tmp_path / "in", tmp_path / "out"
    in_dir.mkdir()
    _write_staff(in_dir / "nhan_su.xlsx", general=False)
    assert main.run(in_dir / "nhan_su.xlsx", out_dir, "TKB.xlsx", thoi_gian_toi_da=20, che_do="tuyen_them") == 0
    capsys.readouterr()
    wb = openpyxl.load_workbook(out_dir / "nhan_su_cap_nhat.xlsx")
    ws = wb["NHÂN SỰ"]
    hire = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "chưa có")
    ws.cell(hire, 4, 5)  # định mức 5 < 8 tiết đang dạy trong TKB đã xếp
    wb.save(in_dir / "sua.xlsx")
    assert main.run(in_dir / "sua.xlsx", tmp_path / "lai", "TKB.xlsx", thoi_gian_toi_da=20,
                    che_do="tuyen_them") == 0
    out = capsys.readouterr().out
    assert "Không dùng lại được TKB đã xếp" in out and "vượt định mức" in out
    assert "kiểm tra luật bắt buộc: ĐẠT" in out
    # Muốn xếp lại từ đầu dù TKB đã xếp vẫn đúng luật: GIU_TKB_DA_XEP = False.
    assert main.run(out_dir / "nhan_su_cap_nhat.xlsx", tmp_path / "lai2", "TKB.xlsx", thoi_gian_toi_da=20,
                    che_do="tuyen_them", giu_tkb_da_xep=False) == 0
    out = capsys.readouterr().out
    assert "Dùng lại TKB đã xếp" not in out and "Không dùng lại" not in out


def test_reloading_overtime_result_rewrites_the_same_files(tmp_path, capsys):
    # Nạp lại nguyên file vào cập nhật của chế độ bù: TKB, tiết bù (ô tô cam), màu dòng và file thống kê y hệt.
    # File vào cập nhật có dòng tô màu bù/dư: style chép từ file vào không được mang theo các màu đó.
    in_dir = tmp_path / "in"
    in_dir.mkdir()
    _write_staff(in_dir / "nhan_su.xlsx", general=False)
    for src, out in ((in_dir / "nhan_su.xlsx", "a"), (tmp_path / "a" / "nhan_su_cap_nhat.xlsx", "b")):
        assert main.run(src, tmp_path / out, "TKB.xlsx", thoi_gian_toi_da=20, che_do="bu_gio",
                        so_tiet_bu_toi_da=4) == 0
    assert "Dùng lại TKB đã xếp" in capsys.readouterr().out

    def cells(path, sheet=None):
        wb = openpyxl.load_workbook(path)
        ws = wb[sheet] if sheet else wb.active
        return [[(c.value, c.fill.start_color.rgb if c.fill.fill_type else None) for c in row]
                for row in ws.iter_rows()]
    for name in ("TKB.xlsx", "TKB_chuc_vu.xlsx", "Thong_Ke.xlsx"):
        assert cells(tmp_path / "a" / name) == cells(tmp_path / "b" / name), name
    assert cells(tmp_path / "a" / "nhan_su_cap_nhat.xlsx", "NHÂN SỰ") == \
        cells(tmp_path / "b" / "nhan_su_cap_nhat_cap_nhat.xlsx", "NHÂN SỰ")
