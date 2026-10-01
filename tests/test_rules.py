"""Quy định trong file vào (tkb/rules.py): các cột quy định của sheet CHƯƠNG TRÌNH HỌC và ba bảng của sheet QUY ĐỊNH;
đọc, kiểm tra, dùng thay giá trị mặc định trong tkb/config.py."""
import shutil

import openpyxl
import pytest

import main
from tkb import config
from tkb.rules import ATTRS, DEFAULTS, LABELS, applied, changed, code, read_rules
from tkb.staff import InputError
from tkb.template import write_staff_template

from .conftest import CURRICULUM, small_staff

MON, CHUNG, NGAY, TIET = "MÔN", "Quy định", "Ngày", "Tiết"  # nơi sửa: sheet chương trình học, ba bảng


def _locate(ws, table):
    """Dòng tiêu đề và dòng cuối của bảng `table` (ô đầu dòng tiêu đề) trong sheet QUY ĐỊNH."""
    top = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == table)
    end = top
    while end + 1 <= ws.max_row and ws.cell(end + 1, 1).value is not None:
        end += 1
    return top, end


def _input(tmp_path, edits=(), name="vao.xlsx"):
    """File vào trường nhỏ (quy định ghi mặc định). `edits`: (nơi, ô đầu dòng, tiêu đề cột, giá trị); nơi là MON
    (sheet CHƯƠNG TRÌNH HỌC) hoặc tên bảng của sheet QUY ĐỊNH; bảng chung ghi vào cột Giá trị. Dòng chưa có thì thêm
    (bảng của sheet QUY ĐỊNH: thêm vào cuối bảng), cột chưa có thì thêm."""
    path = tmp_path / name
    write_staff_template(path, small_staff(general=False), CURRICULUM)
    wb = openpyxl.load_workbook(path)
    for where, key, header, value in edits:
        ws = wb[config.PROGRAM_SHEET] if where == MON else wb[config.RULES_SHEET]
        top, end = (1, ws.max_row) if where == MON else _locate(ws, where)
        heads = {ws.cell(top, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(top, c).value}
        header = "Giá trị" if where == CHUNG else header
        if header not in heads:
            heads[header] = max(heads.values()) + 1
            ws.cell(top, heads[header], header)
        row = next((r for r in range(top + 1, end + 1) if ws.cell(r, 1).value == key), None)
        if row is None:
            if where != MON and end < ws.max_row:
                ws.insert_rows(end + 1)
            row = end + 1
            ws.cell(row, 1, key)
        ws.cell(row, heads[header]).value = value
    wb.save(path)
    return path


def _drop(path, where, header=None):
    """Xóa sheet QUY ĐỊNH (where = None), hoặc một cột quy định của sheet CHƯƠNG TRÌNH HỌC."""
    wb = openpyxl.load_workbook(path)
    if where is None:
        del wb[config.RULES_SHEET]
    else:
        ws = wb[config.PROGRAM_SHEET]
        ws.delete_cols(next(c.column for c in ws[1] if c.value == header))
    wb.save(path)


def _drop_all_rule_columns(path):
    wb = openpyxl.load_workbook(path)
    ws = wb[config.PROGRAM_SHEET]
    keep = sum(1 for c in ws[1] if c.value == "Môn học" or str(c.value).startswith("Khối"))
    ws.delete_cols(keep + 1, ws.max_column - keep)
    del wb[config.RULES_SHEET]
    wb.save(path)


def test_default_rules_equal_config(tmp_path):
    rules = read_rules(_input(tmp_path))
    assert rules == DEFAULTS and changed(rules) == []
    # Cùng thứ tự (thứ tự dựng mô hình): nhóm môn, môn GVCN, ô HĐTN.
    for attr in ("SUBJECT_GROUPS", "HOMEROOM_PRIORITY", "HOMEROOM_CUT_ORDER", "HOMEROOM_FILL_ORDER", "HDTN_FIXED_SLOTS"):
        assert list(rules[attr]) == list(DEFAULTS[attr])


def test_file_without_rules_uses_defaults(tmp_path):
    path = _input(tmp_path)
    _drop_all_rule_columns(path)
    assert read_rules(path) is None


def test_missing_sheet_or_column_keeps_defaults(tmp_path):
    path = _input(tmp_path)
    _drop(path, None)
    _drop(path, MON, "Môn nặng")
    rules = read_rules(path)
    assert "HEAVY_SUBJECTS" not in rules and "HOMEROOM_PERIODS" not in rules and "DAY_SESSIONS" not in rules
    assert rules["MORNING_SUBJECTS"] == DEFAULTS["MORNING_SUBJECTS"]


def test_subject_columns(tmp_path):
    """Có/Không viết kiểu nào cũng được (x, không, trống); GVCN nhận trọn theo thứ tự dòng."""
    warnings = []
    path = _input(tmp_path, [
        (MON, "Tiếng Việt", "Môn nặng", None), (MON, "Thể dục", "Môn nặng", "x"), (MON, "Toán", "Môn nặng", "có"),
        (MON, "Khoa học", "Môn nặng", "không"),
        (MON, "Tiếng Việt", "GVCN nhận trọn", "Không"),
        (MON, "Đạo đức", "Quản lý dạy khối", 5),
        (MON, "Tiếng Anh", "Tối đa tiết mỗi ngày", 2),
        (MON, "Thể dục", "Tên trong TKB", "TD"),
        (MON, "Tiếng Việt", "Môn khó", "Có"),  # cột lạ ở sheet chương trình học: cảnh báo, bỏ qua
    ])
    rules = read_rules(path, warn=warnings.append)
    assert rules["HEAVY_SUBJECTS"] == {config.TOAN, config.TOAN_TC, config.TV_TC, config.TIENG_ANH, config.TIN_HOC,
                                       "Thể dục"}
    assert rules["HOMEROOM_PRIORITY"] == [config.TOAN, config.HDTN, config.KH, config.LSDL, config.DD]
    assert rules["MANAGER_RULES"] == [config.ManagerRule(config.DD, 5), config.ManagerRule(config.KNS, 4)]  # thứ tự dòng
    assert rules["DAILY_LIMITS"] == {config.TOAN: 1, config.TIENG_ANH: 2}
    assert rules["DISPLAY_NAMES"]["Thể dục"] == "TD"
    assert changed(rules) == ["Tên trong TKB", "GVCN nhận trọn", "Quản lý dạy khối", "Tối đa tiết mỗi ngày", "Môn nặng"]
    assert len(warnings) == 1 and "'Môn khó' không phải quy định nào" in warnings[0]


def test_subject_groups(tmp_path):
    """Nhóm môn: môn chính và môn tăng cường cùng số; số nhỏ trước."""
    rules = read_rules(_input(tmp_path, [(MON, "Toán", "Nhóm môn", 1), (MON, "Toán tăng cường", "Nhóm môn", 1),
                                         (MON, "Tiếng Việt", "Nhóm môn", 2),
                                         (MON, "Tiếng Việt tăng cường", "Nhóm môn", 2)]))
    assert list(rules["SUBJECT_GROUPS"].items()) == [(config.TOAN_TC, config.TOAN), (config.TV_TC, config.TV)]


def test_time_frame(tmp_path):
    rules = read_rules(_input(tmp_path, [
        (NGAY, "Thứ 7", "Học buổi sáng", "Có"), (NGAY, "Thứ 4", "Học buổi chiều", "Không"),
        (CHUNG, "Số tiết buổi chiều", None, 4),
        (TIET, 7, "Hạn chế môn nặng", "Không"), (TIET, 8, "Hạn chế môn nặng", "Có"),
    ]))
    assert rules["DAYS"] == ["Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7"]
    assert rules["AFTERNOON"].periods == (5, 6, 7, 8) and rules["MORNING"] == config.MORNING
    sessions = rules["DAY_SESSIONS"]
    assert [d for d, ss in sessions.items() if len(ss) == 2] == [0, 1, 3] and sessions[5] == (rules["MORNING"],)
    assert rules["HEAVY_LATE_PERIODS"] == {8}
    assert changed(rules) == ["Khung giờ", "Hạn chế môn nặng"]


def test_all_errors_at_once(tmp_path):
    with pytest.raises(InputError) as err:
        read_rules(_input(tmp_path, [
            (CHUNG, "Số tiết buổi sáng", None, 0),
            (CHUNG, "Chủ Nhiệm được dạy bù", None, "Có lẽ"),
            (NGAY, "Thứ 4", "Học buổi sáng", "Không"),  # ngày học không liền nhau
            (NGAY, "Thứ 2", "Tiết HĐTN cố định", 9),
            (TIET, 1, "Luôn do GVCN dạy", "Chắc"),
            (TIET, 1, "Môn khó", "Có"),  # cột lạ ở sheet QUY ĐỊNH là lỗi
            (MON, "Toán", "GVCN cắt bớt", 1),  # trùng số với Tiếng Việt
            (MON, "Tin học", "Môn HĐTN", "Có"),  # hai môn HĐTN
            (MON, "Thể dục", "Môn tăng cường", "Có"),  # chưa có Nhóm môn
        ]))
    text = str(err.value)
    parts = ("CHƯƠNG TRÌNH HỌC, dòng 3: cột GVCN cắt bớt: số thứ tự 1 bị lặp (Tiếng Việt)",
             "cột Môn HĐTN chỉ ghi Có ở một môn (đã có Hoạt động trải nghiệm)",
             "môn tăng cường Thể dục chưa ghi Nhóm môn",
             "QUY ĐỊNH, dòng 2: cột Số tiết buổi sáng ghi một số nguyên dương, đang ghi 0",
             "cột Chủ Nhiệm được dạy bù chỉ ghi Có hoặc Không, đang ghi 'Có lẽ'",
             "cột Luôn do GVCN dạy chỉ ghi Có hoặc Không",
             "không có quy định nào tên 'Môn khó'",
             "QUY ĐỊNH: các ngày Học buổi sáng phải liền nhau từ Thứ 2",
             "Xếp tiết HĐTN còn lại vào Thứ 4 nhưng ngày đó không học",
             "Tiết HĐTN cố định Thứ 2 tiết 9 không có trong khung giờ")
    for part in parts:
        assert part in text, part
    assert text.startswith(f"Quy định trong file vào có {len(parts)} lỗi")
    assert text.index("CHƯƠNG TRÌNH HỌC") < text.index("QUY ĐỊNH, dòng")  # sắp theo sheet rồi theo dòng


def test_old_rule_sheets_are_reported(tmp_path):
    path = _input(tmp_path)
    wb = openpyxl.load_workbook(path)
    wb.create_sheet("QUY ĐỊNH MÔN").append(["Môn học", "Môn nặng"])
    wb.save(path)
    with pytest.raises(InputError, match="mẫu cũ"):
        read_rules(path)


def test_applied_restores_config():
    before = config.HEAVY_LATE_PERIODS
    with applied({"HEAVY_LATE_PERIODS": {6}}):
        assert config.HEAVY_LATE_PERIODS == {6}
    assert config.HEAVY_LATE_PERIODS is before
    with applied(None):
        assert config.HEAVY_LATE_PERIODS is before


def test_every_rule_has_a_label():
    assert all(attr in LABELS for attr in ATTRS)


def _code(text):
    return next(line.split()[3] for line in text.splitlines() if line.startswith("Mã kết quả"))


def test_default_rules_give_the_same_timetable(tmp_path, capsys):
    """Có quy định (giá trị mặc định) hay không có thì TKB như nhau."""
    with_rules = _input(tmp_path)
    without = tmp_path / "khong_quy_dinh.xlsx"
    shutil.copy(with_rules, without)
    _drop_all_rule_columns(without)
    codes = []
    for path, note in ((with_rules, "giống mặc định"), (without, "mặc định của chương trình")):
        assert main.run(path, tmp_path / path.stem, thoi_gian_toi_da=10, che_do="tuyen_them") == 0
        out = capsys.readouterr().out
        assert note in out
        codes.append(_code(out))
    assert codes[0] == codes[1]
    # File cập nhật của file vào không ghi quy định thì có thêm các quy định đã dùng.
    updated = tmp_path / without.stem / "khong_quy_dinh_cap_nhat.xlsx"
    assert read_rules(updated) == DEFAULTS
    assert openpyxl.load_workbook(updated).sheetnames == ["NHÂN SỰ", "CHƯƠNG TRÌNH HỌC", "QUY ĐỊNH", "HƯỚNG DẪN",
                                                          "TKB đã xếp"]


def test_rules_from_the_file_are_used(tmp_path, capsys):
    """Thêm sáng Thứ 7, chiều 4 tiết: TKB có cột Thứ 7, kiểm tra luật đạt; chạy xong config trở lại mặc định."""
    path = _input(tmp_path, [(NGAY, "Thứ 7", "Học buổi sáng", "Có"), (CHUNG, "Số tiết buổi chiều", None, 4),
                             (TIET, 7, "Hạn chế môn nặng", "Không"), (TIET, 8, "Hạn chế môn nặng", "Có")])
    assert main.run(path, tmp_path / "out", thoi_gian_toi_da=10, che_do="tuyen_them") == 0
    out = capsys.readouterr().out
    assert "khác mặc định: Khung giờ, Hạn chế môn nặng" in out
    assert "kiểm tra luật bắt buộc: ĐẠT" in out and "Môn nặng ở tiết 8:" in out
    ws = openpyxl.load_workbook(tmp_path / "out" / "TKB.xlsx")["Khối 3"]
    assert [c.value for c in ws[1]][3:] == ["THỨ 2", "THỨ 3", "THỨ 4", "THỨ 5", "THỨ 6", "THỨ 7"]
    assert [ws.cell(r, 3).value for r in range(2, 10)] == list(range(1, 9))  # sáng 1–4, chiều 5–8
    for attr, value in DEFAULTS.items():
        assert getattr(config, attr) == value


def test_changed_rules_in_updated_file_solve_again(tmp_path, capsys):
    """File vào cập nhật lưu mã quy định cùng TKB đã xếp: sửa quy định rồi nạp lại thì xếp lại theo quy định mới,
    không dùng lại TKB cũ."""
    path = _input(tmp_path)
    assert main.run(path, tmp_path / "lan1", thoi_gian_toi_da=10, che_do="tuyen_them") == 0
    updated = tmp_path / "lan1" / "vao_cap_nhat.xlsx"
    ws = openpyxl.load_workbook(updated)["TKB đã xếp"]
    assert (ws["C1"].value, ws["D1"].value) == ("Mã quy định", code())
    wb = openpyxl.load_workbook(updated)
    program = wb[config.PROGRAM_SHEET]
    row = next(r for r in range(2, program.max_row + 1) if program.cell(r, 1).value == "Tiếng Việt")
    col = next(c.column for c in program[1] if c.value == "Môn nặng")
    program.cell(row, col, "Không")
    wb.save(updated)
    capsys.readouterr()
    assert main.run(updated, tmp_path / "lan2", thoi_gian_toi_da=10, che_do="tuyen_them") == 0
    out = capsys.readouterr().out
    assert "khác mặc định: Môn nặng" in out and "xếp lại từ đầu theo quy định mới" in out
    assert "Dùng lại TKB đã xếp" not in out
