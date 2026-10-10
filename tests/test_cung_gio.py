"""Kiểu luật Học cùng giờ (phép đo Cùng giờ của bộ ghép): mỗi lớp học Môn và Môn thứ hai cùng giờ, cùng số tiết, mỗi
môn một người dạy, lớp tính một tiết ở giờ đó (bo_ghep.links, Problem.links). Dùng cho lớp chia nhóm (Âm nhạc và Mỹ
thuật cùng giờ) và dạy kèm (môn riêng Trợ giảng Tiếng Anh cùng giờ với Tiếng Anh). Không có luật này: như trước."""
import dataclasses
from collections import defaultdict

import openpyxl
import pytest

from tkb import bo_mau, config, kich_ban
from tkb.__main__ import main
from tkb.checker import check
from tkb.program import read_program
from tkb.rules import applied, read_rules
from tkb.solver import ConflictError, solve
from tkb.staff import InputError, parse_saved_grid, read_saved_timetable, read_staff
from tkb.template import write_staff_template

from .conftest import CURRICULUM, small_staff, teacher

BLANK = {k: "" for k in kich_ban.RULE_KEYS}
TG = "Trợ giảng Tiếng Anh"
ART = {"kind": "Học cùng giờ", "subject": bo_mau.AM_NHAC, "other": bo_mau.MY_THUAT}
HELP = {"kind": "Học cùng giờ", "subject": bo_mau.TIENG_ANH, "other": TG}
SETTINGS = config.Settings(time_limit=10, workers=4)


def _file(tmp_path, rules=(), staff=None, name="vao", assistant=True):
    """Trường nhỏ (2 lớp Khối 3), `assistant`: có thêm môn Trợ giảng Tiếng Anh (4 tiết, như Tiếng Anh) và GV trợ
    giảng; thêm các dòng luật `rules` (bắt buộc). Ghi qua kịch bản như giao diện."""
    path = tmp_path / f"{name}.xlsx"
    helper = [teacher("TG", "trợ giảng tiếng anh 1", 23, row=10)] if assistant else []
    people = small_staff() + helper if staff is None else staff
    write_staff_template(path, people, {3: {**CURRICULUM[3], **({TG: 4} if assistant else {})}})
    sc, _ = kich_ban.from_excel(path)
    sc["rules"] += [{**BLANK, "hard": True, **r} for r in rules]
    kich_ban.to_excel(sc, path)
    return path


def _read(path):
    curriculum = read_program(path)
    return read_staff(path, subjects=[s for req in curriculum.values() for s in req]), curriculum


@pytest.mark.parametrize("mode", [config.MODE_HIRE, config.MODE_OVERTIME])
def test_split_group_and_assistant_end_to_end(tmp_path, capsys, mode):
    """Âm nhạc cùng giờ Mỹ thuật (chia nhóm), Tiếng Anh cùng giờ Trợ giảng Tiếng Anh (dạy kèm): xếp đạt mọi luật; các
    môn của nhóm ở đúng các giờ như nhau, khác người dạy; ô TKB ghi hai môn; trợ giảng có 8 tiết; nạp lại file cập
    nhật thì dùng lại TKB; giao diện đọc được ô hai môn."""
    path = _file(tmp_path, [ART, HELP])
    argv = [str(path), "-o", str(tmp_path / "TKB.xlsx"), "--time-limit", "10", "--workers", "4", "--mode", mode,
            "--max-overtime", "2"]
    assert main(argv) == 0
    out = capsys.readouterr().out
    assert "kiểm tra luật bắt buộc: ĐẠT" in out and "Cần bổ sung" not in out
    updated = tmp_path / "vao_cap_nhat.xlsx"
    with applied(read_rules(updated)):
        rows = read_saved_timetable(updated).rows
    at = defaultdict(set)
    who = defaultdict(set)
    for cls, day, period, subject, code, *_ in rows:
        at[cls, subject].add((day, period))
        who[cls, subject].add(code)
    for cls in ("3/1", "3/2"):
        assert at[cls, bo_mau.AM_NHAC] == at[cls, bo_mau.MY_THUAT] and len(at[cls, bo_mau.AM_NHAC]) == 1
        assert at[cls, bo_mau.TIENG_ANH] == at[cls, TG] and len(at[cls, TG]) == 4
        assert who[cls, TG] == {"Trợ Giảng Tiếng Anh 1"} and not who[cls, TG] & who[cls, bo_mau.TIENG_ANH]
    cells = [str(c.value) for ws in openpyxl.load_workbook(tmp_path / "TKB.xlsx") for row in ws.iter_rows()
             for c in row if c.value]
    assert sum(v.startswith(f"{bo_mau.AM_NHAC} / {bo_mau.MY_THUAT}\n") for v in cells) == 2
    assert sum(v.startswith(f"{bo_mau.TIENG_ANH} / {TG}\n") for v in cells) == 8
    stats = openpyxl.load_workbook(tmp_path / "Thong_Ke.xlsx")["Thống kê"]
    line = next(r for r in stats.iter_rows(values_only=True) if r[1] and "Trợ Giảng" in str(r[1]))
    assert 8 in line
    again = [str(updated), "-o", str(tmp_path / "lai" / "TKB.xlsx"), "--time-limit", "10", "--workers", "4", "--mode",
             mode, "--max-overtime", "2"]
    assert main(again) == 0 and "Dùng lại TKB đã xếp" in capsys.readouterr().out
    view = kich_ban.timetable(kich_ban.from_excel(updated)[0], mode=mode, overtime_max=2)
    assert view["ok"], view["errors"]
    shared = [c for c in view["cells"] if len(c["codes"]) == 2]
    assert len(shared) == 10 and all(" / " in c["subject"] for c in shared)


def test_rule_errors(tmp_path):
    """Ghi sai: thiếu Môn thứ hai, ưu tiên, có Ngày, Tự ghép không theo Lớp; khác số tiết; cùng một người dạy cả hai."""
    path = _file(tmp_path, [{**ART, "other": ""}, {**ART, "hard": False}, {**ART, "days": "Thứ 2"},
                            {"kind": "Tự ghép", "measure": "Cùng giờ", "scope": "Ngày", "subject": bo_mau.AM_NHAC,
                             "other": bo_mau.MY_THUAT}])
    with pytest.raises(InputError) as err:
        read_rules(path)
    for text in ("dòng 26: kiểu luật Học cùng giờ phải ghi cột Môn thứ hai", "dòng 27: học cùng giờ chỉ ghi Bắt buộc",
                 "dòng 28: cột Ngày không dùng cho kiểu luật Học cùng giờ",
                 "dòng 29: phép đo Cùng giờ: cột Với mỗi ghi Lớp"):
        assert text in str(err.value), str(err.value)
    path = _file(tmp_path, [{**ART, "other": bo_mau.TIN_HOC + ", " + bo_mau.TIENG_ANH}, HELP], name="so_tiet")
    errors = kich_ban.check(kich_ban.from_excel(path)[0], quick=True)["errors"]
    assert any("khối 3 có Âm nhạc 1 tiết, Tin học 1 tiết, Tiếng Anh 4 tiết/tuần: học cùng giờ thì các môn phải cùng "
               "số tiết/tuần" in e for e in errors), errors
    # Không có GV Mỹ thuật; GV Âm nhạc dạy được cả Mỹ thuật: không chia được cho hai người.
    staff = [dataclasses.replace(t, extra_roles=("mỹ thuật",)) if t.role == "âm nhạc" else t
             for t in small_staff() if t.role != "mỹ thuật"] + [teacher("TG", "trợ giảng tiếng anh 1", 23, row=10)]
    path = _file(tmp_path, [ART, HELP], staff=staff, name="mot_nguoi")
    with applied(read_rules(path)):
        with pytest.raises(ConflictError, match="Âm nhạc và Mỹ thuật học cùng giờ .* giao cả hai cho Âm Nhạc 1"):
            solve(*_read(path), SETTINGS, log=lambda *_: None)


def test_checker_and_rules_off(tmp_path):
    """Kiểm tra độc lập: dời tiết Mỹ thuật của 3/1 khỏi giờ Âm nhạc thì báo. Luật Tạm tắt: như không có luật (lớp
    học 32 giờ, mỗi giờ một tiết)."""
    path = _file(tmp_path, [ART], assistant=False)
    with applied(read_rules(path)):
        sol = solve(*_read(path), SETTINGS, log=lambda *_: None)
        assert check(sol.problem, sol.lessons) == [] and sol.problem.links == {
            c: ((bo_mau.AM_NHAC, bo_mau.MY_THUAT),) for c in ("3/1", "3/2")}
        art = next(l for l in sol.lessons if l.class_name == "3/1" and l.subject == bo_mau.MY_THUAT)
        other = next(l for l in sol.lessons if l.class_name == "3/1" and l.subject == bo_mau.TV)
        moved = [dataclasses.replace(l, day=other.day, period=other.period) if l is art else
                 dataclasses.replace(l, day=art.day, period=art.period) if l is other else l for l in sol.lessons]
        errors = check(sol.problem, moved)
    assert any("không có Mỹ thuật cùng giờ" in e for e in errors), errors
    sc, _ = kich_ban.from_excel(path)
    sc["rules"][-1]["off"] = True
    kich_ban.to_excel(sc, path)
    with applied(read_rules(path)):
        staff, curriculum = _read(path)
        assert solve(staff, curriculum, SETTINGS, log=lambda *_: None).problem.links == {}


def test_saved_cell_with_two_subjects():
    """Ô của sheet TKB đã xếp có hai môn: hai cặp "môn", "Mã GV", dấu bù, khóa theo từng cặp; ô một môn như trước."""
    grid = [["Mã kết quả", "X"], [], ["Lớp", "Tiết", "Thứ 2"],
            ["3/1", 1, "Tin học\nTin Học 1 (khóa)\nTiếng Anh\nTiếng Anh 1 (bù) (khóa)"], [None, 2, "Toán\nBộ Môn 1"]]
    saved = parse_saved_grid(grid)
    assert [r[:7] for r in saved.rows] == [
        ("3/1", "Thứ 2", 1, "Tin học", "Tin Học 1", False, True),
        ("3/1", "Thứ 2", 1, "Tiếng Anh", "Tiếng Anh 1", True, True),
        ("3/1", "Thứ 2", 2, "Toán", "Bộ Môn 1", False, False)]
    assert [p[-1] for p in saved.places] == [(0, 1), (2,)]
