"""Khung giờ tự do (tkb/khung_gio.py, bảng Ngày của sheet QUY ĐỊNH): tên ngày, số ngày, tên và số buổi, số tiết từng
buổi của từng ngày đều do trường đặt; cột Buổi Nghỉ, cột Ngày/Buổi của sheet LUẬT đọc theo khung giờ đó."""
import openpyxl
import pytest

from tkb import config, kich_ban, khung_gio
from tkb.__main__ import main
from tkb.rules import applied, read_rules
from tkb.staff import parse_off, read_saved_timetable
from tkb.template import write_staff_template

from .conftest import CURRICULUM, small_staff

NAMES = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat"]
SESSIONS = ["Sáng", "Chiều", "Tối"]
COUNTS = {"Mon": (4, 3, 0), "Tue": (4, 3, 0), "Wed": (4, 3, 2), "Thu": (4, 4, 0), "Fri": (4, 0, 0), "Sat": (4, 0, 0)}


def _free_frame_file(tmp_path):
    """Trường nhỏ với khung giờ 6 ngày tên tiếng Anh, thêm buổi Tối ở Wed, chiều Thu 4 tiết, Sat chỉ học sáng; GV
    tiếng anh nghỉ sáng Sat và tối Wed; luật: Toán không xếp vào Sat, Tiếng Anh không xếp vào buổi Tối."""
    path = tmp_path / "vao.xlsx"
    write_staff_template(path, small_staff(general=False), CURRICULUM)
    sc, _ = kich_ban.from_excel(path)
    sc["sessions"] = SESSIONS
    sc["days"] = [{"name": n, "periods": dict(zip(SESSIONS, COUNTS[n])), "HDTN_FIXED_SLOTS": {"Mon": 1, "Fri": 4}.get(n),
                   "HDTN_FLEX_DAYS": n in ("Tue", "Wed", "Thu")} for n in NAMES]
    while len(sc["periods"]) < 9:
        sc["periods"].append({"HOMEROOM_PERIODS": False, "HEAVY_LATE_PERIODS": False})
    english = next(row for row in sc["staff"] if row["role"] == "Tiếng Anh")
    english["off"] = "Sáng Sat, Tối Wed"
    blank = {k: "" for k in kich_ban.RULE_KEYS}
    sc["rules"] += [{**blank, "kind": "Không xếp vào", "subject": "Toán", "days": "Sat", "hard": True},
                    {**blank, "kind": "Không xếp vào", "subject": "Tiếng Anh", "sessions": "Tối", "hard": True}]
    kich_ban.to_excel(sc, path)
    return path


def test_free_frame_is_read(tmp_path):
    rules = read_rules(_free_frame_file(tmp_path))
    assert rules["DAYS"] == NAMES
    sessions = rules["DAY_SESSIONS"]
    assert [[(s.name, s.periods) for s in sessions[d]] for d in (2, 3, 5)] == [
        [("Sáng", (1, 2, 3, 4)), ("Chiều", (5, 6, 7)), ("Tối", (8, 9))],
        [("Sáng", (1, 2, 3, 4)), ("Chiều", (5, 6, 7, 8))],
        [("Sáng", (1, 2, 3, 4))]]
    with applied(rules):
        assert khung_gio.day_list("Mon-Wed, Sat") == [0, 1, 2, 5] and khung_gio.day_index("thu") == 3
        assert khung_gio.session_at(3, 8).name == "Chiều" and khung_gio.session_at(2, 8).name == "Tối"
        assert khung_gio.max_periods() == 9 and khung_gio.session_names() == SESSIONS
        assert parse_off("Tối Wed, 2 buổi chiều") == (frozenset({(2, "Tối")}), (("Chiều", 2),))
        with pytest.raises(Exception, match="không có ngày học nào"):
            parse_off("Chiều Sun")


def test_free_frame_end_to_end(tmp_path, capsys):
    """Xếp đủ luật với khung giờ đó: luật và buổi nghỉ theo tên ngày, tên buổi mới đều giữ; file TKB có cột từng ngày,
    hàng buổi Tối, tiết 8 của Thu; nạp lại file vào cập nhật thì dùng lại TKB đã xếp (đọc đủ các ngày tên tự đặt)."""
    path = _free_frame_file(tmp_path)
    argv = [str(path), "-o", str(tmp_path / "TKB.xlsx"), "--time-limit", "10", "--workers", "4", "--mode", "tuyen_them"]
    assert main(argv) == 0
    assert "kiểm tra luật bắt buộc: ĐẠT" in capsys.readouterr().out
    rules = read_rules(path)
    with applied(rules):
        saved = read_saved_timetable(tmp_path / "vao_cap_nhat.xlsx")
    lessons = [(cls, config_day, period, subject, teacher)
               for cls, config_day, period, subject, teacher, *_ in saved.rows]
    assert not [row for row in lessons if row[3] == "Toán" and row[1] == "Sat"]
    assert not [row for row in lessons if row[3] == "Tiếng Anh" and row[1] == "Wed" and row[2] >= 8]
    assert {row[1] for row in lessons} == set(NAMES)  # sheet TKB đã xếp đọc lại đủ các ngày tên tự đặt
    english = [row for row in lessons if row[4] == "Tiếng Anh 1"]
    assert english and not [row for row in english if row[1] == "Sat"]
    ws = openpyxl.load_workbook(tmp_path / "TKB.xlsx")["Khối 3"]
    assert [c.value for c in ws[1]][3:] == [n.upper() for n in NAMES]
    labels = [(ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(2, 12)]
    assert [p for _, p in labels] == [1, 2, 3, 4, 5, 6, 7, 8, 8, 9]  # chiều 4 hàng (tiết 8 ở Thu), tối 2 hàng
    assert labels[0][0] == "SÁNG" and labels[4][0] == "CHIỀU" and labels[8][0] == "TỐI"
    assert ws.cell(9, 4).value == config.OFF_LABEL  # hàng tiết 8 buổi chiều: Mon không có
    # Nạp lại file vào cập nhật: dùng lại TKB đã xếp, không xếp lại.
    again = [str(tmp_path / "vao_cap_nhat.xlsx"), "-o", str(tmp_path / "lai" / "TKB.xlsx"), "--time-limit", "10",
             "--workers", "4", "--mode", "tuyen_them"]
    assert main(again) == 0 and "Dùng lại TKB đã xếp" in capsys.readouterr().out
