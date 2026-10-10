"""Sheet PHÒNG (không bắt buộc): phòng học dùng chung. Tiết của môn có phòng phải học ở một phòng hợp (môn, khối, cơ
sở), mỗi giờ không quá sức chứa (tkb/phong_hoc.py: ràng buộc trong mô hình, xếp phòng sau khi xếp giờ, kiểm tra độc
lập). Không có sheet PHÒNG (hay sheet trống): như trước (cùng mã quy định, cùng mã kết quả)."""
import dataclasses
from collections import defaultdict

import openpyxl
import pytest

from tkb import bo_mau, config, kich_ban
from tkb.__main__ import main
from tkb.allocation import build_problem
from tkb.chan_doan import precheck
from tkb.checker import check
from tkb.phong_hoc import assign, labels
from tkb.program import read_program
from tkb.rules import applied, code, read_rules
from tkb.solver import feasible, solve
from tkb.staff import InputError, read_saved_timetable, read_staff
from tkb.template import write_staff_template

from .conftest import CURRICULUM, small_staff

BLANK = {k: "" for k in kich_ban.RULE_KEYS}
ART = {"name": "Phòng nghệ thuật", "subjects": f"{bo_mau.AM_NHAC}, {bo_mau.MY_THUAT}"}
SETTINGS = config.Settings(time_limit=10, workers=4)


def _file(tmp_path, rooms=(), rules=(), name="vao"):
    """Trường nhỏ (2 lớp Khối 3), sheet PHÒNG ghi `rooms` [{name, campus, subjects, grades, capacity}], thêm các dòng
    luật `rules`. Ghi qua kịch bản như giao diện."""
    path = tmp_path / f"{name}.xlsx"
    write_staff_template(path, small_staff(), CURRICULUM)
    sc, _ = kich_ban.from_excel(path)
    sc["rooms"] = [{"campus": "", "grades": "", "capacity": None, **r} for r in rooms]
    sc["rules"] += [{**BLANK, "hard": True, **r} for r in rules]
    kich_ban.to_excel(sc, path)
    return path


def _fixed(subject, cls, day="Thứ 2", period="2"):
    """Dòng luật Cố định vào: môn của lớp học ở giờ đó."""
    return {"kind": "Cố định vào", "subject": subject, "classes": cls, "days": day, "periods": period}


def _problem(path):
    curriculum = read_program(path)
    return build_problem(read_staff(path, subjects=[s for req in curriculum.values() for s in req]), curriculum)


def test_read_rooms_and_rules_code(tmp_path):
    """Đọc sheet PHÒNG; phòng hợp với các tiết theo môn và khối. Sheet trống như không có sheet; thứ tự dòng, hoa
    thường không đổi mã quy định."""
    rooms = [ART, {"name": "Sân", "subjects": "thể dục", "grades": "3", "capacity": 2}]
    path = _file(tmp_path, rooms)
    values = read_rules(path)
    assert values["ROOMS"] == (config.Room("Phòng nghệ thuật", (bo_mau.AM_NHAC, bo_mau.MY_THUAT), "", (), 1, 2),
                               config.Room("Sân", ("thể dục",), "", (3,), 2, 3))
    with applied(values):
        problem = _problem(path)
        with_rooms = code()
    fit = {(problem.courses[cid].class_name, problem.courses[cid].subject): rs for cid, rs in problem.room_fit.items()}
    assert fit == {(c, s): (0,) for c in ("3/1", "3/2") for s in (bo_mau.AM_NHAC, bo_mau.MY_THUAT)} | \
        {(c, "Thể dục"): (1,) for c in ("3/1", "3/2")}
    codes = []
    for rows, name in (((), "trong"), (list(reversed([{**r, "name": r["name"].upper()} for r in rooms])), "dao")):
        with applied(read_rules(_file(tmp_path, rows, name=name))):
            codes.append(code())
    with applied(None):
        assert codes[0] == code() != with_rooms  # sheet trống: như không có sheet (mã cũ)
    assert codes[1] == with_rooms


def test_room_errors(tmp_path):
    path = _file(tmp_path, [ART, {**ART, "name": "phòng NGHỆ THUẬT"}, {"name": "Phòng Tin", "subjects": ""},
                            {"name": "Sân", "subjects": "Thể dục", "capacity": 0},
                            {"name": "Lab", "subjects": "Tin học", "grades": "9"}])
    with pytest.raises(InputError) as err:
        read_rules(path)
    for text in ("PHÒNG, dòng 3: phòng 'phòng NGHỆ THUẬT' đã ghi ở dòng 2",
                 "PHÒNG, dòng 4: phòng 'Phòng Tin' chưa ghi môn ở cột Môn",
                 "PHÒNG, dòng 5: cột Sức chứa ghi một số nguyên dương",
                 "PHÒNG, dòng 6: phòng 'Lab': không có cột Khối 9 ở sheet CHƯƠNG TRÌNH HỌC"):
        assert text in str(err.value), str(err.value)
    # Tên môn, tên cơ sở kiểm khi dựng bài toán: kiểm tra nhanh trên giao diện báo đúng dòng.
    path = _file(tmp_path, [{"name": "Phòng hóa", "subjects": "Hóa học"},
                            {"name": "Sân", "subjects": "Thể dục", "campus": "Điểm Tân Phú"}], name="sai")
    errors = kich_ban.check(kich_ban.from_excel(path)[0], quick=True)["errors"]
    assert "PHÒNG, dòng 2: không có môn hay nhãn môn 'Hóa học' ở sheet CHƯƠNG TRÌNH HỌC" in errors, errors
    assert "PHÒNG, dòng 3: không có lớp nào ở cơ sở 'Điểm Tân Phú' (cột Cơ sở của sheet LỚP)" in errors, errors


def test_shared_room_end_to_end(tmp_path, capsys):
    """Phòng nghệ thuật 1 chỗ cho Âm nhạc và Mỹ thuật (hai GV khác nhau): không giờ nào hai lớp cùng ở phòng; TKB,
    TKB giáo viên ghi tên phòng; file TKB phòng; file vào cập nhật giữ sheet PHÒNG, nạp lại thì dùng lại TKB; giao
    diện ghi phòng ở ô và xem được theo phòng."""
    path = _file(tmp_path, [ART])
    argv = [str(path), "-o", str(tmp_path / "TKB.xlsx"), "--time-limit", "10", "--workers", "4"]
    assert main(argv) == 0
    assert "kiểm tra luật bắt buộc: ĐẠT" in capsys.readouterr().out
    updated = tmp_path / "vao_cap_nhat.xlsx"
    with applied(read_rules(updated)):
        rows = read_saved_timetable(updated).rows
    art = defaultdict(list)
    for cls, day, period, subject, *_ in rows:
        if subject in (bo_mau.AM_NHAC, bo_mau.MY_THUAT):
            art[day, period].append(cls)
    assert len(art) == 4 and all(len(v) == 1 for v in art.values())
    cells = [str(c.value) for ws in openpyxl.load_workbook(tmp_path / "TKB.xlsx") for row in ws.iter_rows()
             for c in row if c.value]
    assert sum(1 for v in cells if v.endswith("\nPhòng nghệ thuật")) == 4
    teachers = openpyxl.load_workbook(tmp_path / "TKB_giao_vien.xlsx").worksheets[0]
    assert any(str(c.value).endswith(f"{bo_mau.AM_NHAC}\nPhòng nghệ thuật") for row in teachers.iter_rows()
               for c in row)
    ws = openpyxl.load_workbook(tmp_path / "TKB_phong.xlsx")["Phòng"]
    assert ws.cell(1, 1).value == "Phòng nghệ thuật: 4 tiết, 1 lớp cùng lúc"
    assert sum(1 for row in ws.iter_rows() for c in row if str(c.value).startswith("3/")) == 4
    room = openpyxl.load_workbook(updated)["PHÒNG"]
    assert [c.value for c in room[2]] == ["Phòng nghệ thuật", None, f"{bo_mau.AM_NHAC}, {bo_mau.MY_THUAT}", None, None]
    assert main([str(updated), "-o", str(tmp_path / "lai" / "TKB.xlsx"), "--time-limit", "10", "--workers", "4"]) == 0
    assert "Dùng lại TKB đã xếp" in capsys.readouterr().out
    view = kich_ban.timetable(kich_ban.from_excel(updated)[0])
    assert view["ok"] and view["rooms"] == ["Phòng nghệ thuật"]
    assert sorted(c["subject"] for c in view["cells"] if c.get("room")) == \
        sorted([bo_mau.AM_NHAC, bo_mau.MY_THUAT] * 2)


def test_rooms_shared_in_part(tmp_path):
    """Phòng đa năng (Âm nhạc, Mỹ thuật, Thể dục) và sân (chỉ Thể dục), mỗi phòng 1 chỗ: Âm nhạc và Thể dục cùng giờ
    được (Thể dục ra sân); Âm nhạc và Mỹ thuật cùng giờ thì không (chẩn đoán chỉ ra sheet PHÒNG); bỏ phòng thì được."""
    rooms = [{"name": "Phòng đa năng", "subjects": f"{bo_mau.AM_NHAC}, {bo_mau.MY_THUAT}, Thể dục"},
             {"name": "Sân", "subjects": "Thể dục"}]
    cases = [(rooms, [_fixed(bo_mau.AM_NHAC, "3/1"), _fixed("Thể dục", "3/2")], True),
             (rooms, [_fixed(bo_mau.AM_NHAC, "3/1"), _fixed(bo_mau.MY_THUAT, "3/2")], False),
             ([], [_fixed(bo_mau.AM_NHAC, "3/1"), _fixed(bo_mau.MY_THUAT, "3/2")], True)]
    for i, (rows, rules, ok) in enumerate(cases):
        path = _file(tmp_path, rows, rules, name=f"vao{i}")
        with applied(read_rules(path)):
            curriculum = read_program(path)
            staff = read_staff(path, subjects=[s for req in curriculum.values() for s in req])
            assert feasible(staff, curriculum, SETTINGS, 30) is ok, (rows, rules)
            if i == 0:
                sol = solve(staff, curriculum, SETTINGS, log=lambda *_: None)
                assert check(sol.problem, sol.lessons) == []
                where = sol.rooms()
                assert where[("3/2", 0, 2, "Thể dục")] == "Sân"  # phòng đa năng đã có lớp học Âm nhạc
                assert where[("3/1", 0, 2, bo_mau.AM_NHAC)] == "Phòng đa năng"


def test_checker_and_precheck(tmp_path):
    """Kiểm tra độc lập: đổi tay để hai lớp cùng cần phòng một chỗ thì báo không đủ phòng. Phép đếm trước khi xếp:
    phòng không đủ chỗ cả tuần."""
    path = _file(tmp_path, [ART])
    with applied(read_rules(path)):
        curriculum = read_program(path)
        staff = read_staff(path, subjects=[s for req in curriculum.values() for s in req])
        sol = solve(staff, curriculum, SETTINGS, log=lambda *_: None)
        lessons = sol.lessons
        music = next(l for l in lessons if l.class_name == "3/1" and l.subject == bo_mau.AM_NHAC)
        art = next(l for l in lessons if l.class_name == "3/2" and l.subject == bo_mau.MY_THUAT)
        other = next(l for l in lessons if l.class_name == "3/2" and (l.day, l.period) == (music.day, music.period))
        moved = [dataclasses.replace(l, day=music.day, period=music.period) if l is art else
                 dataclasses.replace(l, day=art.day, period=art.period) if l is other else l for l in lessons]
        errors = check(sol.problem, moved)
        assert any("không đủ phòng cho 2 tiết 3/1 Âm nhạc, 3/2 Mỹ thuật (các phòng hợp: Phòng nghệ thuật (1 lớp))"
                   in e for e in errors), errors
        assert "" in assign(sol.problem, moved)
    many = {"name": "Phòng chung", "subjects": f"{bo_mau.TV}, {bo_mau.TOAN}, {bo_mau.TIENG_ANH}, {bo_mau.TNXH}"}
    path = _file(tmp_path, [many], name="chat")
    with applied(read_rules(path)):
        found = precheck(_problem(path))
    assert any("cần 36 tiết/tuần ở phòng Phòng chung (sheet PHÒNG) nhưng các phòng đó chỉ chứa 1 lớp × 32 giờ "
               "học = 32 tiết" in line for line in found), found


def test_rooms_per_campus(tmp_path):
    """Phòng ở một cơ sở chỉ dùng cho các lớp ở cơ sở đó; hai cơ sở có phòng cùng tên thì tên kèm cơ sở."""
    from .test_co_so import _file as campus_file
    path = campus_file(tmp_path)
    sc, _ = kich_ban.from_excel(path)
    sc["rooms"] = [{"name": "Sân", "campus": "điểm tân phú", "subjects": "Thể dục", "grades": "", "capacity": None},
                   {"name": "Sân", "campus": "", "subjects": "Thể dục", "grades": "", "capacity": 2}]
    kich_ban.to_excel(sc, path)
    with applied(read_rules(path)):
        problem = _problem(path)
    assert {problem.courses[cid].class_name: rs for cid, rs in problem.room_fit.items()} == {"Lá 1": (0,), "3/1": (1,)}
    assert labels(problem) == ["Sân (Điểm Tân Phú)", "Sân (Cơ sở 1)"]
