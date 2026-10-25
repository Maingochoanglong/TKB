"""Mọi luật là dòng của sheet LUẬT (tkb/luat_co_san.py), xếp bằng bộ ghép: dòng mặc định là giá trị mặc định; sửa số,
điểm, xóa dòng, đổi Bắt buộc có tác dụng; bộ ghép kiểm chéo với bộ kiểm tra độc lập viết tay (tests/kiem_tra_cu.py);
các đường tắt đọc câu bộ ghép ra đúng nghĩa của luật."""
import random
from collections import Counter
from dataclasses import replace

import openpyxl
import pytest

from tkb import allocation, bo_ghep, bo_mau, chan_doan, checker, config, khung_gio, luat_co_san, luat_rieng, solver
from tkb import truong_mau
from tkb.config import CustomRule
from tkb.rules import DEFAULTS, applied, code, read_rules
from tkb.staff import _check_extras
from tkb.template import write_staff_template

from . import kiem_tra_cu
from .conftest import CURRICULUM, plain_rules, small_staff

SETTINGS = dict(time_limit=5, workers=4, overtime_max=4)
NO_FILE_CODE = code()  # mã quy định khi file không ghi gì


def _solve(**values):
    with applied(values):
        sol = solver.solve(small_staff(general=False), CURRICULUM,
                           config.Settings(mode=config.MODE_OVERTIME, **SETTINGS), log=lambda *_: None)
        return sol, checker.check(sol.problem, sol.lessons)


@pytest.fixture(scope="module")
def plain():
    return _solve()[0]


def _hard_rows():
    return [r for r in luat_co_san.default_rows() if r.hard and r.kind == "tu_ghep"]


def _keys(rows) -> set[str]:
    """Các luật có sẵn có dòng ở dạng gốc trong rows."""
    return {n.key for r in rows if (n := luat_co_san.native_of(r)) is not None}


def test_default_rows_are_the_defaults():
    """Mọi luật có sẵn (luật mẫu, luật của file cũ không có sheet LUẬT) xếp bằng bộ ghép. Đó cũng là giá trị mặc định
    của config (tkb/__init__.py): không có file và file ghi các dòng đó cùng một mã quy định. Buổi nghỉ, Cơ sở 2 không
    còn là luật: không có dòng, luôn áp dụng; dòng của chúng ở file cũ bỏ đi."""
    rows = luat_co_san.default_rows()
    assert {r.group_label for r in rows} == set(luat_co_san.GROUPS)
    assert all(luat_co_san.native_of(r) is not None for r in rows)
    values = luat_co_san.apply(rows)
    assert values == {"OFF": frozenset(), "CUSTOM_RULES": rows}
    assert not _keys(rows) & set(luat_co_san.NATIVE_ONLY)
    assert values["CUSTOM_RULES"] == config.CUSTOM_RULES and values["OFF"] == config.OFF
    assert luat_co_san.numbers(rows) == {"DAILY_LIMITS": config.DAILY_LIMITS,
                                         "SESSION_GROUP_LIMIT": config.SESSION_GROUP_LIMIT,
                                         "PAIR_MIN_LESSONS": config.PAIR_MIN_LESSONS}
    old = [*rows, *(r for k in luat_co_san.NATIVE_ONLY for r in luat_co_san.BY_KEY[k].rows())]  # file cũ có hai dòng đó
    assert luat_co_san.apply(old) == values and luat_co_san.apply([]) == {"OFF": frozenset(), "CUSTOM_RULES": []}
    with applied(values):
        assert code() == NO_FILE_CODE


def test_template_sheet_reads_back_to_defaults(tmp_path):
    path = tmp_path / "vao.xlsx"
    write_staff_template(path, small_staff(), CURRICULUM)
    wb = openpyxl.load_workbook(path)
    assert luat_rieng.SHEET not in wb.sheetnames and luat_rieng.RULES_SHEET in wb.sheetnames
    ws = wb[luat_rieng.RULES_SHEET]
    head = [c.value for c in ws[1]]
    assert head[0] == "Nhóm" and head[-1] == "Luật đọc là"
    assert ws.cell(2, len(head)).value.startswith("Mỗi buổi, một lớp học tối đa 2 tiết của một môn")
    rules = read_rules(path)
    assert len(rules.pop("RULES")) == len(luat_co_san.default_rows())
    assert plain_rules(rules) == DEFAULTS
    with applied(read_rules(path)):
        assert code() == NO_FILE_CODE


def _edit(tmp_path, change):
    """File mẫu với sheet LUẬT sửa theo change(ws, dòng tiêu đề) rồi đọc lại."""
    path = tmp_path / "vao.xlsx"
    write_staff_template(path, small_staff(), CURRICULUM)
    wb = openpyxl.load_workbook(path)
    ws = wb[luat_rieng.RULES_SHEET]
    change(ws, [c.value for c in ws[1]])
    wb.save(path)
    return read_rules(path)


def _row_of(ws, head, text):
    say = head.index("Luật đọc là") + 1
    return next(r for r in range(2, ws.max_row + 1) if (ws.cell(r, say).value or "").startswith(text))


def test_edit_number_points_delete_and_soften(tmp_path):
    def change(ws, head):
        col = {h: i + 1 for i, h in enumerate(head)}
        ws.cell(_row_of(ws, head, "Mỗi buổi, một lớp học tối đa"), col["Số"], 3)
        ws.cell(_row_of(ws, head, "Tránh xếp môn có nhãn Môn nặng"), col["Mức"], "Rất cao")
        ws.delete_rows(_row_of(ws, head, "Ở giờ có nhãn Luôn do GVCN dạy"))
        r = _row_of(ws, head, "Các tiết của một môn trong cùng buổi")
        ws.cell(r, col["Bắt buộc"], "Không")
        ws.cell(r, col["Mức"], 3)

    rules = _edit(tmp_path, change)
    custom = rules["CUSTOM_RULES"]
    assert next(r for r in custom if luat_co_san.native_of(r).key == "nhom_buoi").number == 3
    assert _keys(custom) == _keys(config.CUSTOM_RULES) - {"tiet_gvcn", "lien_nhau"}  # dòng xóa: luật tắt
    heavy = next(r for r in custom if r.measure == "vi_tri" and r.op == "ngoai")
    assert (heavy.points, heavy.level) == (None, 4)  # đổi mức: điểm chung của mức Rất cao
    (soft,) = [r for r in custom if luat_co_san.native_of(r) is None]  # đổi thành ưu tiên: khác dạng gốc
    assert soft.measure == "lien" and not soft.hard and soft.level == 3
    with applied(rules):
        assert code() != NO_FILE_CODE
        assert luat_rieng.label(soft) == (f"LUẬT dòng {soft.row}: Mỗi lớp, mỗi môn: các tiết trong một buổi đứng "
                                          f"liền nhau (ưu tiên cao)")  # khác dạng gốc: câu của bộ ghép


def test_levels_keep_the_tuned_points_of_built_in_rules(tmp_path):
    """Luật ưu tiên có sẵn ghi Mức bằng chữ, không có cột Điểm: để nguyên mức mặc định thì dùng điểm đã chỉnh sẵn
    (cùng mã quy định với khi không có file); đổi mức thì theo điểm chung của mức, đổi lại thì về điểm cũ. File cũ có
    cột Điểm vẫn đọc được, và ghi lại thì giữ cột đó."""
    path = tmp_path / "mau.xlsx"
    write_staff_template(path, small_staff(), CURRICULUM)
    ws = openpyxl.load_workbook(path)[luat_rieng.RULES_SHEET]
    head = [c.value for c in ws[1]]
    assert "Điểm" not in head and "Mức" in head
    gaps = _row_of(ws, head, "Giáo viên, trừ GV chủ nhiệm, không có tiết trống")
    assert ws.cell(gaps, head.index("Mức") + 1).value == "Thấp"
    with applied(read_rules(path)):
        assert code() == NO_FILE_CODE

    def level(word):
        def change(ws, head):
            ws.cell(_row_of(ws, head, "Giáo viên, trừ GV chủ nhiệm, không có tiết trống"), head.index("Mức") + 1, word)
        return next(r for r in _edit(tmp_path, change)["CUSTOM_RULES"] if r.measure == "khoang_cach"
                    and r.op == "trong_tiet")
    assert (level("Vừa").points, level("Vừa").level) == (None, 2)
    assert level("Thấp").points == 10

    def old_points(ws, head):
        ws.cell(1, len(head) + 1, "Điểm")
        ws.cell(_row_of(ws, head, "Tránh xếp môn có nhãn Môn nặng"), len(head) + 1, 2000)
    rules = _edit(tmp_path, old_points)
    assert next(r for r in rules["CUSTOM_RULES"] if r.measure == "vi_tri" and r.op == "ngoai").points == 2000
    with applied(rules):
        out = tmp_path / "ghi_lai.xlsx"
        write_staff_template(out, small_staff(), CURRICULUM)
    head = [c.value for c in openpyxl.load_workbook(out)[luat_rieng.RULES_SHEET][1]]
    assert "Điểm" in head  # có dòng ghi điểm riêng: giữ cột


def test_rule_switched_off_is_kept_but_not_used(tmp_path, plain):
    """Cột Tạm tắt = Có: chương trình xếp như khi xóa dòng (cùng mã quy định), dòng vẫn còn (đọc lại, ghi ra file cập
    nhật, bật lại được); luật thêm vào cũng vậy. Bật lại (để trống) thì như mặc định."""
    def change(ws, head, off="Có"):
        col = {h: i + 1 for i, h in enumerate(head)}
        ws.cell(_row_of(ws, head, "Ở giờ có nhãn Luôn do GVCN dạy"), col["Tạm tắt"], off)
        r = ws.max_row + 1
        for key, value in (("Kiểu luật", "Không xếp vào"), ("Môn", "Thể dục"), ("Tiết", 1), ("Bắt buộc", "Có"),
                           ("Tạm tắt", off)):
            ws.cell(r, col[key], value)

    rules = _edit(tmp_path, change)
    assert _keys(rules["CUSTOM_RULES"]) == _keys(config.CUSTOM_RULES) - {"tiet_gvcn"}  # cả luật thêm vào cũng tắt
    kept = [r for r in rules["RULES"] if r.off]
    assert len(kept) == 2 and luat_co_san.native_of(kept[0]).key == "tiet_gvcn"  # vẫn nhận ra luật có sẵn
    assert luat_rieng.cells(kept[1])["off"] == "Có"
    deleted = _edit(tmp_path, lambda ws, head: ws.delete_rows(_row_of(ws, head, "Ở giờ có nhãn Luôn do GVCN dạy")))
    with applied(rules):
        off_code = code()
    with applied(deleted):
        assert code() == off_code  # tạm tắt = xóa dòng
    with applied(_edit(tmp_path, lambda ws, head: change(ws, head, off="Không"))) as _:
        assert code() != NO_FILE_CODE  # bật lại cả luật thêm vào: luật đó có hiệu lực
    with applied(_edit(tmp_path, lambda ws, head: ws.cell(_row_of(ws, head, "Ở giờ"), head.index("Tạm tắt") + 1, None))):
        assert code() == NO_FILE_CODE
    from tkb.writer import quality_rows
    with applied(rules):  # sheet Chất lượng: dòng tạm tắt ghi "tắt", không đếm
        marks = {row[1]: row[3] for row in quality_rows(plain)}
        assert marks[luat_rieng.describe(kept[0]).rsplit(" (", 1)[0]] == "tắt"
        assert luat_rieng.describe(kept[1]).rsplit(" (", 1)[0] in marks


def test_soft_rule_off_drops_its_points(tmp_path):
    """Xóa hay tạm tắt dòng một luật ưu tiên có sẵn thì nó ra khỏi mục tiêu xếp giờ, không chỉ đánh dấu tắt; các luật
    ưu tiên khác giữ mức ghi ở dòng."""
    def change(ws, head):
        ws.delete_rows(_row_of(ws, head, "Giáo viên, trừ GV chủ nhiệm, không có tiết trống"))
        ws.cell(_row_of(ws, head, "Tránh xếp môn có nhãn Môn nặng"), head.index("Tạm tắt") + 1, "Có")
        ws.cell(_row_of(ws, head, "Môn có nhãn Ưu tiên buổi sáng xếp vào"), head.index("Mức") + 1, "Cao")

    rules = _edit(tmp_path, change)
    assert _keys(rules["CUSTOM_RULES"]) == _keys(config.CUSTOM_RULES) - {"tiet_trong", "mon_nang"}
    with applied(rules):
        assert {luat_co_san.native_of(L.rule).key for L in bo_ghep.compiled()} == _keys(rules["CUSTOM_RULES"])
        morning = next(L for L in bo_ghep.compiled() if luat_co_san.native_of(L.rule).key == "buoi_sang")
        assert bo_ghep.weight(morning, config.Weights()) == config.Weights().custom_levels[2]  # mức Cao


def test_built_in_rules_read_as_plain_sentences():
    """Mỗi luật có sẵn có tên viết tay; tên theo số của dòng; câu không có ký hiệu hay chữ kỹ thuật."""
    rows = luat_co_san.default_rows()
    titles = [luat_co_san.title(r) for r in rows]
    assert all(titles) and len(set(titles)) == len(titles)
    assert all(len(n.titles) == len(n.rows()) or n.key == "toi_da_ngay" for n in luat_co_san.NATIVES)
    first = replace(rows[0], number=3)
    assert luat_co_san.title(first).startswith("Mỗi buổi, một lớp học tối đa 3 tiết")
    assert luat_co_san.title(replace(rows[0], hard=False)) is None  # khác dạng gốc: đọc bằng câu của bộ ghép
    for r in rows:
        for rule in (r, replace(r, hard=not r.hard) if r.kind == "tu_ghep" else r):
            text = luat_rieng.describe(rule)
            assert not any(s in text for s in ("<=", ">=", " ô ", "Với mỗi", "mức ")), text
            assert text.count(":") <= 1, text
    levels = [luat_rieng.level_label(r) for r in rows if not r.hard]
    assert levels[:2] == ["Thấp", "Rất cao"]  # 200 điểm: gần Thấp; 10000 điểm: Rất cao


def test_deleted_rule_is_off_when_solving():
    rows = [r for r in luat_co_san.default_rows() if luat_co_san.native_of(r).key != "tiet_gvcn"]
    sol, errors = _solve(**luat_co_san.apply(rows))
    assert errors == []
    with applied({}):  # đủ các dòng mặc định: tiết 1 có thể không do GVCN dạy
        found = checker.check(sol.problem, sol.lessons)
    assert all("lớp học với GV chủ nhiệm" in e for e in found)


def test_legacy_file_rows(tmp_path):
    """File bản trước: số ở các cột cũ, luật ở sheet LUẬT RIÊNG -> các dòng mặc định với số đó, cộng luật riêng."""
    path = tmp_path / "vao.xlsx"
    write_staff_template(path, small_staff(), CURRICULUM)
    wb = openpyxl.load_workbook(path)
    del wb[luat_rieng.RULES_SHEET]
    ws = wb.create_sheet(luat_rieng.SHEET)
    ws.append(["Kiểu luật", "Môn", "Tiết", "Bắt buộc"])
    ws.append(["Không xếp vào", "Thể dục", "1", "Có"])
    rules_ws = wb[config.RULES_SHEET]
    rules_ws.insert_rows(2)
    rules_ws.cell(2, 1, "Số tiết tối đa một nhóm môn mỗi buổi")
    rules_ws.cell(2, 2, 3)
    program = wb[config.PROGRAM_SHEET]
    col = program.max_column + 1
    program.cell(1, col, "Tối đa tiết mỗi ngày")
    row = next(r for r in range(2, program.max_row + 1) if program.cell(r, 1).value == "Tiếng Anh")
    program.cell(row, col, 2)
    wb.save(path)
    rules = read_rules(path)
    assert "RULES" not in rules and rules["SESSION_GROUP_LIMIT"] == 3
    assert rules["DAILY_LIMITS"] == {bo_mau.TIENG_ANH: 2}  # cột cũ có thì là đủ (Toán để trống: không giới hạn)
    (own,) = [r for r in rules["CUSTOM_RULES"] if luat_co_san.native_of(r) is None]
    assert own.kind == "khong_xep"
    with applied(rules):
        rows = luat_co_san.rows()
        assert rows[0].number == 3  # dòng mặc định lấy số từ cột cũ
        assert len(rows) == len(luat_co_san.default_rows()) + 1 and rows[-1] == replace(
            own, group_label=luat_co_san.CUSTOM_GROUP)  # mỗi dòng một lần, theo thứ tự luật có sẵn rồi luật riêng
        assert luat_rieng.label(own).startswith("LUẬT RIÊNG dòng 2")


def test_solved_timetable_keeps_every_hard_row(plain):
    assert checker.check(plain.problem, plain.lessons) == []
    found = bo_ghep.violations(plain.problem, plain.lessons, rules=_hard_rows(), skip_forced=False)
    assert found == [], [luat_rieng.describe(L.rule) + ": " + t for L, _, _, t in found]


def _perturb(rnd: random.Random, problem, lessons):
    """Đổi chỗ một, hai cặp tiết của một lớp; nửa số lần đổi thêm người dạy một tiết."""
    les = list(lessons)
    for _ in range(rnd.randint(1, 2)):
        cls = rnd.choice(problem.classes)
        i, j = rnd.sample([k for k, l in enumerate(les) if l.class_name == cls], 2)
        a, b = les[i], les[j]
        les[i] = replace(a, day=b.day, period=b.period)
        les[j] = replace(b, day=a.day, period=a.period)
    if rnd.random() < 0.5:
        i = rnd.randrange(len(les))
        les[i] = replace(les[i], teacher=rnd.choice(sorted(problem.teachers)))
    return les


@pytest.fixture(scope="module")
def campuses(tmp_path_factory):
    """Ba cơ sở, lớp Lá 1 không có GVCN (tests/test_co_so.py), xếp xong."""
    from tkb.program import read_program
    from tkb.staff import read_staff

    from .test_co_so import _file
    path = _file(tmp_path_factory.mktemp("co_so"))
    values = read_rules(path)
    with applied(values):
        curriculum = read_program(path)
        staff = read_staff(path, subjects=[s for req in curriculum.values() for s in req])
        sol = solver.solve(staff, curriculum, config.Settings(time_limit=5, workers=4), log=lambda *_: None)
    return values, sol


@pytest.mark.parametrize("school", ["nho", "co_so"])
def test_composer_matches_the_old_checker(school, plain, campuses):
    """Đổi chỗ ngẫu nhiên: với từng luật có sẵn, bộ ghép thấy dòng của luật bị vi phạm khi và chỉ khi bộ kiểm tra viết
    tay (tests/kiem_tra_cu.py) thấy lỗi của luật đó; checker (kiểm bằng bộ ghép) báo lỗi đúng lúc đó."""
    values, sol = ({}, plain) if school == "nho" else campuses
    rows = _hard_rows()
    rnd = random.Random(0)
    hits = Counter()
    with applied(values):
        keys = {luat_co_san.native_of(r).key for r in rows}
        assert checker.check(sol.problem, sol.lessons) == []
        for _ in range(80):
            les = _perturb(rnd, sol.problem, sol.lessons)
            old = {k for k, found in kiem_tra_cu.errors(sol.problem, les).items() if found and k in keys}
            new = {luat_co_san.native_of(L.rule).key
                   for L, *_ in bo_ghep.violations(sol.problem, les, rules=rows, skip_forced=False)}
            assert new == old, (sorted(old), sorted(new))
            assert bool(old) <= bool(checker.check(sol.problem, les))
            hits.update(old)
    assert len(hits) >= (10 if school == "nho" else 12), hits  # mọi luật có sẵn bắt buộc (một cơ sở: trừ co_so)


@pytest.mark.parametrize("school", ["nho", "co_so"])
def test_solved_timetable_passes_the_hand_written_checks(school, plain, campuses):
    """TKB xếp bằng bộ ghép đúng mọi luật có sẵn bắt buộc theo bộ kiểm tra viết tay (tests/kiem_tra_cu.py)."""
    values, sol = ({}, plain) if school == "nho" else campuses
    with applied(values):
        assert {k: v for k, v in kiem_tra_cu.errors(sol.problem, sol.lessons).items() if v} == {}


def _with(key, **changes):
    """Giá trị config: các dòng mặc định, dòng của luật có sẵn `key` sửa bằng `changes`."""
    rows = [replace(r, row=i + 2, **(changes if luat_co_san.native_of(r).key == key else {}))
            for i, r in enumerate(luat_co_san.default_rows())]
    return luat_co_san.apply(rows)


@pytest.mark.parametrize("school", ["nho", "mau"])
def test_shortcuts_read_composer_rows(school):
    """Các đường tắt đọc câu bộ ghép của dòng HĐTN, GVCN, ghép cặp ra đúng nghĩa của luật (như bộ kiểm tra viết tay):
    course và ô HĐTN cố định, ô lớp chỉ học với GVCN (cắt khỏi course không do GVCN dạy), môn chỉ GVCN dạy, nhóm ghép
    cặp khi phân công."""
    staff, curriculum = (small_staff(), CURRICULUM) if school == "nho" else \
        (truong_mau.sample_staff(), truong_mau.CURRICULUM)
    problem = allocation.build_problem(staff, curriculum, {})
    hdtn = [c for c in problem.courses if c.subject == config.HDTN and c.fixed_slots]
    assert hdtn and all(c.fixed_slots == tuple(sorted(config.HDTN_FIXED_SLOTS)) for c in hdtn)
    first = [s for s in problem.slots if s[1] in config.HOMEROOM_PERIODS]
    assert allocation.homeroom_periods_slots(problem.slots) == first
    for c in problem.courses:  # không ai trong số người dạy được là GVCN của lớp
        if all(problem.teachers[g].class_name != c.class_name for g in c.teachers):
            assert not set(solver.allowed_slots(c, problem)) & set(first)
    assert allocation.homeroom_only() == set(config.HOMEROOM_ONLY_SUBJECTS) != set()
    pairs = {g: allocation.paired_groups(req, g, assignment=True) for g, req in curriculum.items()}
    assert pairs == {g: kiem_tra_cu._pairs(req) for g, req in curriculum.items()} and any(pairs.values())
    with applied({"CUSTOM_RULES": [r for r in config.CUSTOM_RULES if luat_co_san.native_of(r).key != "tiet_gvcn"]}):
        assert allocation.homeroom_periods_slots(problem.slots) == []  # xóa dòng: không còn ô nào


def test_composer_precheck_counts():
    """Phép đếm trước khi xếp: dòng "tối đa n tiết mỗi nhóm môn mỗi buổi" báo khi số tiết quá sức chứa; dòng Theo cặp
    mỗi buổi báo khi số cặp quá số buổi (trường mẫu: 32 tiết mỗi lớp, 16 cặp, 9 buổi)."""
    problem = allocation.build_problem(truong_mau.sample_staff(), truong_mau.CURRICULUM, {})
    assert chan_doan.precheck(problem) == []
    with applied(_with("nhom_buoi", number=1)):
        assert any("tối đa 1 × 9 buổi" in e for e in chan_doan.precheck(problem))
    pair = CustomRule("tu_ghep", scope=("lop", "buoi"), measure="cap", hard=True, row=2)  # mọi môn
    with applied({"CUSTOM_RULES": [pair]}):
        found = chan_doan.precheck(problem)
    assert any("lớp 1/1 có 32 tiết/tuần, cần 16 buổi có từ 2 tiết nhưng khung giờ chỉ có 9 buổi" in e for e in found)


def test_homeroom_leave_follows_the_rule_row():
    """GVCN không nghỉ được buổi có giờ luôn do GVCN dạy (dòng luật "Ở giờ có nhãn Luôn do GVCN dạy…"); xóa dòng thì
    nghỉ được."""
    cn = replace(small_staff(general=False)[0], off_sessions={(0, khung_gio.first_session_name())})
    assert any("không nghỉ buổi sáng được vì tiết 1 luôn do GVCN dạy" in e for e in _check_extras([cn]))
    with applied({"CUSTOM_RULES": [r for r in config.CUSTOM_RULES if luat_co_san.native_of(r).key != "tiet_gvcn"]}):
        assert _check_extras([cn]) == []
