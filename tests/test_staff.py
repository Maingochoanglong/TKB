import openpyxl
import pytest

from tkb import bo_mau, config
from tkb.program import canonical_subject, read_program, subject_key
from tkb.staff import InputError, class_list, read_staff, validate

from .conftest import CURRICULUM, INPUT_FILE, teacher


def test_bad_lessons():
    with pytest.raises(InputError):
        teacher("X", "bộ môn 1", -1)
    with pytest.raises(InputError):
        teacher("X", "bộ môn 1", "abc")
    assert teacher("X", "bộ môn 1", 23.0).max_lessons == 23


def test_duplicates_rejected():
    with pytest.raises(InputError):
        validate([teacher("A", "chủ nhiệm 1/1", 19), teacher("B", "chủ nhiệm 1/1", 16)])
    with pytest.raises(InputError):
        validate([teacher("A", "chủ nhiệm 1/1", 19), teacher("B", "bộ môn 2", 23), teacher("C", "bộ môn 2", 19)])


def test_read_sample_staff(sample_staff):
    assert len(sample_staff) == 45
    classes = class_list(sample_staff)
    assert len(classes) == 29
    assert classes[0] == "1/1" and classes[-1] == "5/5"
    cn55 = next(t for t in sample_staff if t.class_name == "5/5")
    assert cn55.title == "chủ nhiệm 5/5" and cn55.max_lessons == 16
    bm5 = next(t for t in sample_staff if t.title == "bộ môn 5")
    assert bm5.max_lessons == 19


def test_program_file_is_read_as_written():
    assert list(CURRICULUM) == [1, 2, 3, 4, 5]
    assert CURRICULUM[1]["Tiếng Việt"] == 14 and CURRICULUM[3]["Tin học"] == 1
    assert all(sum(req.values()) == 32 for req in CURRICULUM.values())
    assert read_program(INPUT_FILE) == CURRICULUM and list(read_program(INPUT_FILE)[1]) == list(CURRICULUM[1])


def test_subject_names_match_rules_loosely():
    assert subject_key("Lịch Sử và Địa Lý") == subject_key("Lịch sử - Địa lý") == "lịch sử địa lý"
    assert canonical_subject(" Tự Nhiên và Xã Hội ") == bo_mau.TNXH
    assert canonical_subject("HĐTN") == config.HDTN
    assert canonical_subject("Toán  Tăng Cường") == bo_mau.TOAN_TC
    assert canonical_subject("Múa  dân gian") == "Múa dân gian"  # môn không có luật: giữ tên trong file


# Sheet NHÂN SỰ mẫu V8 có thêm một cột ghi chú (các cột khác bị bỏ qua, ghi gì cũng được).
V8 = ("Họ và Tên", "Chức Vụ", "Lớp", "Ghi chú", "Số Tiết/Tuần")


def _staff_file(path, rows, header=V8):
    wb = openpyxl.Workbook()
    wb.active.title = "NHÂN SỰ"
    wb.active.append(list(header))
    for r in rows:
        wb.active.append(list(r))
    wb.save(path)
    return path


def test_columns_and_auto_numbering(tmp_path):
    path = _staff_file(tmp_path / "ns.xlsx", [
        ("A", "Chủ Nhiệm", "1/1", None, 19), ("B", "Bộ Môn", None, None, 23),
        ("C", "Tiếng Anh", None, None, 23), ("D", "bộ môn", None, "đang nghỉ", 19),
        ("E", "Chủ Nhiệm", " 1 / 2 ", "Có", 16), ("F", "Bộ Môn", None, None, 23)])
    got = [(t.name, t.title, t.max_lessons) for t in read_staff(path)]
    assert got == [("A", "chủ nhiệm 1/1", 19), ("B", "bộ môn 1", 23), ("C", "tiếng anh 1", 23),
                   ("D", "bộ môn 2", 19), ("E", "chủ nhiệm 1/2", 16), ("F", "bộ môn 3", 23)]


@pytest.mark.parametrize("title", ["Bộ Môn 5", "chủ nhiệm 1/1", "Tiếng Anh 2"])
def test_numbered_titles_are_rejected(tmp_path, title):
    # Mẫu cũ (V5–V7) ghi số thứ tự trong Chức Vụ; mẫu V8 không ghi số, chương trình tự đánh số.
    path = _staff_file(tmp_path / "ns.xlsx", [("A", title, None, None, 19)])
    with pytest.raises(InputError, match="không ghi số thứ tự"):
        read_staff(path)


def test_old_headers_are_rejected(tmp_path):
    path = _staff_file(tmp_path / "ns.xlsx", [("A", "Bộ Môn", 23)], header=("Tên", "Chức vụ", "Số tiết"))
    with pytest.raises(InputError, match="mẫu V8"):
        read_staff(path)


@pytest.mark.parametrize("row, message", [
    (("A", "Chủ Nhiệm", None, None, 19), "phải ghi Lớp"),
    (("A", "Bộ Môn", "1/1", None, 23), "chỉ Chủ Nhiệm mới ghi Lớp"),
    (("A", "Chủ Nhiệm", "1-1", None, 19), "khối/số thứ tự"),
])
def test_class_errors(tmp_path, row, message):
    path = _staff_file(tmp_path / "ns.xlsx", [row])
    with pytest.raises(InputError, match=message):
        read_staff(path)


def test_class_turned_into_date(tmp_path):
    import datetime
    path = _staff_file(tmp_path / "ns.xlsx", [("A", "Chủ Nhiệm", datetime.datetime(2026, 1, 1), None, 19)])
    with pytest.raises(InputError, match="ngày tháng"):
        read_staff(path)


def test_program_sheet_aliases_and_total_row(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = "NHÂN SỰ"
    ws = wb.create_sheet("CHƯƠNG TRÌNH HỌC")
    ws.append(["SST", "Môn học", "Khối 1"])
    ws.append([1, "Tự Nhiên và Xã Hội", 2])
    ws.append([2, "Lịch Sử và Địa Lý", 0])
    ws.append([3, "Hoạt Động Trải Nghiệm", 3])
    ws.append([None, "Tổng", "=SUM(C2:C4)"])
    ws.append([4, "Toán Nâng Cao", 1])  # môn lạ vẫn nhận
    wb.save(tmp_path / "in.xlsx")
    assert read_program(tmp_path / "in.xlsx") == {
        1: {"Tự Nhiên và Xã Hội": 2, "Lịch Sử và Địa Lý": 0, "Hoạt Động Trải Nghiệm": 3, "Toán Nâng Cao": 1}}
    ws.append([5, "Lịch sử - Địa lý", 1])  # cùng môn, khác cách viết
    wb.save(tmp_path / "in.xlsx")
    with pytest.raises(InputError, match="bị lặp với dòng 3"):
        read_program(tmp_path / "in.xlsx")


def test_missing_program_sheet_is_an_error(tmp_path):
    path = _staff_file(tmp_path / "ns.xlsx", [("A", "Chủ Nhiệm", "1/1", None, 19)])
    with pytest.raises(InputError, match="thiếu sheet CHƯƠNG TRÌNH HỌC"):
        read_program(path)


def test_all_errors_at_once(tmp_path):
    path = _staff_file(tmp_path / "ns.xlsx", [
        ("A", "Chủ Nhiệm", "1/1", None, 19), ("B", "T. Anh ", None, None, 23),
        ("C", "Chủ Nhiệm", "1/1", None, 19), ("D", "Chủ Nhiệm", "2/1", None, 19),
        ("E", "Chủ Nhiệm", "2/1", None, 19), ("F", "Hiệu Trưởng", None, None, 4)])
    with pytest.raises(InputError) as exc:
        read_staff(path, subjects=["Tiếng Việt", "Tiếng Anh"])
    message = str(exc.value)
    assert "4 lỗi" in message and "Lớp 1/1 có hai Chủ Nhiệm" in message and "Lớp 2/1" in message
    # Chức vụ phải là Chủ Nhiệm/Bộ Môn/Quản Lý hoặc đúng tên một môn trong chương trình học.
    assert "Dòng 3: chức vụ 'T. Anh' không xác định" in message and "'Hiệu Trưởng'" in message
    ok = _staff_file(tmp_path / "ok.xlsx", [("A", "Chủ Nhiệm", "1/1", None, 19),
                                            ("B", "Tiếng Anh ", None, None, 23)])
    b = read_staff(ok, subjects=["Tiếng Việt", "Tiếng Anh"])[1]
    assert (b.title, b.label) == ("tiếng anh 1", "Tiếng Anh")


# Mẫu của trường có thêm các cột không bắt buộc.
EXTRA = ("Họ và Tên", "Chức Vụ", "Lớp", "Số Tiết/Tuần", "Thai Sản", "Hợp Đồng", "Cơ sở 2", "Lớp Đang Dạy", "Buổi Nghỉ")


def test_named_classes_and_optional_columns(tmp_path):
    path = _staff_file(tmp_path / "ns.xlsx", [
        ("A", "Chủ Nhiệm", "1D15", 19, None, "Có", None, None, None),
        ("B", "Chủ Nhiệm", "1d9", 15, "có", None, "x", None, "2 buổi chiều"),
        ("C", "Chủ Nhiệm", "2D16", 19, None, None, None, None, None),
        ("D", "Bộ Môn", None, 19, "Có", None, None, "1D15; 2D16", "Chiều thứ 5, sáng T6, chiều T6"),
        ("E", "Tiếng Anh", None, 23, None, "Không", None, "2D16", None)], header=EXTRA)
    ts = read_staff(path)
    assert class_list(ts) == ["1D9", "1D15", "2D16"]  # sắp theo khối rồi số
    a, b, c, d, e = ts
    assert a.contract and not a.maternity and a.title == "chủ nhiệm 1D15" and a.grade == 1
    assert b.maternity and b.campus2 and b.off_any == (("Chiều", 2),) and not b.campus2_only
    assert d.maternity and d.campus2_only and d.history == {"1D15", "2D16"}
    assert d.off_sessions == {(3, "Chiều"), (4, "Sáng")}  # chiều Thứ 6 vốn nghỉ nên bỏ qua
    assert not e.contract and not e.campus2_only and e.history == {"2D16"} and not c.history


@pytest.mark.parametrize("row, message", [
    (("A", "Bộ Môn", None, 23, "không rõ", None, None, None, None), "chỉ ghi 'Có'"),
    (("A", "Bộ Môn", None, 23, None, None, None, "9D1", None), "Lớp Đang Dạy có lớp không có Chủ Nhiệm nào: 9D1"),
    (("A", "Bộ Môn", None, 23, None, None, None, None, "nghỉ cả tuần"), "Buổi Nghỉ ghi buổi cố định"),
    (("A", "Bộ Môn", None, 23, None, None, None, None, "Chiều T5, 4 buổi chiều"), "chỉ còn 3 buổi chiều"),
    (("A", "Chủ Nhiệm", "1D2", 19, None, None, None, None, "Sáng T2"), "không nghỉ buổi sáng"),
    (("A", "Chủ Nhiệm", "1D2", 19, "Có", None, None, None, None), "không đánh dấu Cơ sở 2"),
])
def test_optional_column_errors(tmp_path, row, message):
    path = _staff_file(tmp_path / "ns.xlsx", [("CN", "Chủ Nhiệm", "1D1", 19, None, None, "Có", None, None), row],
                       header=EXTRA)
    with pytest.raises(InputError, match=message):
        read_staff(path)
